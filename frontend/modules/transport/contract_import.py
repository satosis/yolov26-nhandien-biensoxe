"""Service to import and synchronize transport contracts from Excel workbook.

Handles:
1. Updating customer companies (DSKH sheet): representatives, addresses, bank accounts.
2. Upserting contracts (DATA sheet):
   - If contract exists: updates with new content/value/date/notes.
   - If contract does not exist: creates a new contract.
"""

from __future__ import annotations

import io
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from typing import Any, BinaryIO, Optional, Union

import openpyxl

from backend.services.database import db
from frontend.modules.transport.contract import derive_company_code


def clean_bank_name(raw: str) -> str:
    if not raw:
        return ""
    return re.sub(r"^(tại|tại)\s+(ngân hàng|ngan hang)\s+", "", raw.strip(), flags=re.IGNORECASE).strip()


def sync_contracts_from_workbook(
    file_or_path: Union[str, BinaryIO, bytes],
    dry_run: bool = False,
    seller_company_id: Optional[int] = None,
    request: Optional[Any] = None,
    user: Optional[dict] = None,
) -> dict[str, Any]:
    """Inspect and synchronize contracts and customers from Excel."""
    if isinstance(file_or_path, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(file_or_path), data_only=True)
    elif isinstance(file_or_path, str):
        wb = openpyxl.load_workbook(file_or_path, data_only=True)
    else:
        wb = openpyxl.load_workbook(file_or_path, data_only=True)

    # Xác định công ty đang đăng nhập / bên bán xuất hợp đồng (Bên B)
    from frontend.context.company_context import get_current_company_id
    active_cid = seller_company_id or get_current_company_id() or 33

    seller = None
    if active_cid:
        seller = db._execute_query(
            "SELECT * FROM invoice_companies WHERE id = %s",
            (int(active_cid),),
            fetch=True,
            commit=False,
        )

    if not seller and user:
        try:
            from frontend.deps import get_user_assigned_companies
            assigned = get_user_assigned_companies(user)
            if assigned:
                seller = db._execute_query(
                    "SELECT * FROM invoice_companies WHERE id = %s",
                    (int(assigned[0]["id"]),),
                    fetch=True,
                    commit=False,
                )
        except Exception:
            pass

    if not seller:
        raise ValueError("Vui lòng chọn hoặc thiết lập công ty đang làm việc trước khi đồng bộ hợp đồng.")

    seller_id = int(seller["id"])
    seller_code = derive_company_code(seller["name"]) or "SELLER"

    # --- 1. Sync Customer Companies from DSKH ---
    ws_dskh = wb["DSKH"] if "DSKH" in wb.sheetnames else None
    companies_db = (
        db._execute_query(
            "SELECT id, tax_code, name, address, director_name, bank_account, bank_name FROM invoice_companies",
            fetchall=True,
            commit=False,
        )
        or []
    )
    comp_map = {c["tax_code"]: c for c in companies_db if c.get("tax_code")}

    updated_comps = 0
    created_comps = 0
    updated_bank_accs = 0

    if ws_dskh:
        for r in list(ws_dskh.iter_rows(values_only=True))[6:]:
            mst = str(r[2]).strip() if len(r) > 2 and r[2] else ""
            if not mst or mst.lower() == "none":
                continue
            name = str(r[3]).strip() if len(r) > 3 and r[3] else ""
            addr = str(r[4]).strip() if len(r) > 4 and r[4] else ""
            rep = str(r[5]).strip() if len(r) > 5 and r[5] else ""
            stk1 = str(r[7]).strip() if len(r) > 7 and r[7] else ""
            nh1 = clean_bank_name(str(r[8])) if len(r) > 8 and r[8] else ""
            stk2 = str(r[9]).strip() if len(r) > 9 and r[9] else ""
            nh2 = clean_bank_name(str(r[10])) if len(r) > 10 and r[10] else ""

            comp = comp_map.get(mst)
            if comp:
                cid = comp["id"]
                needs_update = False
                up_fields = {}
                if rep and rep != (comp.get("director_name") or ""):
                    needs_update = True
                    up_fields["director_name"] = rep
                if addr and addr != (comp.get("address") or ""):
                    needs_update = True
                    up_fields["address"] = addr
                if stk1 and stk1 != (comp.get("bank_account") or ""):
                    needs_update = True
                    up_fields["bank_account"] = stk1
                if nh1 and nh1 != (comp.get("bank_name") or ""):
                    needs_update = True
                    up_fields["bank_name"] = nh1

                if needs_update:
                    updated_comps += 1
                    if not dry_run:
                        set_clauses = [f"{k} = %s" for k in up_fields.keys()]
                        set_clauses.append("updated_at = NOW()")
                        params = tuple(up_fields.values()) + (cid,)
                        db._execute_query(
                            f"UPDATE invoice_companies SET {', '.join(set_clauses)} WHERE id = %s",
                            params,
                            commit=True,
                        )
            else:
                created_comps += 1
                if not dry_run:
                    cid = db._execute_query(
                        """INSERT INTO invoice_companies (name, tax_code, address, director_name, bank_account, bank_name, created_at, updated_at)
                           VALUES (%s, %s, %s, %s, %s, %s, NOW(), NOW()) RETURNING id""",
                        (name, mst, addr, rep or None, stk1 or None, nh1 or None),
                        fetch=True,
                        commit=True,
                        return_id=True,
                    )
                    comp = {"id": cid, "tax_code": mst, "name": name}
                    comp_map[mst] = comp

            # Tài khoản phụ (stk2)
            if stk2 and nh2 and comp:
                cid = comp["id"]
                existing_ba = db._execute_query(
                    "SELECT id FROM company_bank_accounts WHERE company_id = %s AND account_number = %s",
                    (cid, stk2),
                    fetch=True,
                    commit=False,
                )
                if not existing_ba:
                    updated_bank_accs += 1
                    if not dry_run:
                        db._execute_query(
                            """INSERT INTO company_bank_accounts (company_id, account_number, bank_name, branch, sort_order)
                               VALUES (%s, %s, %s, '', 1)""",
                            (cid, stk2, nh2),
                            commit=True,
                        )

        # Refresh comp_map sau khi cập nhật danh sách
        companies_db = (
            db._execute_query(
                "SELECT id, tax_code, name, address, director_name, bank_account, bank_name FROM invoice_companies",
                fetchall=True,
                commit=False,
            )
            or []
        )
        comp_map = {c["tax_code"]: c for c in companies_db if c.get("tax_code")}

    # --- 2. Sync Contracts from DATA ---
    if "DATA" not in wb.sheetnames:
        raise ValueError("File Excel không có sheet DATA chứa dữ liệu hợp đồng.")

    ws_data = wb["DATA"]
    contracts_raw = defaultdict(list)
    for r in list(ws_data.iter_rows(values_only=True))[7:]:
        so_hd = r[1]
        if so_hd is None or str(so_hd).strip() == "" or str(so_hd).strip().lower() == "none":
            continue
        contracts_raw[str(so_hd).strip()].append(r)

    # Đếm số hợp đồng cùng ngày + MST để phân biệt số hợp đồng trùng
    date_comp_counts = defaultdict(int)
    for so_hd, rows in contracts_raw.items():
        r = rows[0]
        dt = r[2]
        mst = str(r[3]).strip() if len(r) > 3 and r[3] else ""
        if mst and mst.lower() != "none":
            key = (dt.strftime("%Y-%m-%d") if hasattr(dt, "strftime") else str(dt)[:10], mst)
            date_comp_counts[key] += 1

    existing_contracts = (
        db._execute_query(
            "SELECT * FROM transport_contracts WHERE seller_company_id = %s OR seller_company_id IS NULL",
            (seller_id,),
            fetchall=True,
            commit=False,
        )
        or []
    )
    existing_by_notes = {}
    existing_by_cno = {}
    for ec in existing_contracts:
        notes = ec.get("notes") or ""
        m = re.search(r"HĐ\s*#(\d+)", notes)
        if m:
            existing_by_notes[m.group(1)] = ec
        if ec.get("contract_no"):
            existing_by_cno[ec["contract_no"]] = ec

    created_contracts = 0
    updated_contracts = 0
    skipped_contracts = 0
    unchanged_contracts = 0

    for so_hd, rows in sorted(contracts_raw.items(), key=lambda x: int(x[0]) if x[0].isdigit() else x[0]):
        r = rows[0]
        raw_dt = r[2]
        mst = str(r[3]).strip() if len(r) > 3 and r[3] else ""
        c_name = r[4] or ""
        total_val = sum(Decimal(str(x[10] or 0)) for x in rows if len(x) > 10 and x[10] is not None)

        if isinstance(raw_dt, datetime):
            c_date = raw_dt.date()
        elif isinstance(raw_dt, date):
            c_date = raw_dt
        elif isinstance(raw_dt, str) and raw_dt:
            try:
                c_date = date.fromisoformat(raw_dt[:10])
            except Exception:
                c_date = date.today()
        else:
            c_date = date.today()

        day_str = f"{c_date.day:02d}"
        month_str = f"{c_date.month:02d}"
        short_year = str(c_date.year)[-2:]

        if not mst or mst.lower() == "none":
            skipped_contracts += 1
            continue

        comp = comp_map.get(mst)
        if not comp:
            skipped_contracts += 1
            continue

        cid = comp["id"]
        cust_code = derive_company_code(comp["name"])
        date_key = (c_date.isoformat(), mst)
        is_dup_date = date_comp_counts[date_key] > 1

        if is_dup_date:
            contract_no = f"{day_str}{month_str}{short_year}/HĐNT/{cust_code}-{seller_code}-{so_hd}"
        else:
            contract_no = f"{day_str}{month_str}{short_year}/HĐNT/{cust_code}-{seller_code}"

        title = f"Hợp đồng vận chuyển số {so_hd} - {comp['name']}"
        category = "Hợp đồng vận chuyển"
        eff_date = c_date
        exp_date = date(c_date.year, 12, 31)
        status = "active"

        # Tóm tắt mặt hàng
        item_lines = []
        for x in rows:
            desc = str(x[6]).strip() if len(x) > 6 and x[6] else ""
            dvt = str(x[7]).strip() if len(x) > 7 and x[7] else ""
            sl = str(x[8]).strip() if len(x) > 8 and x[8] else ""
            dg_val = x[9] if len(x) > 9 and x[9] is not None else 0
            tt_val = x[10] if len(x) > 10 and x[10] is not None else 0
            dg = f"{int(dg_val):,}".replace(",", ".")
            tt = f"{int(tt_val):,}".replace(",", ".")
            if desc:
                item_lines.append(f"{desc} (SL: {sl} {dvt} * {dg} đ = {tt} đ)")
            else:
                item_lines.append(f"(SL: {sl} {dvt} * {dg} đ = {tt} đ)")

        items_summary = "; ".join(item_lines)
        notes = f"HĐ #{so_hd} | {items_summary}"

        matched_ec = existing_by_notes.get(so_hd) or existing_by_cno.get(contract_no)

        if matched_ec:
            cid_old = matched_ec.get("company_id")
            val_old = Decimal(str(matched_ec.get("contract_value") or 0))
            s_date_old = matched_ec.get("sign_date")
            title_old = matched_ec.get("title") or ""
            notes_old = matched_ec.get("notes") or ""

            diff = {}
            if cid_old != cid:
                diff["company_id"] = (cid_old, cid)
            if val_old != total_val:
                diff["contract_value"] = (val_old, total_val)
            if s_date_old != c_date:
                diff["sign_date"] = (s_date_old, c_date)
            if title_old != title:
                diff["title"] = (title_old, title)
            if notes_old != notes:
                diff["notes"] = (notes_old, notes)

            if diff:
                updated_contracts += 1
                if not dry_run:
                    db._execute_query(
                        """UPDATE transport_contracts
                           SET contract_no = %s, title = %s, category = %s, company_id = %s,
                               seller_company_id = %s, sign_date = %s, effective_date = %s,
                               expiration_date = %s, contract_value = %s, status = %s,
                               notes = %s, updated_at = NOW()
                           WHERE id = %s""",
                        (
                            contract_no,
                            title,
                            category,
                            cid,
                            seller_id,
                            c_date,
                            eff_date,
                            exp_date,
                            total_val,
                            status,
                            notes,
                            matched_ec["id"],
                        ),
                        commit=True,
                    )
            else:
                unchanged_contracts += 1
        else:
            created_contracts += 1
            if not dry_run:
                db._execute_query(
                    """INSERT INTO transport_contracts
                       (contract_no, title, category, company_id, seller_company_id, sign_date,
                        effective_date, expiration_date, contract_value, status, notes, created_at, updated_at)
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, NOW(), NOW())""",
                    (
                        contract_no,
                        title,
                        category,
                        cid,
                        seller_id,
                        c_date,
                        eff_date,
                        exp_date,
                        total_val,
                        status,
                        notes,
                    ),
                    commit=True,
                )

    return {
        "updated_companies": updated_comps,
        "created_companies": created_comps,
        "added_bank_accounts": updated_bank_accs,
        "created_contracts": created_contracts,
        "updated_contracts": updated_contracts,
        "unchanged_contracts": unchanged_contracts,
        "skipped_contracts": skipped_contracts,
        "total_file_contracts": len(contracts_raw),
    }
