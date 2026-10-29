"""Service for Cash Book (Mau S04a-DNN) and Bank Book (Mau S05a-DNN) under Circular 133/2016/TT-BTC."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import logging
import re
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.services.database import db
from frontend.context.company_context import get_current_company_id

logger = logging.getLogger(__name__)

DAI_LO_XANH_COMPANY_ID = 33

# Số dư ban đầu mặc định
CASH_DEFAULT_OPENING = Decimal("0")               # TK 111: 0 đ
BANK_DEFAULT_OPENING = Decimal("0")               # TK 112: 0 đ


def parse_flexible_date(raw_date: Any) -> date:
    """Parse various date formats (ISO, DD/MM/YYYY, localized VN strings) reliably."""
    if not raw_date:
        return date.today()
    if isinstance(raw_date, date) and not isinstance(raw_date, datetime):
        return raw_date
    if isinstance(raw_date, datetime):
        return raw_date.date()
    s = str(raw_date).strip()
    if not s:
        return date.today()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    # Localized VN strings like "ngày 6 thg 5, 2026" or "06/05/2026"
    m = re.search(r"(\d{1,2})[^\d]+(\d{1,2})[^\d]+(\d{4})", s)
    if m:
        try:
            d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            return date(y, mo, d)
        except Exception:
            pass
    m2 = re.search(r"(\d{4})[^\d]+(\d{1,2})[^\d]+(\d{1,2})", s)
    if m2:
        try:
            y, mo, d = int(m2.group(1)), int(m2.group(2)), int(m2.group(3))
            return date(y, mo, d)
        except Exception:
            pass
    return date.today()


from functools import lru_cache


@lru_cache(maxsize=64)
def _fetch_company_info(company_id: int) -> dict[str, str]:
    comp = db._execute_query(
        "SELECT name, address, tax_code, bank_account, bank_name FROM public.invoice_companies WHERE id = %s",
        (company_id,), fetch=True, commit=False
    ) or {}
    return {
        "name": comp.get("name") or "",
        "address": comp.get("address") or "",
        "tax_code": comp.get("tax_code") or "",
        "bank_account": comp.get("bank_account") or "",
        "bank_name": comp.get("bank_name") or "",
    }


def get_company_info(company_id: int | None = None) -> dict[str, str]:
    """Retrieve company name, address, tax code and bank info from DB (cached)."""
    company_id = company_id or get_current_company_id()
    if not company_id:
        return {"name": "", "address": "", "tax_code": "", "bank_account": "", "bank_name": ""}
    return _fetch_company_info(int(company_id))


def get_opening_balance_setting(book_type: str, company_id: int | None = None) -> Decimal:
    """Retrieve opening balance from app_settings or fallback to default."""
    company_id = company_id or get_current_company_id()
    if not company_id:
        return Decimal(0)
    b_type = "bank" if "bank" in book_type.lower() or "ngân hàng" in book_type.lower() else "cash"
    key = f"transport_{b_type}_opening_balance_{company_id}"
    row = db._execute_query(
        "SELECT value FROM public.app_settings WHERE key = %s",
        (key,), fetch=True, commit=False
    )
    if row and row.get("value") is not None and str(row.get("value")).strip() != "":
        try:
            return Decimal(str(row["value"]).replace(".", "").replace(",", "").strip())
        except Exception:
            pass
    return BANK_DEFAULT_OPENING if b_type == "bank" else CASH_DEFAULT_OPENING


def set_opening_balance_setting(book_type: str, amount: Any, company_id: int | None = None) -> tuple[bool, str]:
    """Save user-configured opening balance into app_settings."""
    company_id = company_id or get_current_company_id()
    if not company_id or int(company_id) <= 0:
        return False, "Chưa chọn công ty làm việc."
    b_type = "bank" if "bank" in book_type.lower() or "ngân hàng" in book_type.lower() else "cash"
    key = f"transport_{b_type}_opening_balance_{company_id}"
    cleaned = str(amount).replace(".", "").replace(",", "").replace(" ", "").strip()
    try:
        val = Decimal(cleaned or "0")
    except Exception:
        return False, "Số tiền số dư đầu kỳ không hợp lệ."
    db._execute_query('''
        INSERT INTO public.app_settings (key, value, updated_at)
        VALUES (%s, %s, NOW())
        ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value, updated_at = NOW()
    ''', (key, str(val)), commit=True)
    b_name = "Sổ tiền gửi ngân hàng (TK 112)" if b_type == "bank" else "Sổ quỹ tiền mặt (TK 111)"
    return True, f"Đã cập nhật số dư đầu kỳ cho {b_name} thành {int(val):,} đ."


def sync_cash_bank_transfers(company_id: int | None = None) -> int:
    """
    Tự động đồng bộ liên kết 2 chiều giữa Sổ quỹ tiền mặt (TK 111) và Sổ tiền gửi NH (TK 112):
    1. Khi Phiếu chi tiền mặt có TK đối 112 (112, 1121...) -> Tự động sinh/cập nhật Báo Có trên Sổ ngân hàng (TK đối 1111).
    2. Khi Phiếu thu tiền mặt có TK đối 112 (112, 1121...) -> Tự động sinh/cập nhật Báo Nợ trên Sổ ngân hàng (TK đối 1111).
    3. Ngược lại, khi Giao dịch ngân hàng có TK đối 111 (111, 1111...) -> Tự động sinh/cập nhật Phiếu thu/chi tiền mặt.
    4. Tự động dọn dẹp các chứng từ liên kết mồ côi nếu phiếu gốc bị xóa hoặc đổi sang TK đối khác.
    """
    company_id = company_id or get_current_company_id()
    if not company_id or int(company_id) <= 0:
        return 0
    comp_info = get_company_info(company_id)
    bank_acc = comp_info.get("bank_account") or ""
    bank_name = comp_info.get("bank_name") or ""
    synced_count = 0

    # 1. Đồng bộ Cash -> Bank
    cash_transfers = db._execute_query('''
        SELECT id, company_id, book_type, entry_date, voucher_date, voucher_no, voucher_type,
               description, offset_account, amount, partner_name, partner_code, note
        FROM public.transport_cash_book
        WHERE company_id = %s 
          AND book_type = 'cash' 
          AND (offset_account LIKE '112%%' OR offset_account = '112')
          AND (source_type IS NULL OR source_type != 'bank_transfer')
        ORDER BY id ASC
    ''', (company_id,), fetchall=True, commit=False) or []

    valid_cash_ids = set()
    for c in cash_transfers:
        cid = c["id"]
        valid_cash_ids.add(str(cid))
        is_payment = (c["voucher_type"] == "payment")
        target_v_type = "receipt" if is_payment else "payment"
        target_offset = "1111" if str(c.get("offset_account") or "").endswith("1") else "111"

        c_vno = (c.get("voucher_no") or "").strip()
        target_vno = f"BC-{c_vno}" if target_v_type == "receipt" else f"UNC-{c_vno}"
        target_desc = c.get("description") or ("Nộp tiền vào TK ngân hàng" if target_v_type == "receipt" else "Rút tiền ngân hàng về quỹ")
        target_amount = c["amount"]
        target_date = c["voucher_date"] or c["entry_date"]
        target_note = c.get("note") or f"Liên kết tự động từ Phiếu {c_vno} (Sổ quỹ tiền mặt)"

        existing_bank = db._execute_query('''
            SELECT id, voucher_no, voucher_type, amount, voucher_date, description, offset_account
            FROM public.transport_cash_book
            WHERE company_id = %s AND book_type = 'bank' AND source_type = 'cash_transfer' AND source_id = %s
        ''', (company_id, str(cid)), fetch=True, commit=False)

        if not existing_bank:
            db._execute_query('''
                INSERT INTO public.transport_cash_book
                (company_id, book_type, entry_date, voucher_date, voucher_no, voucher_type,
                 description, offset_account, amount, bank_account_no, bank_name,
                 partner_name, partner_code, source_type, source_id, note)
                VALUES (%s, 'bank', %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'cash_transfer', %s, %s)
            ''', (
                company_id, target_date, target_date, target_vno, target_v_type,
                target_desc, target_offset, target_amount, bank_acc, bank_name,
                c.get("partner_name") or "", c.get("partner_code") or "",
                str(cid), target_note
            ), commit=True)
            synced_count += 1
        else:
            db._execute_query('''
                UPDATE public.transport_cash_book
                SET entry_date = %s, voucher_date = %s, voucher_type = %s,
                    description = %s, offset_account = %s, amount = %s,
                    bank_account_no = COALESCE(bank_account_no, %s),
                    bank_name = COALESCE(bank_name, %s),
                    partner_name = %s, partner_code = %s, note = %s
                WHERE id = %s
            ''', (
                target_date, target_date, target_v_type,
                target_desc, target_offset, target_amount,
                bank_acc, bank_name,
                c.get("partner_name") or "", c.get("partner_code") or "", target_note,
                existing_bank["id"]
            ), commit=True)

    # Dọn dẹp bank mồ côi
    if valid_cash_ids:
        db._execute_query('''
            DELETE FROM public.transport_cash_book
            WHERE company_id = %s AND book_type = 'bank' AND source_type = 'cash_transfer'
              AND (source_id IS NULL OR source_id NOT IN %s)
        ''', (company_id, tuple(valid_cash_ids)), commit=True)
    else:
        db._execute_query('''
            DELETE FROM public.transport_cash_book
            WHERE company_id = %s AND book_type = 'bank' AND source_type = 'cash_transfer'
        ''', (company_id,), commit=True)

    # 2. Đồng bộ Bank -> Cash
    bank_transfers = db._execute_query('''
        SELECT id, company_id, book_type, entry_date, voucher_date, voucher_no, voucher_type,
               description, offset_account, amount, partner_name, partner_code, note
        FROM public.transport_cash_book
        WHERE company_id = %s 
          AND book_type = 'bank' 
          AND (offset_account LIKE '111%%' OR offset_account = '111')
          AND (source_type IS NULL OR source_type != 'cash_transfer')
        ORDER BY id ASC
    ''', (company_id,), fetchall=True, commit=False) or []

    valid_bank_ids = set()
    for b in bank_transfers:
        bid = b["id"]
        valid_bank_ids.add(str(bid))
        is_receipt = (b["voucher_type"] == "receipt")
        target_v_type = "payment" if is_receipt else "receipt"
        target_offset = "1121" if str(b.get("offset_account") or "").endswith("1") else "112"

        b_vno = (b.get("voucher_no") or "").strip()
        target_vno = f"PC-{b_vno}" if target_v_type == "payment" else f"PT-{b_vno}"
        target_desc = b.get("description") or ("Nộp tiền vào TK ngân hàng" if target_v_type == "payment" else "Rút tiền ngân hàng về quỹ")
        target_amount = b["amount"]
        target_date = b["voucher_date"] or b["entry_date"]
        target_note = b.get("note") or f"Liên kết tự động từ Giao dịch {b_vno} (Sổ tiền gửi NH)"

        existing_cash = db._execute_query('''
            SELECT id FROM public.transport_cash_book
            WHERE company_id = %s AND book_type = 'cash' AND source_type = 'bank_transfer' AND source_id = %s
        ''', (company_id, str(bid)), fetch=True, commit=False)

        if not existing_cash:
            db._execute_query('''
                INSERT INTO public.transport_cash_book
                (company_id, book_type, entry_date, voucher_date, voucher_no, voucher_type,
                 description, offset_account, amount, partner_name, partner_code, source_type, source_id, note)
                VALUES (%s, 'cash', %s, %s, %s, %s, %s, %s, %s, %s, %s, 'bank_transfer', %s, %s)
            ''', (
                company_id, target_date, target_date, target_vno, target_v_type,
                target_desc, target_offset, target_amount,
                b.get("partner_name") or "", b.get("partner_code") or "",
                str(bid), target_note
            ), commit=True)
            synced_count += 1
        else:
            db._execute_query('''
                UPDATE public.transport_cash_book
                SET entry_date = %s, voucher_date = %s, voucher_type = %s,
                    description = %s, offset_account = %s, amount = %s,
                    partner_name = %s, partner_code = %s, note = %s
                WHERE id = %s
            ''', (
                target_date, target_date, target_v_type,
                target_desc, target_offset, target_amount,
                b.get("partner_name") or "", b.get("partner_code") or "", target_note,
                existing_cash["id"]
            ), commit=True)

    if valid_bank_ids:
        db._execute_query('''
            DELETE FROM public.transport_cash_book
            WHERE company_id = %s AND book_type = 'cash' AND source_type = 'bank_transfer'
              AND (source_id IS NULL OR source_id NOT IN %s)
        ''', (company_id, tuple(valid_bank_ids)), commit=True)
    else:
        db._execute_query('''
            DELETE FROM public.transport_cash_book
            WHERE company_id = %s AND book_type = 'cash' AND source_type = 'bank_transfer'
        ''', (company_id,), commit=True)

    return synced_count


def get_cash_book_data(
    company_id: int | None = None,
    year: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> dict[str, Any]:
    """Retrieve Cash Book data (Mau S04a-DNN) with running balance."""
    company_id = company_id or get_current_company_id()
    today = date.today()
    if date_from is None:
        y = year or today.year
        date_from = date(y, 1, 1)
    if date_to is None:
        y = year or today.year
        date_to = date(y, 12, 31)

    if not company_id or int(company_id) <= 0:
        return {
            "company": {"name": "", "address": "", "tax_code": "", "bank_account": "", "bank_name": ""},
            "year": year or today.year,
            "date_from": date_from,
            "date_to": date_to,
            "opening_balance": 0,
            "initial_opening": 0,
            "total_receipt": 0,
            "total_payment": 0,
            "closing_balance": 0,
            "entries": [],
        }

    # Tự động đồng bộ các khoản chuyển tiền mặt <-> ngân hàng (111 <-> 112)
    sync_cash_bank_transfers(company_id=company_id)

    # 1. Tính số dư đầu kỳ (Trước date_from)
    # Lấy các phát sinh trước date_from
    prev_sums = db._execute_query('''
        SELECT 
            COALESCE(SUM(CASE WHEN voucher_type = 'receipt' THEN amount ELSE 0 END), 0) as total_receipt,
            COALESCE(SUM(CASE WHEN voucher_type = 'payment' THEN amount ELSE 0 END), 0) as total_payment
        FROM public.transport_cash_book
        WHERE company_id = %s AND book_type = 'cash' AND entry_date < %s
    ''', (company_id, date_from), fetch=True, commit=False) or {}

    prev_receipt = Decimal(str(prev_sums.get("total_receipt") or 0))
    prev_payment = Decimal(str(prev_sums.get("total_payment") or 0))
    initial_opening = get_opening_balance_setting("cash")
    opening_balance = initial_opening + prev_receipt - prev_payment

    # 2. Lấy các phát sinh trong kỳ [date_from, date_to]
    rows_db = db._execute_query('''
        SELECT id, entry_date, voucher_date, voucher_no, voucher_type,
               description, offset_account, amount, partner_name, note,
               COALESCE(partner_code, '') as partner_code
        FROM public.transport_cash_book
        WHERE company_id = %s AND book_type = 'cash' AND entry_date >= %s AND entry_date <= %s
        ORDER BY entry_date ASC, id ASC
    ''', (company_id, date_from, date_to), fetchall=True, commit=False) or []

    running_bal = opening_balance
    entries = []
    total_receipt = Decimal(0)
    total_payment = Decimal(0)

    for r in rows_db:
        amt = Decimal(str(r["amount"] or 0))
        is_rec = (r["voucher_type"] == "receipt")
        receipt_amt = amt if is_rec else Decimal(0)
        payment_amt = Decimal(0) if is_rec else amt

        if is_rec:
            running_bal += amt
            total_receipt += amt
        else:
            running_bal -= amt
            total_payment += amt

        entries.append({
            "id": r["id"],
            "entry_date": r["entry_date"].strftime("%d/%m/%Y") if r["entry_date"] else "",
            "voucher_date": r["voucher_date"].strftime("%d/%m/%Y") if r["voucher_date"] else "",
            "receipt_no": r["voucher_no"] if is_rec else "",
            "payment_no": "" if is_rec else r["voucher_no"],
            "partner_code": r.get("partner_code") or "",
            "description": r["description"] or "",
            "offset_account": r["offset_account"] or ("131" if is_rec else "331"),
            "receipt_amount": int(round(receipt_amt)),
            "payment_amount": int(round(payment_amt)),
            "balance": int(round(running_bal)),
            "partner_name": r["partner_name"] or "",
            "note": r["note"] or "",
        })

    closing_balance = running_bal
    comp_info = get_company_info()

    return {
        "title": "SỔ QUỸ TIỀN MẶT",
        "form_code": "Mẫu số S04a-DNN",
        "circular": "Ban hành theo Thông tư số 133/2016/TT-BTC ngày 26/8/2016 của Bộ Tài chính",
        "currency": "VND",
        "company_name": comp_info["name"],
        "company_address": comp_info["address"],
        "tax_code": comp_info["tax_code"],
        "date_from": date_from.strftime("%Y-%m-%d"),
        "date_to": date_to.strftime("%Y-%m-%d"),
        "date_str": f"Từ ngày {date_from.strftime('%d/%m/%Y')} đến ngày {date_to.strftime('%d/%m/%Y')}",
        "opening_balance": int(round(opening_balance)),
        "initial_opening": int(round(initial_opening)),
        "total_receipt": int(round(total_receipt)),
        "total_payment": int(round(total_payment)),
        "closing_balance": int(round(closing_balance)),
        "entries": entries,
    }


def get_bank_book_data(
    company_id: int | None = None,
    year: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    bank_account_no: str = "",
    bank_name: str = "",
) -> dict[str, Any]:
    """Retrieve Bank Deposit Book data (Mau S05a-DNN) with running balance."""
    company_id = company_id or get_current_company_id()
    today = date.today()
    if date_from is None:
        y = year or today.year
        date_from = date(y, 1, 1)
    if date_to is None:
        y = year or today.year
        date_to = date(y, 12, 31)

    if not company_id or int(company_id) <= 0:
        return {
            "company": {"name": "", "address": "", "tax_code": "", "bank_account": "", "bank_name": ""},
            "bank_account": "",
            "bank_name": "",
            "year": year or today.year,
            "date_from": date_from,
            "date_to": date_to,
            "opening_balance": 0,
            "initial_opening": 0,
            "total_deposit": 0,
            "total_withdraw": 0,
            "closing_balance": 0,
            "entries": [],
            "bank_accounts": [],
        }

    comp_info = get_company_info(company_id)
    actual_bank_account = bank_account_no or comp_info["bank_account"]
    actual_bank_name = bank_name or comp_info["bank_name"]

    # Tự động đồng bộ các khoản chuyển tiền mặt <-> ngân hàng (111 <-> 112)
    sync_cash_bank_transfers(company_id=company_id)

    prev_sums = db._execute_query('''
        SELECT 
            COALESCE(SUM(CASE WHEN voucher_type = 'receipt' THEN amount ELSE 0 END), 0) as total_receipt,
            COALESCE(SUM(CASE WHEN voucher_type = 'payment' THEN amount ELSE 0 END), 0) as total_payment
        FROM public.transport_cash_book
        WHERE company_id = %s AND book_type = 'bank' AND entry_date < %s
    ''', (company_id, date_from), fetch=True, commit=False) or {}

    prev_receipt = Decimal(str(prev_sums.get("total_receipt") or 0))
    prev_payment = Decimal(str(prev_sums.get("total_payment") or 0))
    initial_opening = get_opening_balance_setting("bank")
    opening_balance = initial_opening + prev_receipt - prev_payment

    rows_db = db._execute_query('''
        SELECT id, entry_date, voucher_date, voucher_no, voucher_type,
               description, offset_account, amount, bank_account_no, bank_name, partner_name, note,
               COALESCE(partner_code, '') as partner_code
        FROM public.transport_cash_book
        WHERE company_id = %s AND book_type = 'bank' AND entry_date >= %s AND entry_date <= %s
        ORDER BY entry_date ASC, id ASC
    ''', (company_id, date_from, date_to), fetchall=True, commit=False) or []

    running_bal = opening_balance
    entries = []
    total_deposit = Decimal(0)
    total_withdraw = Decimal(0)

    for r in rows_db:
        amt = Decimal(str(r["amount"] or 0))
        is_rec = (r["voucher_type"] == "receipt")
        deposit_amt = amt if is_rec else Decimal(0)
        withdraw_amt = Decimal(0) if is_rec else amt

        if is_rec:
            running_bal += amt
            total_deposit += amt
        else:
            running_bal -= amt
            total_withdraw += amt

        entries.append({
            "id": r["id"],
            "entry_date": r["entry_date"].strftime("%d/%m/%Y") if r["entry_date"] else "",
            "voucher_date": r["voucher_date"].strftime("%d/%m/%Y") if r["voucher_date"] else "",
            "voucher_no": r["voucher_no"] or "",
            "voucher_type": r["voucher_type"],
            "partner_code": r.get("partner_code") or "",
            "description": r["description"] or "",
            "offset_account": r["offset_account"] or ("131" if is_rec else "331"),
            "deposit_amount": int(round(deposit_amt)),
            "withdraw_amount": int(round(withdraw_amt)),
            "balance": int(round(running_bal)),
            "bank_account_no": r["bank_account_no"] or actual_bank_account,
            "bank_name": r["bank_name"] or actual_bank_name,
            "partner_name": r["partner_name"] or "",
            "note": r["note"] or "",
        })

    closing_balance = running_bal

    return {
        "title": "SỔ TIỀN GỬI NGÂN HÀNG",
        "form_code": "Mẫu số S05a-DNN",
        "circular": "Ban hành theo Thông tư số 133/2016/TT-BTC ngày 26/8/2016 của Bộ Tài chính",
        "currency": "VND",
        "bank_account_no": actual_bank_account,
        "bank_name": actual_bank_name,
        "company_name": comp_info["name"],
        "company_address": comp_info["address"],
        "tax_code": comp_info["tax_code"],
        "date_from": date_from.strftime("%Y-%m-%d"),
        "date_to": date_to.strftime("%Y-%m-%d"),
        "date_str": f"Từ ngày {date_from.strftime('%d/%m/%Y')} đến ngày {date_to.strftime('%d/%m/%Y')}",
        "opening_balance": int(round(opening_balance)),
        "initial_opening": int(round(initial_opening)),
        "total_deposit": int(round(total_deposit)),
        "total_withdraw": int(round(total_withdraw)),
        "closing_balance": int(round(closing_balance)),
        "entries": entries,
    }


def record_cash_or_bank_entry(
    company_id: int | None = None,
    book_type: str = "cash",   # 'cash' hoặc 'bank'
    voucher_type: str = "receipt", # 'receipt' (Thu / Báo Có) hoặc 'payment' (Chi / Báo Nợ)
    amount: Decimal | float | int = 0,
    voucher_date: date | str | None = None,
    voucher_no: str = "",
    description: str = "",
    offset_account: str = "131",
    partner_name: str = "",
    partner_code: str = "",
    bank_account_no: str = "",
    bank_name: str = "",
    source_type: str = "manual",
    source_id: str | int | None = None,
    note: str = "",
) -> tuple[bool, str, int]:
    """Record a cash/bank transaction when debt is collected or paid."""
    company_id = company_id or get_current_company_id()
    if not company_id:
        return False, "Chưa chọn công ty làm việc.", 0
    try:
        cleaned_str = str(amount).replace(".", "").replace(",", "").replace(" ", "").strip()
        val = Decimal(cleaned_str)
    except Exception:
        return False, "Số tiền không hợp lệ.", 0

    if val <= 0:
        return False, "Số tiền phải lớn hơn 0.", 0

    v_date = parse_flexible_date(voucher_date)

    b_type = "bank" if "ngân hàng" in book_type.lower() or "bank" in book_type.lower() or "chuyển khoản" in book_type.lower() or "ck" in book_type.lower() else "cash"
    if b_type == "bank" and (not bank_account_no or not bank_name):
        comp_info = get_company_info(company_id)
        bank_account_no = bank_account_no or comp_info["bank_account"]
        bank_name = bank_name or comp_info["bank_name"]

    if not voucher_no:
        # Tự sinh số chứng từ
        prefix = "PT" if (b_type == "cash" and voucher_type == "receipt") else (
            "PC" if (b_type == "cash" and voucher_type == "payment") else (
                "BC-" if voucher_type == "receipt" else "UNC-"
            )
        )
        count_row = db._execute_query('''
            SELECT COUNT(*) as cnt FROM public.transport_cash_book
            WHERE company_id = %s AND book_type = %s AND voucher_type = %s
        ''', (company_id, b_type, voucher_type), fetch=True, commit=False) or {}
        num = (count_row.get("cnt") or 0) + 1
        voucher_no = f"{prefix}{num:04d}" if "-" in prefix else f"{prefix}{num:03d}"

    row_id = db._execute_query('''
        INSERT INTO public.transport_cash_book
        (company_id, book_type, entry_date, voucher_date, voucher_no, voucher_type,
         description, offset_account, amount, bank_account_no, bank_name, partner_name,
         partner_code, source_type, source_id, note)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        RETURNING id
    ''', (
        company_id, b_type, v_date, v_date, voucher_no, voucher_type,
        description.strip() or ("Thu tiền" if voucher_type == "receipt" else "Chi tiền"),
        offset_account, val,
        bank_account_no if b_type == "bank" else None,
        bank_name if b_type == "bank" else None,
        partner_name.strip(),
        partner_code.strip() if partner_code else None,
        source_type, str(source_id) if source_id else None,
        note.strip()
    ), fetch=True, commit=True)

    ins_id = row_id["id"] if row_id else 0
    # Đồng bộ liên kết nếu đây là giao dịch chuyển tiền giữa 111 và 112
    sync_cash_bank_transfers(company_id=company_id)

    # TỰ ĐỘNG LIÊN THÔNG VỚI BẢNG KHẤU HAO TSCĐ (TK 211) & BẢNG PHÂN BỔ CCDC (TK 242)
    clean_offset = str(offset_account or "").strip()
    desc_lower = (description or "").lower()
    book_name = "Sổ quỹ tiền mặt (TK 111)" if b_type == "cash" else "Sổ tiền gửi ngân hàng (TK 112)"

    if voucher_type == "payment":
        # 1. Liên thông ghi nhận Tài sản cố định (Nợ 211)
        if clean_offset.startswith("211") or ("mua xe" in desc_lower or "mua ô tô" in desc_lower):
            try:
                exist_asset = db._execute_query('''
                    SELECT id FROM public.enterprise_assets
                    WHERE company_id = %s AND (serial_number = %s OR note LIKE %s)
                ''', (company_id, f"CT-{ins_id}", f"%{voucher_no}%"), fetch=True, commit=False)
                if not exist_asset:
                    asset_cat = "Phương tiện vận tải" if ("xe" in desc_lower or "ô tô" in desc_lower or "cẩu" in desc_lower) else "Máy móc thiết bị"
                    depr_yrs = 6 if ("xe" in desc_lower or "ô tô" in desc_lower) else 5
                    db._execute_query('''
                        INSERT INTO public.enterprise_assets
                        (company_id, asset_code, name, category, serial_number, acquired_date,
                         purchase_cost, assigned_to, status, note, depreciation_years)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 'in_use', %s, %s)
                    ''', (
                        company_id, f"TSCD-{ins_id}", description.strip(), asset_cat,
                        f"CT-{ins_id}", v_date, val, partner_name.strip() or "Bộ phận vận tải",
                        f"Tự động ghi nhận từ chứng từ {voucher_no} ({book_name})", depr_yrs
                    ), commit=True)
                    logger.info(f"Auto-created enterprise_asset for company_id={company_id}: {description} ({val})")
            except Exception as e:
                logger.warning(f"Auto-sync asset failed: {e}")

        # 2. Liên thông ghi nhận Công cụ dụng cụ / Chi phí trả trước (Nợ 242)
        elif clean_offset.startswith("242") or ("quả cân" in desc_lower or "mua lốp" in desc_lower):
            try:
                from frontend.modules.transport.prepaid_expense_service import add_prepaid_expense
                add_prepaid_expense(
                    company_id=company_id,
                    item_code=f"CCDC-{ins_id}",
                    name=description.strip(),
                    unit="Cái" if ("quả cân" in desc_lower or "lốp" in desc_lower) else "Bộ",
                    quantity=1,
                    original_cost=val,
                    start_date=v_date,
                    total_months=24 if "quả cân" in desc_lower else 12,
                    note=f"Tự động ghi nhận từ chứng từ {voucher_no} ({book_name})",
                )
                logger.info(f"Auto-created prepaid_expense for company_id={company_id}: {description} ({val})")
            except Exception as e:
                logger.warning(f"Auto-sync prepaid expense failed: {e}")

    return True, f"Đã ghi sổ vào {book_name} thành công (Số CT: {voucher_no}).", ins_id


def get_cash_or_bank_entry(entry_id: int, company_id: int | None = None) -> dict[str, Any] | None:
    """Fetch single entry by ID and company_id for viewing/editing."""
    company_id = company_id or get_current_company_id()
    if not company_id or int(company_id) <= 0:
        return None
    row = db._execute_query('''
        SELECT id, company_id, book_type, entry_date, voucher_date, voucher_no, voucher_type,
               description, offset_account, amount, bank_account_no, bank_name, partner_name,
               COALESCE(partner_code, '') as partner_code,
               source_type, source_id, note
        FROM public.transport_cash_book
        WHERE id = %s AND company_id = %s
    ''', (entry_id, company_id), fetch=True, commit=False)
    if not row:
        return None
    return {
        "id": row["id"],
        "company_id": row["company_id"],
        "book_type": row["book_type"],
        "voucher_type": row["voucher_type"],
        "voucher_no": row["voucher_no"] or "",
        "voucher_date": row["voucher_date"].strftime("%Y-%m-%d") if row["voucher_date"] else "",
        "entry_date": row["entry_date"].strftime("%Y-%m-%d") if row["entry_date"] else "",
        "amount": int(round(Decimal(str(row["amount"] or 0)))),
        "description": row["description"] or "",
        "offset_account": row["offset_account"] or "",
        "partner_name": row["partner_name"] or "",
        "partner_code": row.get("partner_code") or "",
        "bank_account_no": row["bank_account_no"] or "",
        "bank_name": row["bank_name"] or "",
        "note": row["note"] or "",
        "source_type": row["source_type"] or "",
    }


def update_cash_or_bank_entry(
    entry_id: int,
    company_id: int | None = None,
    voucher_type: str | None = None,
    amount: Decimal | float | int | str | None = None,
    voucher_date: date | str | None = None,
    voucher_no: str | None = None,
    description: str | None = None,
    offset_account: str | None = None,
    partner_name: str | None = None,
    partner_code: str | None = None,
    bank_account_no: str | None = None,
    bank_name: str | None = None,
    note: str | None = None,
) -> tuple[bool, str]:
    """Update an existing cash or bank voucher entry."""
    company_id = company_id or get_current_company_id()
    if not company_id or int(company_id) <= 0:
        return False, "Chưa chọn công ty làm việc."
    existing = db._execute_query('''
        SELECT id, book_type, voucher_type, amount, voucher_date, voucher_no, description, offset_account, partner_name, note,
               source_type, source_id
        FROM public.transport_cash_book
        WHERE id = %s AND company_id = %s
    ''', (entry_id, company_id), fetch=True, commit=False)
    if not existing:
        return False, "Không tìm thấy chứng từ cần cập nhật."

    updates = []
    params = []
    val = None
    v_date = None

    if voucher_type is not None and voucher_type in ("receipt", "payment"):
        updates.append("voucher_type = %s")
        params.append(voucher_type)

    if amount is not None and str(amount).strip() != "":
        try:
            cleaned = str(amount).replace(".", "").replace(",", "").replace(" ", "").strip()
            val = Decimal(cleaned)
            if val <= 0:
                return False, "Số tiền phải lớn hơn 0."
            updates.append("amount = %s")
            params.append(val)
        except Exception:
            return False, "Số tiền không hợp lệ."

    if voucher_date is not None and str(voucher_date).strip() != "":
        v_date = parse_flexible_date(voucher_date)
        updates.append("voucher_date = %s")
        params.append(v_date)
        updates.append("entry_date = %s")
        params.append(v_date)

    if voucher_no is not None:
        updates.append("voucher_no = %s")
        params.append(voucher_no.strip())

    if description is not None:
        updates.append("description = %s")
        params.append(description.strip())

    if offset_account is not None:
        updates.append("offset_account = %s")
        params.append(offset_account.strip())

    if partner_name is not None:
        updates.append("partner_name = %s")
        params.append(partner_name.strip())

    if partner_code is not None:
        updates.append("partner_code = %s")
        params.append(partner_code.strip())

    if bank_account_no is not None:
        updates.append("bank_account_no = %s")
        params.append(bank_account_no.strip())

    if bank_name is not None:
        updates.append("bank_name = %s")
        params.append(bank_name.strip())

    if note is not None:
        updates.append("note = %s")
        params.append(note.strip())

    if not updates:
        return True, "Không có thay đổi nào."

    # Nếu người dùng sửa trực tiếp trên bản ghi counterpart liên kết (cash_transfer hoặc bank_transfer)
    if existing.get("source_type") in ("cash_transfer", "bank_transfer") and existing.get("source_id"):
        parent_id = existing["source_id"]
        parent_updates = []
        parent_params = []
        if val is not None:
            parent_updates.append("amount = %s")
            parent_params.append(val)
        if v_date is not None:
            parent_updates.append("voucher_date = %s")
            parent_updates.append("entry_date = %s")
            parent_params.extend([v_date, v_date])
        if description is not None:
            parent_updates.append("description = %s")
            parent_params.append(description.strip())
        if parent_updates:
            parent_params.extend([parent_id, company_id])
            db._execute_query(f'''
                UPDATE public.transport_cash_book
                SET {", ".join(parent_updates)}
                WHERE id = %s AND company_id = %s
            ''', tuple(parent_params), commit=True)

    params.extend([entry_id, company_id])
    db._execute_query(f'''
        UPDATE public.transport_cash_book
        SET {", ".join(updates)}
        WHERE id = %s AND company_id = %s
    ''', tuple(params), commit=True)

    # Tự động đồng bộ liên kết 111 <-> 112
    sync_cash_bank_transfers(company_id=company_id)
    return True, "Cập nhật chứng từ thành công."


def delete_cash_or_bank_entry(entry_id: int, company_id: int | None = None) -> tuple[bool, str]:
    """Delete an existing cash or bank voucher entry."""
    company_id = company_id or get_current_company_id()
    if not company_id or int(company_id) <= 0:
        return False, "Chưa chọn công ty làm việc."
    existing = db._execute_query('''
        SELECT id, voucher_no, source_type, source_id FROM public.transport_cash_book
        WHERE id = %s AND company_id = %s
    ''', (entry_id, company_id), fetch=True, commit=False)
    if not existing:
        return False, "Không tìm thấy chứng từ cần xóa."

    v_no = existing.get("voucher_no") or f"ID {entry_id}"

    # 1. Xóa bản ghi counterpart nếu đây là nguồn transfer
    db._execute_query('''
        DELETE FROM public.transport_cash_book
        WHERE company_id = %s AND source_id = %s AND source_type IN ('cash_transfer', 'bank_transfer')
    ''', (company_id, str(entry_id)), commit=True)

    # 2. Xóa bản ghi parent nếu người dùng bấm xóa trên bản ghi counterpart
    if existing.get("source_type") in ("cash_transfer", "bank_transfer") and existing.get("source_id"):
        db._execute_query('''
            DELETE FROM public.transport_cash_book
            WHERE company_id = %s AND id = %s
        ''', (company_id, existing["source_id"]), commit=True)

    # 3. Xóa chính bản ghi này
    db._execute_query('''
        DELETE FROM public.transport_cash_book
        WHERE id = %s AND company_id = %s
    ''', (entry_id, company_id), commit=True)

    sync_cash_bank_transfers(company_id=company_id)
    return True, f"Đã xóa chứng từ {v_no} thành công."


def export_cash_book_excel(
    company_id: int | None = None,
    year: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> bytes:
    """Export Cash Book to Excel matching Mau S04a-DNN."""
    company_id = company_id or get_current_company_id()
    if not company_id or int(company_id) <= 0:
        return b""
    data = get_cash_book_data(company_id=company_id, year=year, date_from=date_from, date_to=date_to)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S04a-DNN_So_Quy_Tien_Mat"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Times New Roman", size=14, bold=True, color="1E3A8A")
    font_bold = Font(name="Times New Roman", size=10, bold=True)
    font_regular = Font(name="Times New Roman", size=10)
    font_italic = Font(name="Times New Roman", size=9, italic=True)

    border_thin = Side(border_style="thin", color="000000")
    border_double = Side(border_style="double", color="000000")
    box_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    fill_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_total = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    # Header info
    ws["A1"] = f"Đơn vị: {data['company_name']}"
    ws["A1"].font = font_bold
    ws["A2"] = f"Địa chỉ: {data['company_address']}"
    ws["A2"].font = font_regular
    ws["A3"] = f"Mã số thuế: {data['tax_code']}"
    ws["A3"].font = font_regular

    ws["I1"] = data["form_code"]
    ws["I1"].font = font_bold
    ws["I1"].alignment = Alignment(horizontal="right")
    ws["I2"] = data["circular"]
    ws["I2"].font = font_italic
    ws["I2"].alignment = Alignment(horizontal="right")

    ws.merge_cells("A5:K5")
    ws["A5"] = data["title"]
    ws["A5"].font = font_title
    ws["A5"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A6:K6")
    ws["A6"] = data["date_str"]
    ws["A6"].font = font_italic
    ws["A6"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A7:K7")
    ws["A7"] = f"Loại quỹ: {data['currency']}"
    ws["A7"].font = font_regular
    ws["A7"].alignment = Alignment(horizontal="center")

    # Headers table (Rows 9-11)
    ws.merge_cells("A9:A10")
    ws["A9"] = "Ngày, tháng\nghi sổ"
    ws.merge_cells("B9:B10")
    ws["B9"] = "Ngày, tháng\nchứng từ"
    ws.merge_cells("C9:D9")
    ws["C9"] = "Số hiệu chứng từ"
    ws["C10"] = "Thu"
    ws["D10"] = "Chi"
    ws.merge_cells("E9:E10")
    ws["E9"] = "Mã khách"
    ws.merge_cells("F9:F10")
    ws["F9"] = "Tên khách"
    ws.merge_cells("G9:G10")
    ws["G9"] = "Diễn giải"
    ws.merge_cells("H9:H10")
    ws["H9"] = "TK đối ứng"
    ws.merge_cells("I9:K9")
    ws["I9"] = "Số tiền"
    ws["I10"] = "Thu"
    ws["J10"] = "Chi"
    ws["K10"] = "Tồn"

    letters = ["A", "B", "C", "D", "E", "F", "G", "H", "1", "2", "3"]
    for c_idx, ltr in enumerate(letters, 1):
        cell = ws.cell(row=11, column=c_idx, value=ltr)
        cell.font = font_bold
        cell.alignment = Alignment(horizontal="center")
        cell.fill = fill_header
        cell.border = box_border

    for r in range(9, 11):
        for c in range(1, 12):
            cell = ws.cell(row=r, column=c)
            cell.font = font_bold
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = fill_header
            cell.border = box_border

    # Dòng Số dư đầu kỳ (Row 12)
    for c in range(1, 12):
        ws.cell(row=12, column=c).border = box_border
    ws["G12"] = "Số dư đầu kỳ"
    ws["G12"].font = font_bold
    ws["K12"] = data["opening_balance"]
    ws["K12"].number_format = "#,##0"
    ws["K12"].font = font_bold
    ws["K12"].alignment = Alignment(horizontal="right")

    curr_row = 13
    for e in data["entries"]:
        ws.cell(row=curr_row, column=1, value=e["entry_date"]).alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=2, value=e["voucher_date"]).alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=3, value=e["receipt_no"]).alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=4, value=e["payment_no"]).alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=5, value=e.get("partner_code") or "").alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=6, value=e.get("partner_name") or "").alignment = Alignment(horizontal="left")
        ws.cell(row=curr_row, column=7, value=e["description"]).alignment = Alignment(horizontal="left")
        ws.cell(row=curr_row, column=8, value=e.get("offset_account") or "").alignment = Alignment(horizontal="center")

        c_rec = ws.cell(row=curr_row, column=9, value=e["receipt_amount"] if e["receipt_amount"] > 0 else "")
        c_rec.number_format = "#,##0"
        c_rec.alignment = Alignment(horizontal="right")

        c_pay = ws.cell(row=curr_row, column=10, value=e["payment_amount"] if e["payment_amount"] > 0 else "")
        c_pay.number_format = "#,##0"
        c_pay.alignment = Alignment(horizontal="right")

        c_bal = ws.cell(row=curr_row, column=11, value=e["balance"])
        c_bal.number_format = "#,##0"
        c_bal.alignment = Alignment(horizontal="right")

        for c in range(1, 12):
            c_cell = ws.cell(row=curr_row, column=c)
            c_cell.font = font_regular
            c_cell.border = box_border

        curr_row += 1

    # Dòng Cộng số phát sinh
    ws.merge_cells(f"A{curr_row}:H{curr_row}")
    ws[f"A{curr_row}"] = "Cộng số phát sinh trong kỳ"
    ws[f"A{curr_row}"].font = font_bold
    c_tot_rec = ws.cell(row=curr_row, column=9, value=data["total_receipt"])
    c_tot_rec.number_format = "#,##0"
    c_tot_rec.font = font_bold
    c_tot_rec.alignment = Alignment(horizontal="right")

    c_tot_pay = ws.cell(row=curr_row, column=10, value=data["total_payment"])
    c_tot_pay.number_format = "#,##0"
    c_tot_pay.font = font_bold
    c_tot_pay.alignment = Alignment(horizontal="right")

    ws.cell(row=curr_row, column=11, value="")
    for c in range(1, 12):
        ws.cell(row=curr_row, column=c).border = box_border
        ws.cell(row=curr_row, column=c).fill = fill_total

    curr_row += 1
    # Dòng Số dư cuối kỳ
    ws.merge_cells(f"A{curr_row}:H{curr_row}")
    ws[f"A{curr_row}"] = "Số dư cuối kỳ"
    ws[f"A{curr_row}"].font = font_bold
    c_close = ws.cell(row=curr_row, column=11, value=data["closing_balance"])
    c_close.number_format = "#,##0"
    c_close.font = font_bold
    c_close.alignment = Alignment(horizontal="right")
    for c in range(1, 12):
        ws.cell(row=curr_row, column=c).border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_double)
        ws.cell(row=curr_row, column=c).fill = fill_total

    col_widths = {1: 13, 2: 13, 3: 11, 4: 11, 5: 15, 6: 28, 7: 35, 8: 12, 9: 17, 10: 17, 11: 18}
    for c, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def export_bank_book_excel(
    company_id: int | None = None,
    year: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> bytes:
    """Export Bank Deposit Book to Excel matching Mau S05a-DNN."""
    company_id = company_id or get_current_company_id()
    if not company_id or int(company_id) <= 0:
        return b""
    data = get_bank_book_data(company_id=company_id, year=year, date_from=date_from, date_to=date_to)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S05a-DNN_So_Tien_Gui_NH"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Times New Roman", size=14, bold=True, color="1E3A8A")
    font_bold = Font(name="Times New Roman", size=10, bold=True)
    font_regular = Font(name="Times New Roman", size=10)
    font_italic = Font(name="Times New Roman", size=9, italic=True)

    border_thin = Side(border_style="thin", color="000000")
    border_double = Side(border_style="double", color="000000")
    box_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    fill_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_total = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")

    ws["A1"] = f"Đơn vị: {data['company_name']}"
    ws["A1"].font = font_bold
    ws["A2"] = f"Địa chỉ: {data['company_address']}"
    ws["A2"].font = font_regular
    ws["A3"] = f"Mã số thuế: {data['tax_code']}"
    ws["A3"].font = font_regular

    ws["G1"] = data["form_code"]
    ws["G1"].font = font_bold
    ws["G1"].alignment = Alignment(horizontal="right")
    ws["G2"] = data["circular"]
    ws["G2"].font = font_italic
    ws["G2"].alignment = Alignment(horizontal="right")

    ws.merge_cells("A5:J5")
    ws["A5"] = data["title"]
    ws["A5"].font = font_title
    ws["A5"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A6:J6")
    ws["A6"] = f"Số tài khoản: {data['bank_account_no']} tại {data['bank_name']}"
    ws["A6"].font = font_bold
    ws["A6"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A7:J7")
    ws["A7"] = data["date_str"]
    ws["A7"].font = font_italic
    ws["A7"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A9:A10")
    ws["A9"] = "Ngày, tháng\nghi sổ"
    ws.merge_cells("B9:C9")
    ws["B9"] = "Chứng từ"
    ws["B10"] = "Số hiệu"
    ws["C10"] = "Ngày tháng"
    ws.merge_cells("D9:D10")
    ws["D9"] = "Mã khách"
    ws.merge_cells("E9:E10")
    ws["E9"] = "Tên khách"
    ws.merge_cells("F9:F10")
    ws["F9"] = "Diễn giải"
    ws.merge_cells("G9:G10")
    ws["G9"] = "TK đối ứng"
    ws.merge_cells("H9:J9")
    ws["H9"] = "Số tiền"
    ws["H10"] = "Gửi vào"
    ws["I10"] = "Rút ra"
    ws["J10"] = "Còn lại (Tồn)"

    for r in range(9, 11):
        for c in range(1, 11):
            cell = ws.cell(row=r, column=c)
            cell.font = font_bold
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = fill_header
            cell.border = box_border

    # Row 11: Letters
    letters = ["A", "B", "C", "D", "E", "F", "G", "1", "2", "3"]
    for c_idx, ltr in enumerate(letters, 1):
        cell = ws.cell(row=11, column=c_idx, value=ltr)
        cell.font = font_bold
        cell.alignment = Alignment(horizontal="center")
        cell.fill = fill_header
        cell.border = box_border

    # Row 12: Số dư đầu kỳ
    ws["F12"] = "Số dư đầu kỳ"
    ws["F12"].font = font_bold
    ws["J12"] = data["opening_balance"]
    ws["J12"].number_format = "#,##0"
    ws["J12"].font = font_bold
    ws["J12"].alignment = Alignment(horizontal="right")
    for c in range(1, 11):
        ws.cell(row=12, column=c).border = box_border

    curr_row = 13
    for e in data["entries"]:
        ws.cell(row=curr_row, column=1, value=e["entry_date"]).alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=2, value=e["voucher_no"]).alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=3, value=e["voucher_date"]).alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=4, value=e.get("partner_code") or "").alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=5, value=e.get("partner_name") or "").alignment = Alignment(horizontal="left")
        ws.cell(row=curr_row, column=6, value=e["description"]).alignment = Alignment(horizontal="left")
        ws.cell(row=curr_row, column=7, value=e["offset_account"]).alignment = Alignment(horizontal="center")

        c_dep = ws.cell(row=curr_row, column=8, value=e["deposit_amount"] if e["deposit_amount"] > 0 else "")
        c_dep.number_format = "#,##0"
        c_dep.alignment = Alignment(horizontal="right")

        c_wth = ws.cell(row=curr_row, column=9, value=e["withdraw_amount"] if e["withdraw_amount"] > 0 else "")
        c_wth.number_format = "#,##0"
        c_wth.alignment = Alignment(horizontal="right")

        c_bal = ws.cell(row=curr_row, column=10, value=e["balance"])
        c_bal.number_format = "#,##0"
        c_bal.alignment = Alignment(horizontal="right")

        for c in range(1, 11):
            c_cell = ws.cell(row=curr_row, column=c)
            c_cell.font = font_regular
            c_cell.border = box_border

        curr_row += 1

    # Dòng Cộng phát sinh
    ws.merge_cells(f"A{curr_row}:G{curr_row}")
    ws[f"A{curr_row}"] = "Cộng số phát sinh trong kỳ"
    ws[f"A{curr_row}"].font = font_bold
    c_tot_dep = ws.cell(row=curr_row, column=8, value=data["total_deposit"])
    c_tot_dep.number_format = "#,##0"
    c_tot_dep.font = font_bold
    c_tot_dep.alignment = Alignment(horizontal="right")

    c_tot_wth = ws.cell(row=curr_row, column=9, value=data["total_withdraw"])
    c_tot_wth.number_format = "#,##0"
    c_tot_wth.font = font_bold
    c_tot_wth.alignment = Alignment(horizontal="right")

    for c in range(1, 11):
        ws.cell(row=curr_row, column=c).border = box_border
        ws.cell(row=curr_row, column=c).fill = fill_total

    curr_row += 1
    # Dòng Số dư cuối kỳ
    ws.merge_cells(f"A{curr_row}:G{curr_row}")
    ws[f"A{curr_row}"] = "Số dư cuối kỳ"
    ws[f"A{curr_row}"].font = font_bold
    c_close = ws.cell(row=curr_row, column=10, value=data["closing_balance"])
    c_close.number_format = "#,##0"
    c_close.font = font_bold
    c_close.alignment = Alignment(horizontal="right")
    for c in range(1, 11):
        ws.cell(row=curr_row, column=c).border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_double)
        ws.cell(row=curr_row, column=c).fill = fill_total

    col_widths = {1: 13, 2: 12, 3: 13, 4: 15, 5: 35, 6: 38, 7: 11, 8: 18, 9: 18, 10: 18}
    for c, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


# =====================================================================
# TẠO FILE MẪU EXCEL IMPORT VÀ IMPORT GIAO DỊCH HÀNG LOẠT
# =====================================================================

def generate_bank_import_template() -> bytes:
    """Tạo file mẫu Excel chuẩn để người dùng nhập giao dịch ngân hàng hàng loạt."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Giao_Dich_Ngan_Hang"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Times New Roman", size=13, bold=True, color="1E3A8A")
    font_header = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Times New Roman", size=11)
    font_guide = Font(name="Times New Roman", size=10, italic=True, color="475569")

    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_sample = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")

    thin = Side(border_style="thin", color="CBD5E1")
    box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "MẪU NHẬP GIAO DỊCH TIỀN GỬI NGÂN HÀNG"
    ws["A1"].font = font_title
    ws.merge_cells("A1:I1")

    ws["A2"] = "Hướng dẫn: Nhập giao dịch từ dòng 5 trở đi. Cột 'Loại giao dịch': Ghi 'Thu' (hoặc 'Báo Có') / 'Chi' (hoặc 'Báo Nợ' / 'UNC'). Cột Ngày: DD/MM/YYYY. Số tiền: nhập số tiền dương."
    ws["A2"].font = font_guide
    ws.merge_cells("A2:I2")

    headers = [
        "STT",
        "Ngày giao dịch (*)",
        "Loại giao dịch (*)",
        "Số chứng từ / UNC",
        "Số tiền (VNĐ) (*)",
        "Tài khoản đối ứng",
        "Khách hàng / Đối tác / NCC",
        "Diễn giải nội dung (*)",
        "Ghi chú kèm theo",
    ]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = box_border

    # 3 dòng mẫu
    samples = [
        (1, "20/05/2026", "Thu (Báo Có)", "BC-0001", 37260000, "131", "CÔNG TY CỔ PHẦN CÂN ĐIỆN TỬ VĨNH PHÚC", "Khách hàng chuyển khoản cước vận tải HĐ 92", "Kèm sao kê VietinBank"),
        (2, "22/05/2026", "Chi (Báo Nợ / UNC)", "UNC-0001", 15000000, "331", "CÔNG TY TNHH XĂNG DẦU PHÚ THỌ", "Thanh toán tiền dầu xe 19H-147.98", "HĐ xăng dầu số 00124"),
        (3, "25/05/2026", "Thu (Báo Có)", "", 50000000, "131", "CÔNG TY TNHH KHOA HỌC TÂN THỊNH PHÁT", "Chuyển khoản thanh toán tiền cước vận chuyển", "Tự sinh số CT nếu để trống"),
    ]

    for row_offset, r_data in enumerate(samples, start=5):
        for col_idx, val in enumerate(r_data, start=1):
            cell = ws.cell(row=row_offset, column=col_idx, value=val)
            cell.font = font_regular
            cell.fill = fill_sample
            cell.border = box_border
            if col_idx == 5:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")
            elif col_idx in (1, 2, 3, 4, 6):
                cell.alignment = Alignment(horizontal="center")

    col_widths = {1: 7, 2: 18, 3: 22, 4: 18, 5: 18, 6: 16, 7: 38, 8: 48, 9: 25}
    for c, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    ws.row_dimensions[4].height = 28

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_cash_import_template() -> bytes:
    """Tạo file mẫu Excel chuẩn để người dùng nhập thu chi tiền mặt hàng loạt."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "So_Quy_Tien_Mat"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Times New Roman", size=13, bold=True, color="1E3A8A")
    font_header = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Times New Roman", size=11)
    font_guide = Font(name="Times New Roman", size=10, italic=True, color="475569")

    fill_header = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
    fill_sample = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")

    thin = Side(border_style="thin", color="CBD5E1")
    box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "MẪU NHẬP THU CHI SỔ QUỸ TIỀN MẶT (TK 111)"
    ws["A1"].font = font_title
    ws.merge_cells("A1:I1")

    ws["A2"] = "Hướng dẫn: Nhập giao dịch từ dòng 5 trở đi. Cột 'Loại giao dịch': Ghi 'Thu' (Phiếu thu) / 'Chi' (Phiếu chi). Cột Ngày: DD/MM/YYYY. Số tiền: nhập số tiền dương."
    ws["A2"].font = font_guide
    ws.merge_cells("A2:I2")

    headers = [
        "STT",
        "Ngày chứng từ (*)",
        "Loại chứng từ (*)",
        "Số chứng từ",
        "Số tiền (VNĐ) (*)",
        "Tài khoản đối ứng",
        "Đối tác / Người nộp - nhận",
        "Diễn giải nội dung (*)",
        "Ghi chú kèm theo",
    ]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = box_border

    samples = [
        (1, "15/05/2026", "Thu (Phiếu Thu)", "PT001", 13500000, "131", "Cty CĐT Vĩnh Phúc", "Thu tiền cước vận chuyển HĐ 61", "Phiếu thu tiền mặt"),
        (2, "18/05/2026", "Chi (Phiếu Chi)", "PC001", 5000000, "642", "Văn phòng phẩm Ánh Dương", "Chi mua văn phòng phẩm quản lý", "Hóa đơn VPP số 0023"),
    ]

    for row_offset, r_data in enumerate(samples, start=5):
        for col_idx, val in enumerate(r_data, start=1):
            cell = ws.cell(row=row_offset, column=col_idx, value=val)
            cell.font = font_regular
            cell.fill = fill_sample
            cell.border = box_border
            if col_idx == 5:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")
            elif col_idx in (1, 2, 3, 4, 6):
                cell.alignment = Alignment(horizontal="center")

    col_widths = {1: 7, 2: 18, 3: 22, 4: 18, 5: 18, 6: 16, 7: 35, 8: 45, 9: 25}
    for c, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    ws.row_dimensions[4].height = 28

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def import_bank_transactions_from_excel(
    file_content: bytes,
    bank_account_no: str = "",
    bank_name: str = "",
    company_id: int | None = None,
) -> tuple[bool, str, int]:
    """Đọc file Excel và nạp hàng loạt giao dịch ngân hàng vào public.transport_cash_book."""
    company_id = company_id or get_current_company_id()
    if not company_id or int(company_id) <= 0:
        return False, "Chưa chọn công ty làm việc.", 0
    if not bank_account_no or not bank_name:
        comp_info = get_company_info(company_id)
        bank_account_no = bank_account_no or comp_info["bank_account"]
        bank_name = bank_name or comp_info["bank_name"]

    try:
        wb = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
    except Exception as e:
        return False, f"Không mở được file Excel: {e}", 0

    ws = wb.active
    success_count = 0

    # Tìm dòng bắt đầu dữ liệu (sau dòng tiêu đề cột)
    start_row = 5
    for r in range(1, 15):
        txt_b = str(ws.cell(row=r, column=2).value or "").lower()
        txt_c = str(ws.cell(row=r, column=3).value or "").lower()
        if "ngày" in txt_b or "loại" in txt_c or "ngày" in txt_c:
            start_row = r + 1
            break

    for r in range(start_row, ws.max_row + 1):
        raw_date = ws.cell(row=r, column=2).value
        raw_type = str(ws.cell(row=r, column=3).value or "").strip()
        raw_v_no = str(ws.cell(row=r, column=4).value or "").strip()
        raw_amount = ws.cell(row=r, column=5).value
        raw_offset = str(ws.cell(row=r, column=6).value or "").strip()
        raw_partner = str(ws.cell(row=r, column=7).value or "").strip()
        raw_desc = str(ws.cell(row=r, column=8).value or "").strip()
        raw_note = str(ws.cell(row=r, column=9).value or "").strip()

        # Dòng trống hoàn toàn -> bỏ qua
        if raw_date is None and not raw_type and raw_amount is None and not raw_desc:
            continue

        # Dòng hướng dẫn hoặc ghi chú mẫu -> bỏ qua
        if "hướng dẫn" in str(raw_date).lower() or "tự sinh số ct nếu để trống" in raw_note.lower():
            continue

        # Parse số tiền
        if raw_amount is None:
            continue
        try:
            cleaned_amt = str(raw_amount).replace(".", "").replace(",", "").replace(" ", "").strip()
            amt = Decimal(cleaned_amt)
            if amt <= 0:
                continue
        except Exception:
            continue

        # Parse ngày tháng
        v_date = date.today()
        if isinstance(raw_date, datetime):
            v_date = raw_date.date()
        elif isinstance(raw_date, date):
            v_date = raw_date
        elif isinstance(raw_date, str) and raw_date.strip():
            s = raw_date.strip()
            for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
                try:
                    v_date = datetime.strptime(s, fmt).date()
                    break
                except ValueError:
                    pass

        # Parse loại giao dịch: Thu (Báo Có) hay Chi (Báo Nợ / UNC)
        t_low = raw_type.lower()
        is_payment = (
            "chi" in t_low or "nợ" in t_low or "unc" in t_low or
            "rút" in t_low or "payment" in t_low or "out" in t_low
        )
        v_type = "payment" if is_payment else "receipt"

        if not raw_offset:
            raw_offset = "331" if is_payment else "131"

        if not raw_desc:
            raw_desc = "Rút ra thanh toán tiền gửi NH" if is_payment else "Gửi vào tài khoản ngân hàng"

        v_no = raw_v_no
        if not v_no:
            prefix = "UNC-" if is_payment else "BC-"
            count_row = db._execute_query('''
                SELECT COUNT(*) as cnt FROM public.transport_cash_book
                WHERE company_id = %s AND book_type = 'bank' AND voucher_type = %s
            ''', (company_id, v_type), fetch=True, commit=False) or {}
            num = (count_row.get("cnt") or 0) + success_count + 1
            v_no = f"{prefix}{num:04d}"

        db._execute_query('''
            INSERT INTO public.transport_cash_book
            (company_id, book_type, entry_date, voucher_date, voucher_no, voucher_type,
             description, offset_account, amount, bank_account_no, bank_name, partner_name,
             source_type, note)
            VALUES (%s, 'bank', %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'excel_import', %s)
        ''', (
            company_id, v_date, v_date, v_no, v_type,
            raw_desc, raw_offset, amt, bank_account_no, bank_name, raw_partner, raw_note
        ), commit=True)

        success_count += 1

    if success_count == 0:
        return False, "Không tìm thấy dòng giao dịch hợp lệ nào trong file Excel để nạp.", 0

    return True, f"Đã nạp thành công {success_count} giao dịch vào Sổ tiền gửi ngân hàng!", success_count


def import_cash_transactions_from_excel(
    file_content: bytes,
    company_id: int | None = None,
) -> tuple[bool, str, int]:
    """Đọc file Excel và nạp hàng loạt thu chi tiền mặt vào public.transport_cash_book."""
    company_id = company_id or get_current_company_id()
    if not company_id or int(company_id) <= 0:
        return False, "Chưa chọn công ty làm việc.", 0
    try:
        wb = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
    except Exception as e:
        return False, f"Không mở được file Excel: {e}", 0

    ws = wb.active
    success_count = 0

    start_row = 5
    for r in range(1, 15):
        txt_b = str(ws.cell(row=r, column=2).value or "").lower()
        txt_c = str(ws.cell(row=r, column=3).value or "").lower()
        if "ngày" in txt_b or "loại" in txt_c or "chứng từ" in txt_c:
            start_row = r + 1
            break

    for r in range(start_row, ws.max_row + 1):
        raw_date = ws.cell(row=r, column=2).value
        raw_type = str(ws.cell(row=r, column=3).value or "").strip()
        raw_v_no = str(ws.cell(row=r, column=4).value or "").strip()
        raw_amount = ws.cell(row=r, column=5).value
        raw_offset = str(ws.cell(row=r, column=6).value or "").strip()
        raw_partner = str(ws.cell(row=r, column=7).value or "").strip()
        raw_desc = str(ws.cell(row=r, column=8).value or "").strip()
        raw_note = str(ws.cell(row=r, column=9).value or "").strip()

        if raw_date is None and not raw_type and raw_amount is None and not raw_desc:
            continue
        if "hướng dẫn" in str(raw_date).lower():
            continue

        if raw_amount is None:
            continue
        try:
            cleaned_amt = str(raw_amount).replace(".", "").replace(",", "").replace(" ", "").strip()
            amt = Decimal(cleaned_amt)
            if amt <= 0:
                continue
        except Exception:
            continue

        v_date = date.today()
        if isinstance(raw_date, datetime):
            v_date = raw_date.date()
        elif isinstance(raw_date, date):
            v_date = raw_date
        elif isinstance(raw_date, str) and raw_date.strip():
            s = raw_date.strip()
            for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
                try:
                    v_date = datetime.strptime(s, fmt).date()
                    break
                except ValueError:
                    pass

        t_low = raw_type.lower()
        is_payment = ("chi" in t_low or "payment" in t_low or "out" in t_low or "pc" in t_low)
        v_type = "payment" if is_payment else "receipt"

        if not raw_offset:
            raw_offset = "331" if is_payment else "131"

        if not raw_desc:
            raw_desc = "Chi tiền mặt" if is_payment else "Thu tiền mặt"

        v_no = raw_v_no
        if not v_no:
            prefix = "PC" if is_payment else "PT"
            count_row = db._execute_query('''
                SELECT COUNT(*) as cnt FROM public.transport_cash_book
                WHERE company_id = %s AND book_type = 'cash' AND voucher_type = %s
            ''', (company_id, v_type), fetch=True, commit=False) or {}
            num = (count_row.get("cnt") or 0) + success_count + 1
            v_no = f"{prefix}{num:03d}"

        db._execute_query('''
            INSERT INTO public.transport_cash_book
            (company_id, book_type, entry_date, voucher_date, voucher_no, voucher_type,
             description, offset_account, amount, partner_name, source_type, note)
            VALUES (%s, 'cash', %s, %s, %s, %s, %s, %s, %s, %s, 'excel_import', %s)
        ''', (
            company_id, v_date, v_date, v_no, v_type,
            raw_desc, raw_offset, amt, raw_partner, raw_note
        ), commit=True)

        success_count += 1

    if success_count == 0:
        return False, "Không tìm thấy dòng chứng từ hợp lệ nào trong file Excel để nạp.", 0

    return True, f"Đã nạp thành công {success_count} chứng từ vào Sổ quỹ tiền mặt!", success_count
