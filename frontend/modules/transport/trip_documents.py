"""Trip documents service based on '1.3_Hợp đồng Đại Lộ Xanh 2026.xlsx'.

Provides:
1. Exporting the exact master Excel workbook configured for an individual invoice/trip:
   - DATA sheet updated with the invoice
   - DSKH sheet updated with the customer
   - BBVC, ĐNTT, BBTLHĐ, Hợp đồng sheets configured for the invoice
2. Rendering standard A4 HTML for:
   - BBVC: Biên bản vận chuyển
   - ĐNTT: Đề nghị thanh toán
   - BBTLHĐ: Biên bản nghiệm thu kiêm thanh lý hợp đồng
   - HDVC: Hợp đồng vận chuyển
   - ALL: Trọn bộ hồ sơ chuyến hàng
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from html import escape
import io
import logging
from pathlib import Path
import re
from typing import Any
from frontend.modules.gender_utils import split_rep_honorific, get_honorific

import openpyxl
import docx
from docx.shared import Pt, Mm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

from backend.services import config as cfg
from frontend.modules.transport.contract import (
    derive_company_code,
    _to_data_uri_if_local,
    _contract_bank_accounts,
    _remove_accents,
    _signature_print_width_mm,
)
from frontend.modules.transport.hdnt_bbbg import (
    number_to_vietnamese_words,
    _format_bank_lines,
)


def _sig_slot_html(sig_url: str = "", stamp_url: str = "") -> str:
    url = sig_url or stamp_url
    sig_data = _to_data_uri_if_local(url) if url else ""
    if sig_data:
        width_mm = _signature_print_width_mm(url)
        return f'<div class="sig-slot-inner"><img src="{sig_data}" alt="Chữ ký" style="width:{width_mm:g}mm"></div>'
    return '<div class="sig-slot-inner"></div>'

logger = logging.getLogger(__name__)

TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "assets" / "DLX_chuyen_hang_template.xlsx"


def _extract_trip_data(invoice: dict, lines: list[dict], customer: dict, seller: dict | None = None) -> dict[str, Any]:
    raw_date = invoice.get("invoice_date")
    if isinstance(raw_date, datetime):
        c_date = raw_date.date()
    elif isinstance(raw_date, date):
        c_date = raw_date
    elif isinstance(raw_date, str) and raw_date:
        try:
            c_date = date.fromisoformat(raw_date[:10])
        except Exception:
            c_date = date.today()
    else:
        c_date = date.today()

    inv_num = int(invoice.get("invoice_number") or 1)
    inv_num_fmt = f"{inv_num:08d}"
    inv_serial = str(invoice.get("invoice_serial") or "C26TAA").strip()

    seller = seller or {}
    seller_name = str(seller.get("name") or "").strip()
    seller_code = derive_company_code(seller_name) or "SELLER"
    seller_tax = str(seller.get("tax_code") or "").strip()
    seller_rep = str(seller.get("director_name") or "").strip()
    seller_pos = str(seller.get("director_title") or "Giám đốc").strip()
    seller_addr = str(seller.get("address") or "").strip()
    seller_phone = str(seller.get("phone") or "").strip()
    seller_email = str(seller.get("email") or "").strip()
    seller_bank_acc = str(seller.get("bank_account") or "").strip()
    seller_bank_name = str(seller.get("bank_name") or "").strip()
    seller_sig_url = str(seller.get("signature_url") or "").strip()
    seller_stamp_url = str(seller.get("stamp_url") or "").strip()

    customer = customer or {}
    customer_name = str(customer.get("name") or "").strip()
    cust_code = derive_company_code(customer_name) or "KH"
    customer_tax = str(customer.get("tax_code") or invoice.get("buyer_tax_code") or "").strip()
    customer_rep = str(customer.get("director_name") or "").strip()
    customer_pos = str(customer.get("director_title") or "Giám đốc").strip()
    customer_addr = str(customer.get("address") or "").strip()
    customer_phone = str(customer.get("phone") or "").strip()
    customer_bank_acc = str(customer.get("bank_account") or "").strip()
    customer_bank_name = str(customer.get("bank_name") or "").strip()
    customer_sig_url = str(customer.get("signature_url") or "").strip()
    customer_stamp_url = str(customer.get("stamp_url") or "").strip()

    seller_gender = seller.get("director_gender") or seller.get("gender")
    customer_gender = customer.get("director_gender") or customer.get("gender")
    seller_prefix, seller_clean_name = split_rep_honorific(seller_rep, seller_gender)
    customer_prefix, customer_clean_name = split_rep_honorific(customer_rep, customer_gender)
    seller_honorific = get_honorific(seller_rep, seller_gender)
    customer_honorific = get_honorific(customer_rep, customer_gender)

    items: list[dict[str, Any]] = []
    plate_numbers: list[str] = []
    origins: list[str] = []
    destinations: list[str] = []

    for idx, line in enumerate(lines or []):
        amount = int(round(float(line.get("amount") or 0)))
        desc = str(line.get("description") or line.get("cargo_description") or "").strip()
        plate = str(line.get("plate_number") or "").strip()
        if plate and plate not in plate_numbers:
            plate_numbers.append(plate)
        orig = str(line.get("origin") or "").strip()
        dest = str(line.get("destination") or "").strip()
        if orig and orig not in origins:
            origins.append(orig)
        if dest and dest not in destinations:
            destinations.append(dest)
        items.append({
            "index": idx + 1,
            "description": desc or "Cước vận chuyển hàng hóa",
            "plate_number": plate,
            "unit": "Chuyến",
            "quantity": 1,
            "price": amount,
            "amount": amount,
        })

    subtotal = sum(it["amount"] for it in items)
    if subtotal == 0:
        subtotal = int(round(float(invoice.get("subtotal") or 0)))
    vat_rate = float(invoice.get("vat_rate") if invoice.get("vat_rate") is not None else 0.08)
    vat_percent = int(round(vat_rate * 100))
    vat_amount = int(round(subtotal * vat_rate))
    total_amount = subtotal + vat_amount
    total_in_words = number_to_vietnamese_words(total_amount)

    primary_plate = ", ".join(plate_numbers) if plate_numbers else ""
    primary_origin = ", ".join(origins) if origins else ""
    primary_dest = ", ".join(destinations) if destinations else ""

    contract_no = invoice.get("contract_no") or f"{c_date.day:02d}{c_date.month:02d}{str(c_date.year)[-2:]}/HĐNT/{cust_code}-{seller_code}"
    bbvc_no = invoice.get("bbvc_no") or (contract_no if "BBVC" in str(contract_no) else f"{c_date.day:02d}{c_date.month:02d}/{c_date.year}/BBVC")
    dntt_no = invoice.get("dntt_no") or (contract_no if ("ĐNTT" in str(contract_no) or "DNTT" in str(contract_no)) else f"{inv_num}/ĐNTT-ĐLX")

    return {
        "c_date": c_date,
        "day": f"{c_date.day:02d}",
        "month": f"{c_date.month:02d}",
        "year": str(c_date.year),
        "invoice_number": inv_num,
        "invoice_number_formatted": inv_num_fmt,
        "invoice_serial": inv_serial,
        "contract_no": contract_no,
        "bbvc_no": bbvc_no,
        "dntt_no": dntt_no,
        "seller_name": seller_name,
        "seller_code": seller_code,
        "seller_tax": seller_tax,
        "seller_rep": seller_rep,
        "seller_gender": seller_gender,
        "seller_honorific": seller_honorific,
        "seller_prefix": seller_prefix,
        "seller_clean_rep": seller_clean_name or seller_rep,
        "seller_pos": seller_pos,
        "seller_addr": seller_addr,
        "seller_phone": seller_phone,
        "seller_email": seller_email,
        "seller_bank_acc": seller_bank_acc,
        "seller_bank_name": seller_bank_name,
        "seller_sig_url": seller_sig_url,
        "seller_stamp_url": seller_stamp_url,
        "customer_name": customer_name,
        "cust_code": cust_code,
        "customer_tax": customer_tax,
        "customer_rep": customer_rep,
        "customer_gender": customer_gender,
        "customer_honorific": customer_honorific,
        "customer_prefix": customer_prefix,
        "customer_clean_rep": customer_clean_name or customer_rep,
        "customer_pos": customer_pos,
        "customer_addr": customer_addr,
        "customer_phone": customer_phone,
        "customer_bank_acc": customer_bank_acc,
        "customer_bank_name": customer_bank_name,
        "customer_sig_url": customer_sig_url,
        "customer_stamp_url": customer_stamp_url,
        "items": items,
        "primary_plate": primary_plate,
        "primary_origin": primary_origin,
        "primary_dest": primary_dest,
        "subtotal": subtotal,
        "vat_rate": vat_rate,
        "vat_percent": vat_percent,
        "vat_amount": vat_amount,
        "total_amount": total_amount,
        "total_in_words": total_in_words,
    }


def build_dlx_trip_workbook(
    invoice: dict,
    lines: list[dict],
    customer: dict,
    seller: dict | None = None,
) -> bytes:
    """Build and return master Excel workbook customized for this specific trip/invoice."""
    info = _extract_trip_data(invoice, lines, customer, seller)
    wb = openpyxl.load_workbook(str(TEMPLATE_PATH), data_only=False)

    # 1. Update or Insert invoice into DATA sheet
    if "DATA" in wb.sheetnames:
        ws_data = wb["DATA"]
        matched_row = None
        for r in range(7, ws_data.max_row + 1):
            val = ws_data.cell(r, 2).value
            if val == info["invoice_number"]:
                matched_row = r
                break

        first_desc = info["items"][0]["description"] if info["items"] else "Cước vận chuyển hàng hóa"
        first_price = info["items"][0]["amount"] if info["items"] else info["subtotal"]

        if matched_row:
            target_row = matched_row
        else:
            # Insert a new row right before the Total row (if row 111 is Total)
            last_r = ws_data.max_row
            total_r = None
            for r in range(last_r, 6, -1):
                if str(ws_data.cell(r, 1).value).lower().startswith("total"):
                    total_r = r
                    break
            if total_r:
                ws_data.insert_rows(total_r, amount=1)
                target_row = total_r
            else:
                target_row = last_r + 1

        # Populate DATA row
        ws_data.cell(target_row, 1, f'=IF(B{target_row}=BBVC!$R$2,1,"")')
        ws_data.cell(target_row, 2, info["invoice_number"])
        ws_data.cell(target_row, 3, datetime.combine(info["c_date"], datetime.min.time()))
        ws_data.cell(target_row, 3).number_format = "yyyy-mm-dd hh:mm:ss"
        ws_data.cell(target_row, 4, info["customer_tax"])
        ws_data.cell(target_row, 5, info["customer_name"])
        ws_data.cell(target_row, 6, info["customer_addr"])
        ws_data.cell(target_row, 7, first_desc)
        ws_data.cell(target_row, 8, "Chuyến")
        ws_data.cell(target_row, 9, 1)
        ws_data.cell(target_row, 10, first_price)
        ws_data.cell(target_row, 11, first_price)

    # 2. Update or Insert customer in DSKH sheet
    if "DSKH" in wb.sheetnames and info["customer_tax"]:
        ws_dskh = wb["DSKH"]
        found_cust = False
        for r in range(7, ws_dskh.max_row + 1):
            if str(ws_dskh.cell(r, 3).value or "").strip() == info["customer_tax"]:
                found_cust = True
                ws_dskh.cell(r, 4, info["customer_name"])
                ws_dskh.cell(r, 5, info["customer_addr"])
                ws_dskh.cell(r, 6, info["customer_rep"])
                ws_dskh.cell(r, 8, info["customer_bank_acc"])
                ws_dskh.cell(r, 9, info["customer_bank_name"])
                break
        if not found_cust:
            new_r = ws_dskh.max_row + 1
            ws_dskh.cell(new_r, 1, new_r - 6)
            ws_dskh.cell(new_r, 3, info["customer_tax"])
            ws_dskh.cell(new_r, 4, info["customer_name"])
            ws_dskh.cell(new_r, 5, info["customer_addr"])
            ws_dskh.cell(new_r, 6, info["customer_rep"])
            ws_dskh.cell(new_r, 7, info["customer_pos"])
            ws_dskh.cell(new_r, 8, info["customer_bank_acc"])
            ws_dskh.cell(new_r, 9, info["customer_bank_name"])

    # 3. Configure BBVC sheet
    if "BBVC" in wb.sheetnames:
        ws_bbvc = wb["BBVC"]
        ws_bbvc["R2"] = info["invoice_number"]
        ws_bbvc["S4"] = int(info["day"])
        ws_bbvc["E12"] = info["primary_plate"]
        # Populate explicit line values in BBVC rows 24..26
        for idx, it in enumerate(info["items"][:3]):
            row_idx = 24 + idx
            ws_bbvc.cell(row_idx, 3, it["description"])
            ws_bbvc.cell(row_idx, 12, it["unit"])
            ws_bbvc.cell(row_idx, 13, it["quantity"])
            ws_bbvc.cell(row_idx, 14, it["price"])
            ws_bbvc.cell(row_idx, 15, it["amount"])
        ws_bbvc["O27"] = info["subtotal"]
        ws_bbvc["O28"] = info["vat_amount"]
        ws_bbvc["O29"] = info["total_amount"]
        ws_bbvc["F30"] = info["total_in_words"]

    # 4. Configure ĐNTT sheet
    if "ĐNTT" in wb.sheetnames:
        ws_dntt = wb["ĐNTT"]
        ws_dntt["N2"] = info["invoice_number"]
        ws_dntt["J5"] = f"ngày {info['day']} tháng {info['month']} năm {info['year']}"
        first_desc = info["items"][0]["description"] if info["items"] else "hàng hóa"
        ws_dntt["N5"] = first_desc
        ws_dntt["N6"] = info["primary_origin"]
        ws_dntt["N7"] = info["primary_dest"]
        ws_dntt["P4"] = info["total_amount"]
        ws_dntt["E12"] = info["total_amount"]
        ws_dntt["K14"] = info["total_amount"]
        ws_dntt["D15"] = info["total_in_words"]

    # 5. Configure Hợp đồng sheet
    if "Hợp đồng " in wb.sheetnames:
        ws_hd = wb["Hợp đồng "]
        ws_hd["T1"] = int(info["day"])
        ws_hd["T2"] = info["month"]
        ws_hd["T5"] = info["customer_tax"]
        ws_hd["T6"] = info["cust_code"]
        ws_hd["F15"] = info["customer_name"]
        ws_hd["C16"] = info["customer_rep"]
        ws_hd["B17"] = info["customer_addr"]
        ws_hd["B19"] = info["customer_tax"]
        ws_hd["B20"] = info["customer_bank_acc"]
        ws_hd["M20"] = info["customer_bank_name"]

    # 6. Configure BBTLHĐ sheet
    if "BBTLHĐ" in wb.sheetnames:
        ws_bbtl = wb["BBTLHĐ"]
        ws_bbtl["Q1"] = int(info["day"])
        ws_bbtl["Q2"] = info["month"]
        ws_bbtl["Q5"] = info["customer_tax"]
        ws_bbtl["B10"] = datetime.combine(info["c_date"], datetime.min.time())
        ws_bbtl["F8"] = info["customer_name"]
        ws_bbtl["B11"] = info["customer_name"]
        ws_bbtl["C12"] = info["customer_rep"]
        ws_bbtl["B13"] = info["customer_addr"]
        ws_bbtl["B14"] = info["customer_bank_acc"]
        ws_bbtl["K14"] = info["customer_bank_name"]
        ws_bbtl["B15"] = info["customer_tax"]
        job_summary = "; ".join(it["description"] for it in info["items"])
        ws_bbtl["A26"] = f"1. Bên B đã thực hiện các công việc như sau: {job_summary}"
        ws_bbtl["A30"] = (
            f"thanh toán cho Bên B chi phí thực hiện hợp đồng: {info['total_amount']:,} đồng "
            f"({info['total_in_words']}./.)(Giá trên đã bao gồm thuế GTGT)."
        ).replace(",", ".")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# HTML DOCUMENT RENDERERS (A4 Standard Print-Ready)
# ---------------------------------------------------------------------------

_COMMON_CSS = """
  @page {
    size: A4 portrait;
    margin: 15mm 15mm 15mm 15mm;
  }
  * { box-sizing: border-box; }
  body {
    margin: 0;
    padding: 0;
    background: #525659;
    font-family: "Times New Roman", Times, serif;
    font-size: 13pt;
    line-height: 1.35;
    color: #000;
  }
  .a4-page {
    width: 210mm;
    min-height: 297mm;
    margin: 15px auto;
    padding: 20mm 18mm 20mm 20mm;
    background: #fff;
    box-shadow: 0 4px 15px rgba(0,0,0,0.3);
    position: relative;
    page-break-after: always;
  }
  .a4-page:last-child {
    page-break-after: avoid;
  }
  @media print {
    body { background: #fff !important; }
    .a4-page {
      margin: 0 !important;
      padding: 0 !important;
      box-shadow: none !important;
      width: 100% !important;
      min-height: auto !important;
    }
    .no-print { display: none !important; }
  }
  .text-center { text-align: center; }
  .text-right { text-align: right; }
  .text-justify { text-align: justify; }
  .bold { font-weight: bold; }
  .italic { font-style: italic; }
  .uppercase { text-transform: uppercase; }
  
  table.layout-table {
    width: 100%;
    border-collapse: collapse;
    border: none;
    margin-bottom: 6px;
  }
  table.layout-table td {
    padding: 2px 4px;
    vertical-align: top;
    font-size: 13pt;
  }
  table.grid-table {
    width: 100%;
    border-collapse: collapse;
    margin: 10px 0;
  }
  table.grid-table th, table.grid-table td {
    border: 1px solid #000;
    padding: 5px 6px;
    font-size: 12pt;
    vertical-align: middle;
  }
  table.grid-table th {
    background: #f1f5f9;
    text-align: center;
    font-weight: bold;
  }
  @media print {
    table.grid-table th { background: transparent !important; }
  }
  .national-header {
    width: 100%;
    margin-bottom: 12px;
  }
  .doc-title {
    font-size: 16pt;
    font-weight: bold;
    text-align: center;
    margin: 12px 0 4px 0;
    text-transform: uppercase;
  }
  .doc-subtitle {
    font-size: 12pt;
    text-align: center;
    margin-bottom: 14px;
    font-style: italic;
  }
  .sig-table {
    width: 100%;
    margin-top: 15px;
    border-collapse: collapse;
    page-break-inside: avoid;
    break-inside: avoid;
  }
  .sig-header {
    width: 50%;
    text-align: center;
    vertical-align: top;
    padding-bottom: 0;
    font-size: 12pt;
  }
  .sig-slot {
    width: 50%;
    height: 180px;
    min-height: 180px;
    max-height: 180px;
    text-align: center;
    vertical-align: middle;
    padding: 0;
  }
  .sig-slot-inner {
    height: 180px;
    display: flex;
    align-items: center;
    justify-content: center;
  }
  .sig-slot-inner img {
    max-height: none;
    max-width: none;
    object-fit: contain;
  }
"""


def _render_toolbar(info: dict, active_kind: str, invoice_guid: str | None = None) -> str:
    guid_val = invoice_guid or ""
    return f"""
    <div class="no-print" style="position: sticky; top: 0; z-index: 1000; background: #0f172a; color: #fff; padding: 10px 18px; display: flex; align-items: center; justify-content: space-between; gap: 12px; box-shadow: 0 4px 12px rgba(0,0,0,0.3); font-family: system-ui, -apple-system, sans-serif;">
      <div style="display: flex; align-items: center; gap: 12px; flex-wrap: wrap;">
        <span style="background: #2563eb; color: #fff; font-size: 11px; font-weight: 700; padding: 3px 8px; border-radius: 4px; text-transform: uppercase;">Chứng từ chuyến hàng</span>
        <span style="font-weight: 600; font-size: 13.5px; color: #e2e8f0;">HĐ: <strong style="color: #60a5fa;">{info['invoice_number_formatted']}</strong> &bull; KH: <strong style="color: #f1f5f9;">{escape(info['customer_name'])}</strong></span>
        
        <div style="display: flex; align-items: center; gap: 6px; margin-left: 8px;">
          <label for="kind-select" style="font-size: 12px; color: #94a3b8; margin: 0;">Loại chứng từ:</label>
          <select id="kind-select" onchange="changeDocKind(this.value)" style="background: #1e293b; color: #f8fafc; border: 1px solid #475569; padding: 5px 10px; border-radius: 5px; font-size: 13px; font-weight: 500; cursor: pointer; outline: none;">
            <option value="hdvc" {"selected" if active_kind == "hdvc" else ""}>Hợp đồng vận chuyển</option>
            <option value="bbvc" {"selected" if active_kind == "bbvc" else ""}>Biên bản vận chuyển (BBVC)</option>
            <option value="dntt" {"selected" if active_kind == "dntt" else ""}>Đề nghị thanh toán (ĐNTT)</option>
            <option value="bbtl" {"selected" if active_kind == "bbtl" else ""}>Biên bản nghiệm thu & thanh lý</option>
          </select>
        </div>
      </div>
      
      <div style="display: flex; align-items: center; gap: 8px; flex-wrap: wrap;">
        <button type="button" onclick="window.print()" style="background: #2563eb; color: #fff; border: none; padding: 6px 14px; border-radius: 5px; font-weight: 600; cursor: pointer; font-size: 12.5px; display: inline-flex; align-items: center; gap: 5px;">
          <svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><path d="M6 9V2h12v7M6 18H4a2 2 0 0 1-2-2v-5a2 2 0 0 1 2-2h16a2 2 0 0 1 2 2v5a2 2 0 0 1-2 2h-2"/><rect x="6" y="14" width="12" height="8"/></svg>
          In / Lưu PDF
        </button>
        <button type="button" onclick="downloadDocAsset('pdf')" style="background: #059669; color: #fff; border: none; padding: 6px 14px; border-radius: 5px; font-weight: 600; cursor: pointer; font-size: 12.5px; display: inline-flex; align-items: center; gap: 5px;">
          <svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4"/><polyline points="7 10 12 15 17 10"/><line x1="12" y1="15" x2="12" y2="3"/></svg>
          Tải PDF
        </button>
        <button type="button" onclick="downloadDocAsset('docx')" style="background: #0284c7; color: #fff; border: none; padding: 6px 14px; border-radius: 5px; font-weight: 600; cursor: pointer; font-size: 12.5px; display: inline-flex; align-items: center; gap: 5px;">
          <svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"/><polyline points="14 2 14 8 20 8"/><line x1="16" y1="13" x2="8" y2="13"/><line x1="16" y1="17" x2="8" y2="17"/></svg>
          Tải Word (.docx)
        </button>
      </div>
    </div>
    
    <script>
      function changeDocKind(k) {{
        var url = new URL(window.location.href);
        url.searchParams.set('kind', k);
        url.searchParams.set('view', 'html');
        window.location.href = url.toString();
      }}
      function downloadDocAsset(fmt) {{
        var url = new URL(window.location.href);
        url.searchParams.delete('view');
        url.searchParams.set('format', fmt);
        url.searchParams.set('download', fmt);
        window.location.href = url.toString();
      }}
    </script>
    """


def render_bbvc_page(info: dict) -> str:
    """Render Biên bản vận chuyển (BBVC) page."""
    seller_sig_url = info.get("seller_sig_url") or info.get("seller_signature_url") or info.get("seller_stamp_url") or ""
    customer_sig_url = info.get("customer_sig_url") or info.get("customer_signature_url") or info.get("customer_stamp_url") or ""
    seller_sig_html = _sig_slot_html(seller_sig_url)
    customer_sig_html = _sig_slot_html(customer_sig_url)

    item_rows = []
    for it in info["items"]:
        item_rows.append(f"""
        <tr>
          <td class="text-center">{it['index']}</td>
          <td>{escape(it['description'])}</td>
          <td class="text-center bold">{escape(it['plate_number'] or info['primary_plate'])}</td>
          <td class="text-center">{escape(it['unit'])}</td>
          <td class="text-center">{it['quantity']}</td>
          <td class="text-right">{it['price']:,}</td>
          <td class="text-right bold">{it['amount']:,}</td>
        </tr>
        """.replace(",", "."))

    return f"""
    <div class="a4-page">
      <div class="national-header text-center">
        <div class="bold uppercase" style="font-size: 12pt;">CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</div>
        <div class="bold" style="font-size: 13pt;">Độc lập - Tự do - Hạnh phúc</div>
        <div style="letter-spacing: -1px; margin-top: -2px;">---------------o0o---------------</div>
      </div>

      <div class="doc-title">BIÊN BẢN VẬN CHUYỂN</div>
      <div class="doc-subtitle">Số : {info['bbvc_no']}</div>

      <div style="margin-bottom: 8px;">
        Hôm nay, ngày {info['day']} tháng {info['month']} năm {info['year']}. Tại văn phòng quý khách hàng, chúng tôi gồm có:
      </div>

      <table class="layout-table">
        <tr>
          <td colspan="2" class="bold uppercase" style="padding-top: 4px;">BÊN VẬN CHUYỂN (BÊN A): {escape(info['seller_name'])}</td>
        </tr>
        <tr>
          <td style="width: 130px;">MST:</td>
          <td class="bold">{escape(info['seller_tax'])}</td>
        </tr>
        <tr>
          <td>Địa chỉ:</td>
          <td>{escape(info['seller_addr'])}</td>
        </tr>
        <tr>
          <td>Đại diện:</td>
          <td>{info['seller_honorific']} <span class="bold">{escape(info['seller_clean_rep'])}</span> &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Chức vụ: <span class="bold">{escape(info['seller_pos'])}</span></td>
        </tr>
        <tr>
          <td>Biển số xe:</td>
          <td class="bold">{escape(info['primary_plate'])}</td>
        </tr>

        <tr>
          <td colspan="2" class="bold uppercase" style="padding-top: 10px;">BÊN THUÊ VẬN CHUYỂN (BÊN B): {escape(info['customer_name'])}</td>
        </tr>
        <tr>
          <td>MST:</td>
          <td class="bold">{escape(info['customer_tax'])}</td>
        </tr>
        <tr>
          <td>Địa chỉ:</td>
          <td>{escape(info['customer_addr'])}</td>
        </tr>
        <tr>
          <td>Đại diện:</td>
          <td>{info['customer_honorific']} <span class="bold">{escape(info['customer_clean_rep'])}</span> &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Chức vụ: <span class="bold">{escape(info['customer_pos'])}</span></td>
        </tr>
      </table>

      <div class="bold" style="margin-top: 10px;">I. NỘI DUNG VẬN CHUYỂN:</div>
      <div class="text-justify" style="margin-bottom: 6px;">
        Bên B đồng ý giao, Bên A đồng ý nhận vận chuyển hàng hóa, máy móc thiết bị, vật tư với số lượng và đơn giá vận chuyển cụ thể như sau:
      </div>

      <table class="grid-table">
        <thead>
          <tr>
            <th style="width: 40px;">STT</th>
            <th>Tên hàng hóa, dịch vụ</th>
            <th style="width: 120px;">Biển số xe</th>
            <th style="width: 60px;">ĐVT</th>
            <th style="width: 55px;">SL</th>
            <th style="width: 100px;">Đơn giá</th>
            <th style="width: 110px;">Thành tiền</th>
          </tr>
        </thead>
        <tbody>
          {"".join(item_rows)}
          <tr>
            <td colspan="6" class="text-right bold">Cộng:</td>
            <td class="text-right bold">{info['subtotal']:,}</td>
          </tr>
          <tr>
            <td colspan="6" class="text-right bold">Thuế GTGT {info['vat_percent']}%:</td>
            <td class="text-right bold">{info['vat_amount']:,}</td>
          </tr>
          <tr>
            <td colspan="6" class="text-right bold" style="color: #b91c1c;">Tổng thanh toán:</td>
            <td class="text-right bold" style="color: #b91c1c;">{info['total_amount']:,}</td>
          </tr>
        </tbody>
      </table>

      <div class="italic" style="margin: 6px 0 10px 0;">
        <span class="bold">Số tiền bằng chữ:</span> {info['total_in_words']}
      </div>

      <div class="bold">II. KẾT LUẬN:</div>
      <div class="text-justify" style="margin-bottom: 4px;">
        Hai bên cùng nhau xác nhận đã giao nhận đúng, đủ số lượng hàng vận chuyển theo đúng lộ trình đã thống nhất.
      </div>
      <div class="text-justify">
        Biên bản vận chuyển được lập thành 02 bản, có giá trị pháp lý như nhau, mỗi bên giữ 01 bản.
      </div>

      <table class="sig-table">
        <tr>
          <td class="sig-header">
            <div class="bold uppercase">ĐẠI DIỆN BÊN A</div>
            <div class="italic small">(Ký, ghi rõ họ tên, đóng dấu)</div>
          </td>
          <td class="sig-header">
            <div class="bold uppercase">ĐẠI DIỆN BÊN B</div>
            <div class="italic small">(Ký, ghi rõ họ tên)</div>
          </td>
        </tr>
        <tr>
          <td class="sig-slot">
            {seller_sig_html}
          </td>
          <td class="sig-slot">
            {customer_sig_html}
          </td>
        </tr>
      </table>
    </div>
    """.replace(":,}", ":,}".replace(",", "."))


def render_dntt_page(info: dict) -> str:
    """Render Đề nghị thanh toán (ĐNTT) page."""
    seller_sig_url = info.get("seller_sig_url") or info.get("seller_signature_url") or info.get("seller_stamp_url") or ""
    seller_sig_html = _sig_slot_html(seller_sig_url)

    first_item = info["items"][0] if info["items"] else {}
    cargo_name = first_item.get("description") or "hàng hóa"

    return f"""
    <div class="a4-page">
      <table class="layout-table" style="margin-bottom: 15px;">
        <tr>
          <td style="width: 50%;">
            <div class="bold uppercase" style="font-size: 11pt;">{escape(info['seller_name'])}</div>
            <div style="font-size: 11pt;">Số: <span class="bold">{info['dntt_no']}</span></div>
            <div style="font-size: 11pt;">V/v: Thanh toán hóa đơn</div>
          </td>
          <td style="width: 50%;" class="text-center">
            <div class="bold uppercase" style="font-size: 11pt;">CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</div>
            <div class="bold" style="font-size: 12pt;">Độc lập – Tự do - Hạnh phúc</div>
            <div style="letter-spacing: -1px; margin-top: -2px;">--------------------</div>
            <div class="italic" style="font-size: 11pt; margin-top: 4px;">Phú Thọ, ngày {info['day']} tháng {info['month']} năm {info['year']}</div>
          </td>
        </tr>
      </table>

      <div class="doc-title" style="margin: 25px 0 15px 0;">ĐỀ NGHỊ THANH TOÁN</div>

      <div style="margin: 15px 0 10px 0; font-size: 13.5pt;">
        <span class="bold">Kính gửi:</span> <span class="bold uppercase">{escape(info['customer_name'])}</span>
      </div>

      <div class="text-justify" style="margin-bottom: 12px; text-indent: 25px;">
        Trước tiên xin chân thành cảm ơn quý công ty đã tin tưởng và lựa chọn Công ty chúng tôi.
      </div>

      <div class="text-justify" style="margin-bottom: 12px; text-indent: 25px;">
        Chúng tôi đã vận chuyển <strong>{escape(cargo_name)}</strong> từ <strong>{escape(info['primary_origin'])}</strong> đi <strong>{escape(info['primary_dest'])}</strong> BKS <strong>{escape(info['primary_plate'])}</strong>.
      </div>

      <div class="text-justify" style="margin-bottom: 12px; text-indent: 25px;">
        Tổng cước vận chuyển là: <strong style="color: #b91c1c;">{info['total_amount']:,} đồng</strong> (đã bao gồm thuế GTGT) và đã phát hành Hóa đơn số: <strong>{info['invoice_number_formatted']}</strong> ngày {info['day']}/{info['month']}/{info['year']}.
      </div>

      <div class="text-justify" style="margin-bottom: 10px; text-indent: 25px;">
        Bằng công văn này, kính đề nghị Quý Công ty thanh toán cho chúng tôi số tiền là: <strong style="color: #b91c1c;">{info['total_amount']:,} đồng</strong>.
      </div>

      <div class="italic" style="margin-bottom: 20px; padding-left: 25px;">
        (Số tiền bằng chữ: <strong>{info['total_in_words']}</strong>).
      </div>

      <div class="bold" style="margin-bottom: 6px;">Thông Tin Thụ Hưởng:</div>
      <div style="padding-left: 20px; line-height: 1.6;">
        <div>- Đơn vị thụ hưởng: <span class="bold uppercase">{escape(info['seller_name'])}</span>.</div>
        <div>- Số tài khoản: <span class="bold" style="font-size: 14pt;">{escape(info['seller_bank_acc'])}</span></div>
        <div>- Tại ngân hàng: <span class="bold">{escape(info['seller_bank_name'])}</span>.</div>
      </div>

      <div style="margin: 20px 0 35px 25px;">
        Xin trân trọng cảm ơn!
      </div>

      <table class="sig-table">
        <tr>
          <td style="width: 50%; text-align: left; padding-left: 15px; vertical-align: top;">
            <div class="bold italic">Nơi nhận:</div>
            <div style="font-size: 11pt;">- Như trên</div>
            <div style="font-size: 11pt;">- Lưu VP Công ty</div>
          </td>
          <td class="sig-header" style="width: 50%;">
            <div class="bold uppercase">{escape(info['seller_name'])}</div>
            <div class="italic small">(Ký, ghi rõ họ tên, đóng dấu)</div>
          </td>
        </tr>
        <tr>
          <td></td>
          <td class="sig-slot">
            {seller_sig_html}
          </td>
        </tr>
      </table>
    </div>
    """.replace(":,}", ":,}".replace(",", "."))


def render_bbtl_page(info: dict) -> str:
    """Render Biên bản nghiệm thu kiêm thanh lý hợp đồng (BBTLHĐ) page."""
    seller_sig_url = info.get("seller_sig_url") or info.get("seller_signature_url") or info.get("seller_stamp_url") or ""
    customer_sig_url = info.get("customer_sig_url") or info.get("customer_signature_url") or info.get("customer_stamp_url") or ""
    seller_sig_html = _sig_slot_html(seller_sig_url)
    customer_sig_html = _sig_slot_html(customer_sig_url)

    job_lines = []
    for it in info["items"]:
        job_lines.append(escape(it["description"]))
    job_desc = "; ".join(job_lines)

    return f"""
    <div class="a4-page">
      <div class="national-header text-center">
        <div class="bold uppercase" style="font-size: 12pt;">CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM</div>
        <div class="bold" style="font-size: 13pt;">Độc lập – Tự do – Hạnh phúc</div>
        <div style="letter-spacing: -1px; margin-top: -2px;">---------------o0o---------------</div>
        <div class="italic text-right" style="font-size: 11.5pt; margin-top: 4px;">Phú Thọ, ngày {info['day']} tháng {info['month']} năm {info['year']}</div>
      </div>

      <div class="doc-title" style="margin-top: 10px;">BIÊN BẢN NGHIỆM THU KIÊM THANH LÝ HỢP ĐỒNG</div>
      <div class="doc-subtitle">Số: {info['day']}{info['month']}{info['year']}/BBTL-{info['cust_code']}</div>

      <div class="italic" style="margin-bottom: 6px;">
        - Căn cứ Hợp đồng vận chuyển số <strong>{info['contract_no']}</strong> ký ngày {info['day']} tháng {info['month']} năm {info['year']} giữa {escape(info['seller_name'])} và {escape(info['customer_name'])}.<br>
        - Căn cứ kết quả thực tế trong quá trình thực hiện hợp đồng đã ký.
      </div>

      <div style="margin: 8px 0;">
        Hôm nay, ngày {info['day']} tháng {info['month']} năm {info['year']}, tại văn phòng {escape(info['customer_name'])}, chúng tôi gồm có:
      </div>

      <table class="layout-table">
        <tr>
          <td colspan="2" class="bold uppercase">BÊN A: {escape(info['customer_name'])}</td>
        </tr>
        <tr>
          <td style="width: 110px;">Đại diện:</td>
          <td>{info['customer_honorific']} <span class="bold">{escape(info['customer_clean_rep'])}</span> &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Chức vụ: <span class="bold">{escape(info['customer_pos'])}</span></td>
        </tr>
        <tr>
          <td>Địa chỉ:</td>
          <td>{escape(info['customer_addr'])}</td>
        </tr>
        <tr>
          <td>STK:</td>
          <td>{escape(info['customer_bank_acc'])} tại ngân hàng {escape(info['customer_bank_name'])}</td>
        </tr>
        <tr>
          <td>MST:</td>
          <td class="bold">{escape(info['customer_tax'])}</td>
        </tr>

        <tr>
          <td colspan="2" class="bold uppercase" style="padding-top: 8px;">BÊN B: {escape(info['seller_name'])}</td>
        </tr>
        <tr>
          <td>Đại diện:</td>
          <td>{info['seller_honorific']} <span class="bold">{escape(info['seller_clean_rep'])}</span> &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Chức vụ: <span class="bold">{escape(info['seller_pos'])}</span></td>
        </tr>
        <tr>
          <td>Địa chỉ:</td>
          <td>{escape(info['seller_addr'])}</td>
        </tr>
        <tr>
          <td>STK:</td>
          <td>{escape(info['seller_bank_acc'])} tại ngân hàng {escape(info['seller_bank_name'])}</td>
        </tr>
        <tr>
          <td>MST:</td>
          <td class="bold">{escape(info['seller_tax'])}</td>
        </tr>
      </table>

      <div class="italic text-justify" style="margin: 8px 0 6px 0;">
        Hai bên đã cùng nhau kiểm tra xác nhận, đánh giá việc thực hiện các nội dung công việc thỏa thuận theo hợp đồng đã ký như sau:
      </div>

      <div class="text-justify" style="margin-bottom: 6px;">
        <strong>1. Bên B đã thực hiện các công việc như sau:</strong> {job_desc}.
      </div>

      <div class="text-justify" style="margin-bottom: 6px;">
        <strong>2. Chi phí thực hiện hợp đồng:</strong> Căn cứ đơn giá theo thỏa thuận tại hợp đồng số {info['contract_no']}, Bên A thanh toán cho Bên B chi phí thực hiện hợp đồng là: <strong style="color: #b91c1c;">{info['total_amount']:,} đồng</strong> (Bằng chữ: <em>{info['total_in_words']}</em>). <em>(Giá trên đã bao gồm thuế GTGT)</em>.
      </div>

      <div class="text-justify" style="margin-bottom: 6px;">
        - Hình thức thanh toán: Tiền mặt hoặc Chuyển khoản.<br>
        - Bên B đã bàn giao cho Bên A hoá đơn tài chính số <strong>{info['invoice_number_formatted']}</strong> để Bên A thực hiện việc thanh quyết toán chi phí thực hiện hợp đồng theo thỏa thuận.
      </div>

      <div class="text-justify" style="margin-bottom: 6px;">
        <strong>3. Hai bên thống nhất xác nhận</strong> thực hiện hoàn thành các công việc theo thỏa thuận tại Hợp đồng đã ký để phục vụ việc thanh quyết toán kinh phí thực hiện hợp đồng.
      </div>

      <div class="text-justify" style="margin-bottom: 6px;">
        Hợp đồng này sẽ hết hiệu lực sau khi Bên A hoàn thành việc thanh toán toàn bộ kinh phí cho Bên B.
      </div>

      <div class="text-justify">
        Biên bản này được lập thành 02 bản mỗi bên giữ 01 bản có giá trị pháp lý như nhau.
      </div>

      <table class="sig-table">
        <tr>
          <td class="sig-header">
            <div class="bold uppercase">ĐẠI DIỆN BÊN A</div>
            <div class="italic small">(Ký, ghi rõ họ tên, đóng dấu)</div>
          </td>
          <td class="sig-header">
            <div class="bold uppercase">ĐẠI DIỆN BÊN B</div>
            <div class="italic small">(Ký, ghi rõ họ tên, đóng dấu)</div>
          </td>
        </tr>
        <tr>
          <td class="sig-slot">
            {customer_sig_html}
          </td>
          <td class="sig-slot">
            {seller_sig_html}
          </td>
        </tr>
      </table>
    </div>
    """.replace(":,}", ":,}".replace(",", "."))


def render_hdvc_pages(info: dict) -> str:
    """Render Hợp đồng vận chuyển mẫu Đại Lộ Xanh (2 pages)."""
    seller_sig_url = info.get("seller_sig_url") or info.get("seller_signature_url") or info.get("seller_stamp_url") or ""
    customer_sig_url = info.get("customer_sig_url") or info.get("customer_signature_url") or info.get("customer_stamp_url") or ""
    seller_sig_html = _sig_slot_html(seller_sig_url)
    customer_sig_html = _sig_slot_html(customer_sig_url)

    return f"""
    <!-- PAGE 1: HỢP ĐỒNG VẬN CHUYỂN -->
    <div class="a4-page">
      <div class="national-header text-center">
        <div class="bold uppercase" style="font-size: 12pt;">CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</div>
        <div class="bold" style="font-size: 13pt;">Độc lập - Tự do - Hạnh phúc</div>
        <div style="letter-spacing: -1px; margin-top: -2px;">----------o0o----------</div>
      </div>

      <div class="doc-title" style="margin-top: 10px;">HỢP ĐỒNG VẬN CHUYỂN</div>
      <div class="doc-subtitle">Số : {info['contract_no']}</div>

      <div class="italic text-justify" style="font-size: 11.5pt; margin-bottom: 8px;">
        - Căn cứ vào luật thương mại có hiệu lực ngày 01/01/2006.<br>
        - Căn cứ vào bộ luật dân sự có hiệu lực ngày 01/01/2006.<br>
        - Căn cứ vào nhu cầu và khả năng của hai bên.
      </div>

      <div style="margin-bottom: 8px;">
        Hôm nay, ngày {info['day']} tháng {info['month']} năm {info['year']}, tại văn phòng {escape(info['seller_name'])}, chúng tôi đại diện cho các bên tham gia ký kết hợp đồng gồm:
      </div>

      <table class="layout-table">
        <tr>
          <td colspan="2" class="bold uppercase">BÊN THUÊ VẬN CHUYỂN (BÊN A): {escape(info['customer_name'])}</td>
        </tr>
        <tr>
          <td style="width: 120px;">Đại diện:</td>
          <td>{info['customer_honorific']} <span class="bold">{escape(info['customer_clean_rep'])}</span> &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Chức vụ: <span class="bold">{escape(info['customer_pos'])}</span></td>
        </tr>
        <tr>
          <td>Địa chỉ:</td>
          <td>{escape(info['customer_addr'])}</td>
        </tr>
        <tr>
          <td>Mã số thuế:</td>
          <td class="bold">{escape(info['customer_tax'])}</td>
        </tr>
        <tr>
          <td>Tài khoản:</td>
          <td>{escape(info['customer_bank_acc'])} tại ngân hàng {escape(info['customer_bank_name'])}</td>
        </tr>

        <tr>
          <td colspan="2" class="bold uppercase" style="padding-top: 8px;">BÊN VẬN CHUYỂN (BÊN B): {escape(info['seller_name'])}</td>
        </tr>
        <tr>
          <td>Đại diện:</td>
          <td>{info['seller_honorific']} <span class="bold">{escape(info['seller_clean_rep'])}</span> &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Chức vụ: <span class="bold">{escape(info['seller_pos'])}</span></td>
        </tr>
        <tr>
          <td>Địa chỉ:</td>
          <td>{escape(info['seller_addr'])}</td>
        </tr>
        <tr>
          <td>Điện thoại:</td>
          <td>{escape(info['seller_phone'])}</td>
        </tr>
        <tr>
          <td>Mã số thuế:</td>
          <td class="bold">{escape(info['seller_tax'])}</td>
        </tr>
        <tr>
          <td>Tài khoản:</td>
          <td>{escape(info['seller_bank_acc'])} tại ngân hàng {escape(info['seller_bank_name'])}</td>
        </tr>
        <tr>
          <td>Email:</td>
          <td>{escape(info['seller_email'])}</td>
        </tr>
      </table>

      <div class="italic text-justify" style="margin: 8px 0 6px 0;">
        Sau khi thương thảo, bàn bạc, hai bên thống nhất ký kết hợp đồng vận chuyển với các điều khoản sau:
      </div>

      <div class="bold uppercase" style="margin-top: 6px; font-size: 12pt;">Điều 1: NỘI DUNG CỦA HỢP ĐỒNG</div>
      <div class="text-justify" style="margin-bottom: 4px;">
        Bên A thuê vận chuyển và Bên B nhận cung cấp, điều động phương tiện để vận chuyển hàng hóa theo yêu cầu của Bên A như sau:
      </div>
      <div style="padding-left: 20px; line-height: 1.5;">
        - Tên hàng hóa: bàn cân, quả chuẩn hoặc các mặt hàng do bên thuê vận chuyển yêu cầu.<br>
        - Đơn vị tính: chuyến.<br>
        - Số lượng: Theo thông báo từng chuyến/đợt.<br>
        - Cự ly vận chuyển: Căn cứ theo thỏa thuận của từng lô hàng.
      </div>

      <div class="bold uppercase" style="margin-top: 8px; font-size: 12pt;">Điều 2: ĐƠN GIÁ, KHỐI LƯỢNG, ĐỊA ĐIỂM GIAO NHẬN HÀNG HÓA</div>
      <div class="text-justify" style="line-height: 1.5;">
        2.1. Đơn giá áp dụng theo chuyến do hai bên thỏa thuận theo từng thời điểm.<br>
        2.2. Khối lượng hàng hoá chở tối thiểu trên một chuyến theo định mức của bên vận chuyển.<br>
        2.3. Địa điểm giao nhận hàng hóa cụ thể từng chuyến sẽ thông báo theo yêu cầu của bên A.
      </div>
    </div>

    <!-- PAGE 2: HỢP ĐỒNG VẬN CHUYỂN (TIẾP THEO) -->
    <div class="a4-page">
      <div class="bold uppercase" style="margin-top: 4px; font-size: 12pt;">Điều 3: PHƯƠNG THỨC THANH TOÁN</div>
      <div class="text-justify" style="line-height: 1.5; margin-bottom: 8px;">
        3.1. Thanh toán bằng chuyển khoản, tiền mặt hoặc bù trừ công nợ.<br>
        3.2. Thanh toán theo chuyến áp theo giá đã được hai bên thỏa thuận từng thời điểm.<br>
        3.3. Thanh toán: Sau khi vận chuyển 2 bên tiến hành đối chiếu khối lượng hàng hóa đã vận chuyển và xuất hóa đơn thuế GTGT (VAT). Trong vòng 30 ngày làm việc sau khi nhận đầy đủ chứng từ, Bên A có trách nhiệm thanh toán cước cho Bên B.<br>
        3.4. Nếu Bên A chậm thanh toán, Bên B có quyền tạm dừng vận chuyển hàng.
      </div>

      <div class="bold uppercase" style="margin-top: 8px; font-size: 12pt;">Điều 4: TRÁCH NHIỆM MỖI BÊN</div>
      <div class="text-justify" style="line-height: 1.5; margin-bottom: 8px;">
        <strong>4.1. Bên A:</strong><br>
        - Thông báo lịch vận chuyển ít nhất là 06 tiếng để Bên B bố trí phương tiện.<br>
        - Cung cấp cho Bên B các hồ sơ chứng từ cần thiết cho việc vận chuyển hàng hóa.<br>
        - Chịu trách nhiệm dọn bãi cho xe Bên B vào cẩu và trả hàng nhanh chóng.<br>
        - Thanh toán cho Bên B theo đúng Điều 3 của hợp đồng.<br>
        <strong>4.2. Bên B:</strong><br>
        - Bố trí phương tiện vận chuyển đúng thời gian, không làm lỡ kế hoạch của Bên A.<br>
        - Đảm bảo người vận hành các xe, thiết bị phục vụ nhu cầu vận chuyển xuất trình đầy đủ giấy tờ hợp lệ.<br>
        - Đảm bảo chất lượng, số lượng hàng hóa khi vận chuyển.<br>
        - Xuất hóa đơn GTGT hợp pháp cho Bên A.
      </div>

      <div class="bold uppercase" style="margin-top: 8px; font-size: 12pt;">Điều 5: ĐIỀU KHOẢN CHUNG</div>
      <div class="text-justify" style="line-height: 1.5; margin-bottom: 8px;">
        - Hai bên cam kết thực hiện đúng những điều khoản ghi trong hợp đồng. Nếu bên nào vi phạm sẽ phải bồi thường thiệt hại cho bên kia theo quy định pháp luật.<br>
        - Mọi tranh chấp phát sinh trong quá trình thực hiện hợp đồng sẽ được hai bên giải quyết thông qua thương lượng hòa giải. Nếu không giải quyết được sẽ đưa ra Tòa án có thẩm quyền tại Tỉnh Phú Thọ để giải quyết.
      </div>

      <div class="bold uppercase" style="margin-top: 8px; font-size: 12pt;">Điều 6: HIỆU LỰC CỦA HỢP ĐỒNG</div>
      <div class="text-justify" style="line-height: 1.5; margin-bottom: 12px;">
        - Hợp đồng này có hiệu lực kể từ ngày ký đến hết ngày 31/12/2026.<br>
        - Hợp đồng này được lập thành 02 bản, mỗi bên giữ 01 bản có giá trị pháp lý như nhau.
      </div>

      <table class="sig-table" style="margin-top: 15px;">
        <tr>
          <td class="sig-header">
            <div class="bold uppercase">ĐẠI DIỆN BÊN A</div>
            <div class="italic small">(Ký, đóng dấu)</div>
          </td>
          <td class="sig-header">
            <div class="bold uppercase">ĐẠI DIỆN BÊN B</div>
            <div class="italic small">(Ký, đóng dấu)</div>
          </td>
        </tr>
        <tr>
          <td class="sig-slot">
            {customer_sig_html}
          </td>
          <td class="sig-slot">
            {seller_sig_html}
          </td>
        </tr>
      </table>
    </div>
    """


def render_dlx_document_html(
    invoice: dict,
    lines: list[dict],
    customer: dict,
    seller: dict | None = None,
    kind: str = "hdvc",
    preview: bool = True,
    invoice_guid: str | None = None,
) -> str:
    """Render documents as full HTML with toolbar and standard A4 pages."""
    info = _extract_trip_data(invoice, lines, customer, seller)

    if kind == "bbvc":
        body_content = render_bbvc_page(info)
        doc_title = f"Biên bản vận chuyển - {info['bbvc_no']}"
    elif kind == "dntt":
        body_content = render_dntt_page(info)
        doc_title = f"Đề nghị thanh toán - {info['dntt_no']}"
    elif kind == "bbtl":
        body_content = render_bbtl_page(info)
        doc_title = f"Thanh lý hợp đồng - {info['contract_no']}"
    elif kind == "all":
        body_content = "\n".join([
            render_hdvc_pages(info),
            render_bbvc_page(info),
            render_bbtl_page(info),
            render_dntt_page(info),
        ])
        doc_title = f"Trọn bộ hồ sơ chuyến hàng - HĐ {info['invoice_number_formatted']}"
    else:  # "hdvc"
        body_content = render_hdvc_pages(info)
        doc_title = f"Hợp đồng vận chuyển - {info['contract_no']}"

    toolbar_html = _render_toolbar(info, kind, invoice_guid) if preview else ""

    return f"""<!DOCTYPE html>
<html lang="vi">
<head>
  <meta charset="utf-8">
  <title>{escape(doc_title)}</title>
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <style>
    {_COMMON_CSS}
  </style>
</head>
<body>
  {toolbar_html}
  <div class="doc-body-container" style="padding: 10px 0;">
    {body_content}
  </div>
</body>
</html>"""


def _set_run_font(run, name: str = "Times New Roman", size_pt: float = 11.0, bold: bool = False, italic: bool = False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic


def _add_p(doc, text: str = "", align=WD_ALIGN_PARAGRAPH.LEFT, space_before: float = 0, space_after: float = 3, line_spacing: float = 1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        _set_run_font(r)
    return p


def render_dlx_document_docx(
    invoice: dict,
    lines: list[dict],
    customer: dict,
    seller: dict | None = None,
    kind: str = "hdvc",
) -> tuple[bytes, str]:
    """Generate Word (.docx) document for the requested trip document(s)."""
    import docx
    from docx.shared import Pt, Mm, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

    info = _extract_trip_data(invoice, lines, customer, seller)
    doc = docx.Document()

    for sec in doc.sections:
        sec.top_margin = Mm(18)
        sec.bottom_margin = Mm(18)
        sec.left_margin = Mm(20)
        sec.right_margin = Mm(15)

    def write_national_header():
        p1 = _add_p(doc, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", align=WD_ALIGN_PARAGRAPH.CENTER)
        p1.runs[0].bold = True
        p1.runs[0].font.size = Pt(11)
        p2 = _add_p(doc, "Độc lập - Tự do - Hạnh phúc", align=WD_ALIGN_PARAGRAPH.CENTER)
        p2.runs[0].bold = True
        p2.runs[0].font.size = Pt(12)
        p3 = _add_p(doc, "----------o0o----------", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
        p3.runs[0].font.size = Pt(10)

    def write_signatures(
        left_title: str,
        left_url: str,
        right_title: str,
        right_url: str,
        left_sub: str = "(Ký, đóng dấu)",
        right_sub: str = "(Ký, đóng dấu)",
    ):
        sig_tbl = doc.add_table(rows=2, cols=2)
        sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        sig_tbl.autofit = False
        for row in sig_tbl.rows:
            row.cells[0].width = Inches(3.5)
            row.cells[1].width = Inches(3.5)

        c00, c01 = sig_tbl.rows[0].cells
        p00 = c00.paragraphs[0]
        p00.paragraph_format.space_before = Pt(0)
        p00.paragraph_format.space_after = Pt(2)
        if left_title.upper().startswith("NƠI NHẬN"):
            p00.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r1 = p00.add_run("Nơi nhận:\n")
            _set_run_font(r1, bold=True, italic=True, size_pt=10.5)
            r2 = p00.add_run("- Như trên\n- Lưu VP Công ty")
            _set_run_font(r2, size_pt=10)
        else:
            p00.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p00.add_run(left_title + "\n")
            _set_run_font(r1, bold=True, size_pt=11)
            r2 = p00.add_run(left_sub)
            _set_run_font(r2, italic=True, size_pt=9.5)

        p01 = c01.paragraphs[0]
        p01.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p01.paragraph_format.space_before = Pt(0)
        p01.paragraph_format.space_after = Pt(2)
        r4 = p01.add_run(right_title + "\n")
        _set_run_font(r4, bold=True, size_pt=11)
        r5 = p01.add_run(right_sub)
        _set_run_font(r5, italic=True, size_pt=9.5)

        # Hàng 2: Ảnh chữ ký (chỉ điền ảnh chữ ký, cân đối độ cao cả 2 bên)
        sig_tbl.rows[1].height = Pt(130)
        c10, c11 = sig_tbl.rows[1].cells
        c10.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        c11.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        def _set_docx_cell_sig(cell, url: str):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            valid_path = None
            if url:
                clean_path = url.split("?")[0].lstrip("/")
                test_p = Path(cfg.BASE_DIR) / clean_path
                if test_p.exists() and test_p.is_file():
                    valid_path = test_p
            if valid_path:
                try:
                    p.add_run().add_picture(
                        str(valid_path), width=Mm(_signature_print_width_mm(url))
                    )
                    return
                except Exception:
                    pass
            p.paragraph_format.line_spacing = Pt(20)
            p.add_run("\n\n\n")

        if not left_title.upper().startswith("NƠI NHẬN"):
            _set_docx_cell_sig(c10, left_url)
        _set_docx_cell_sig(c11, right_url)

    sections_to_render = []
    if kind == "bbvc":
        sections_to_render = ["bbvc"]
        filename_prefix = f"BBVC_{info['invoice_number_formatted']}"
    elif kind == "dntt":
        sections_to_render = ["dntt"]
        filename_prefix = f"DNTT_{info['invoice_number_formatted']}"
    elif kind == "bbtl":
        sections_to_render = ["bbtl"]
        filename_prefix = f"BBTL_{info['invoice_number_formatted']}"
    elif kind == "all":
        sections_to_render = ["hdvc", "bbvc", "bbtl", "dntt"]
        filename_prefix = f"Ho_so_chuyen_hang_HD{info['invoice_number_formatted']}"
    else:  # "hdvc"
        sections_to_render = ["hdvc"]
        filename_prefix = f"HDVC_{info['invoice_number_formatted']}"

    for s_idx, sec_name in enumerate(sections_to_render):
        if s_idx > 0:
            doc.add_page_break()

        if sec_name == "hdvc":
            write_national_header()
            p_t = _add_p(doc, "HỢP ĐỒNG VẬN CHUYỂN", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4)
            _set_run_font(p_t.runs[0], bold=True, size_pt=14)
            p_no = _add_p(doc, f"Số : {info['contract_no']}", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
            _set_run_font(p_no.runs[0], italic=True, size_pt=11)

            _add_p(doc, f"Hôm nay, ngày {info['day']} tháng {info['month']} năm {info['year']}, tại văn phòng {info['seller_name']}, chúng tôi gồm có:", space_after=4)
            p_a = _add_p(doc, f"BÊN THUÊ VẬN CHUYỂN (BÊN A): {info['customer_name']}", space_before=4)
            p_a.runs[0].bold = True
            _add_p(doc, f"- Đại diện: {info['customer_honorific']} {info['customer_clean_rep']}      Chức vụ: {info['customer_pos']}")
            _add_p(doc, f"- Địa chỉ: {info['customer_addr']}")
            _add_p(doc, f"- MST: {info['customer_tax']}      Tài khoản: {info['customer_bank_acc']} tại {info['customer_bank_name']}")

            p_b = _add_p(doc, f"BÊN VẬN CHUYỂN (BÊN B): {info['seller_name']}", space_before=6)
            p_b.runs[0].bold = True
            _add_p(doc, f"- Đại diện: {info['seller_honorific']} {info['seller_clean_rep']}      Chức vụ: {info['seller_pos']}")
            _add_p(doc, f"- Địa chỉ: {info['seller_addr']}")
            _add_p(doc, f"- MST: {info['seller_tax']}      Tài khoản: {info['seller_bank_acc']} tại {info['seller_bank_name']}")

            _add_p(doc, "Điều 1: NỘI DUNG CỦA HỢP ĐỒNG", space_before=8).runs[0].bold = True
            _add_p(doc, "Bên A thuê vận chuyển và Bên B nhận cung cấp, điều động phương tiện vận chuyển hàng hóa theo từng chuyến thông báo.")
            _add_p(doc, "Điều 2: ĐƠN GIÁ VÀ ĐỊA ĐIỂM GIAO NHẬN", space_before=6).runs[0].bold = True
            _add_p(doc, "Đơn giá áp dụng theo chuyến do hai bên thỏa thuận từng thời điểm. Địa điểm giao nhận theo yêu cầu cụ thể của Bên A.")
            _add_p(doc, "Điều 3: PHƯƠNG THỨC THANH TOÁN", space_before=6).runs[0].bold = True
            _add_p(doc, "Thanh toán bằng chuyển khoản hoặc tiền mặt trong vòng 30 ngày làm việc sau khi nhận đầy đủ chứng từ hóa đơn.")
            _add_p(doc, "Điều 4: TRÁCH NHIỆM MỖI BÊN", space_before=6).runs[0].bold = True
            _add_p(doc, "Bên A thông báo kế hoạch vận chuyển trước 06 tiếng và thanh toán đúng hạn. Bên B bố trí xe đúng thời gian và đảm bảo an toàn hàng hóa.")
            _add_p(doc, "Điều 5: HIỆU LỰC HỢP ĐỒNG", space_before=6).runs[0].bold = True
            _add_p(doc, f"Hợp đồng có hiệu lực từ ngày ký đến hết ngày 31/12/2026, lập thành 02 bản có giá trị pháp lý như nhau.", space_after=14)

            customer_sig = info.get("customer_sig_url") or info.get("customer_stamp_url") or ""
            seller_sig = info.get("seller_sig_url") or info.get("seller_stamp_url") or ""
            write_signatures("ĐẠI DIỆN BÊN A", customer_sig, "ĐẠI DIỆN BÊN B", seller_sig)

        elif sec_name == "bbvc":
            write_national_header()
            p_t = _add_p(doc, "BIÊN BẢN VẬN CHUYỂN", align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_run_font(p_t.runs[0], bold=True, size_pt=14)
            p_no = _add_p(doc, f"Số : {info['bbvc_no']}", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
            _set_run_font(p_no.runs[0], italic=True, size_pt=11)

            _add_p(doc, f"Hôm nay, ngày {info['day']} tháng {info['month']} năm {info['year']}. Tại văn phòng quý khách hàng, chúng tôi gồm có:")
            p_a = _add_p(doc, f"BÊN VẬN CHUYỂN (BÊN A): {info['seller_name']}", space_before=4)
            p_a.runs[0].bold = True
            _add_p(doc, f"MST: {info['seller_tax']}      Địa chỉ: {info['seller_addr']}")
            _add_p(doc, f"Đại diện: {info['seller_honorific']} {info['seller_clean_rep']} - Giám đốc      BKS: {info['primary_plate']}")

            p_b = _add_p(doc, f"BÊN THUÊ VẬN CHUYỂN (BÊN B): {info['customer_name']}", space_before=6)
            p_b.runs[0].bold = True
            _add_p(doc, f"MST: {info['customer_tax']}      Địa chỉ: {info['customer_addr']}")
            _add_p(doc, f"Đại diện: {info['customer_honorific']} {info['customer_clean_rep']} - {info['customer_pos']}")

            _add_p(doc, "I. NỘI DUNG VẬN CHUYỂN:", space_before=8).runs[0].bold = True

            # Items table
            tbl = doc.add_table(rows=1, cols=6)
            tbl.style = 'Table Grid'
            hdr = tbl.rows[0].cells
            col_names = ["STT", "Tên hàng hóa, dịch vụ", "BKS", "ĐVT", "Đơn giá", "Thành tiền"]
            for i, name in enumerate(col_names):
                hdr[i].text = name
                hdr[i].paragraphs[0].runs[0].bold = True
                hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

            for it in info["items"]:
                r_cells = tbl.add_row().cells
                r_cells[0].text = str(it["index"])
                r_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_cells[1].text = it["description"]
                r_cells[2].text = it["plate_number"] or info["primary_plate"]
                r_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_cells[3].text = it["unit"]
                r_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_cells[4].text = f"{it['price']:,}".replace(",", ".")
                r_cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r_cells[5].text = f"{it['amount']:,}".replace(",", ".")
                r_cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

            # Total rows
            r_c = tbl.add_row().cells
            r_c[0].merge(r_c[4])
            r_c[0].text = "Cộng:"
            r_c[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_c[0].paragraphs[0].runs[0].bold = True
            r_c[5].text = f"{info['subtotal']:,}".replace(",", ".")
            r_c[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

            r_v = tbl.add_row().cells
            r_v[0].merge(r_v[4])
            r_v[0].text = f"Thuế GTGT {info['vat_percent']}%:"
            r_v[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_v[0].paragraphs[0].runs[0].bold = True
            r_v[5].text = f"{info['vat_amount']:,}".replace(",", ".")
            r_v[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

            r_t = tbl.add_row().cells
            r_t[0].merge(r_t[4])
            r_t[0].text = "Tổng thanh toán:"
            r_t[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_t[0].paragraphs[0].runs[0].bold = True
            r_t[5].text = f"{info['total_amount']:,}".replace(",", ".")
            r_t[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_t[5].paragraphs[0].runs[0].bold = True

            p_words = _add_p(doc, f"Số tiền bằng chữ: {info['total_in_words']}", space_before=6)
            p_words.runs[0].italic = True

            _add_p(doc, "II. KẾT LUẬN:", space_before=6).runs[0].bold = True
            _add_p(doc, "Hai bên cùng nhau xác nhận đã giao nhận đúng, đủ số lượng hàng vận chuyển theo đúng lộ trình đã thống nhất. Biên bản được lập thành 02 bản có giá trị như nhau.", space_after=14)

            customer_sig = info.get("customer_sig_url") or info.get("customer_stamp_url") or ""
            seller_sig = info.get("seller_sig_url") or info.get("seller_stamp_url") or ""
            write_signatures(
                "ĐẠI DIỆN BÊN A",
                seller_sig,
                "ĐẠI DIỆN BÊN B",
                customer_sig,
                left_sub="(Ký, ghi rõ họ tên, đóng dấu)",
                right_sub="(Ký, ghi rõ họ tên)",
            )

        elif sec_name == "bbtl":
            write_national_header()
            p_t = _add_p(doc, "BIÊN BẢN NGHIỆM THU KIÊM THANH LÝ HỢP ĐỒNG", align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_run_font(p_t.runs[0], bold=True, size_pt=13)
            p_no = _add_p(doc, f"Số: {info['day']}{info['month']}{info['year']}/BBTL-{info['cust_code']}", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
            _set_run_font(p_no.runs[0], italic=True, size_pt=10.5)

            p_cc = _add_p(doc, f"- Căn cứ Hợp đồng vận chuyển số {info['contract_no']} ký ngày {info['day']} tháng {info['month']} năm {info['year']};\n- Căn cứ kết quả thực tế giao nhận hàng hóa.")
            p_cc.runs[0].italic = True

            _add_p(doc, f"Hôm nay, ngày {info['day']} tháng {info['month']} năm {info['year']}, tại văn phòng {info['customer_name']}, chúng tôi gồm có:", space_after=4)
            p_a = _add_p(doc, f"BÊN A: {info['customer_name']}", space_before=4)
            p_a.runs[0].bold = True
            _add_p(doc, f"Đại diện: {info['customer_honorific']} {info['customer_clean_rep']} - {info['customer_pos']}      MST: {info['customer_tax']}")
            _add_p(doc, f"Địa chỉ: {info['customer_addr']}")
            _add_p(doc, f"STK: {info['customer_bank_acc']} tại {info['customer_bank_name']}")

            p_b = _add_p(doc, f"BÊN B: {info['seller_name']}", space_before=6)
            p_b.runs[0].bold = True
            _add_p(doc, f"Đại diện: {info['seller_honorific']} {info['seller_clean_rep']} - Giám đốc      MST: {info['seller_tax']}")
            _add_p(doc, f"Địa chỉ: {info['seller_addr']}")
            _add_p(doc, f"STK: {info['seller_bank_acc']} tại {info['seller_bank_name']}")

            job_desc = "; ".join(it["description"] for it in info["items"])
            _add_p(doc, "Hai bên cùng nhau kiểm tra, xác nhận và đánh giá kết quả thực hiện như sau:", space_before=8).runs[0].italic = True
            _add_p(doc, f"1. Bên B đã thực hiện công việc: {job_desc}.")
            _add_p(doc, f"2. Chi phí thực hiện hợp đồng: Căn cứ thỏa thuận tại hợp đồng số {info['contract_no']}, Bên A thanh toán cho Bên B số tiền: {info['total_amount']:,} đồng (Bằng chữ: {info['total_in_words']}) (Đã bao gồm thuế GTGT).".replace(",", "."))
            _add_p(doc, f"Bên B đã bàn giao cho Bên A hóa đơn tài chính số {info['invoice_number_formatted']}.")
            _add_p(doc, "3. Hai bên thống nhất nghiệm thu và thanh lý toàn bộ hợp đồng. Biên bản lập thành 02 bản có giá trị pháp lý như nhau.", space_after=14)

            customer_sig = info.get("customer_sig_url") or info.get("customer_stamp_url") or ""
            seller_sig = info.get("seller_sig_url") or info.get("seller_stamp_url") or ""
            write_signatures(
                "ĐẠI DIỆN BÊN A",
                customer_sig,
                "ĐẠI DIỆN BÊN B",
                seller_sig,
                left_sub="(Ký, ghi rõ họ tên, đóng dấu)",
                right_sub="(Ký, ghi rõ họ tên, đóng dấu)",
            )

        elif sec_name == "dntt":
            tbl_hdr = doc.add_table(rows=1, cols=2)
            c0 = tbl_hdr.cell(0, 0)
            c1 = tbl_hdr.cell(0, 1)
            p0 = c0.paragraphs[0]
            p0.add_run(f"{info['seller_name']}\n").bold = True
            p0.add_run(f"Số: {info['dntt_no']}\n")
            p0.add_run("V/v: Thanh toán hóa đơn")

            p1 = c1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n").bold = True
            p1.add_run("Độc lập – Tự do - Hạnh phúc\n").bold = True
            p1.add_run("--------------------\n")
            p1.add_run(f"Phú Thọ, ngày {info['day']} tháng {info['month']} năm {info['year']}").italic = True

            p_t = _add_p(doc, "ĐỀ NGHỊ THANH TOÁN", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16, space_after=10)
            _set_run_font(p_t.runs[0], bold=True, size_pt=14)

            p_kg = _add_p(doc, f"Kính gửi: {info['customer_name']}", space_after=8)
            p_kg.runs[0].bold = True

            first_item = info["items"][0] if info["items"] else {}
            cargo_name = first_item.get("description") or "hàng hóa"
            _add_p(doc, "Trước tiên xin chân thành cảm ơn quý công ty đã tin tưởng và lựa chọn Công ty chúng tôi.")
            _add_p(doc, f"Chúng tôi đã vận chuyển {cargo_name} từ {info['primary_origin']} đi {info['primary_dest']} BKS {info['primary_plate']}.")
            _add_p(doc, f"Tổng cước vận chuyển là: {info['total_amount']:,} đồng (đã bao gồm thuế GTGT) và đã phát hành Hóa đơn số: {info['invoice_number_formatted']} ngày {info['day']}/{info['month']}/{info['year']}.".replace(",", "."))
            _add_p(doc, f"Bằng công văn này, kính đề nghị Quý Công ty thanh toán cho chúng tôi số tiền là: {info['total_amount']:,} đồng.".replace(",", "."))
            p_w = _add_p(doc, f"(Số tiền bằng chữ: {info['total_in_words']}).", space_after=8)
            p_w.runs[0].italic = True

            _add_p(doc, "Thông Tin Thụ Hưởng:", space_before=6).runs[0].bold = True
            _add_p(doc, f"- Đơn vị thụ hưởng: {info['seller_name']}.")
            _add_p(doc, f"- Số tài khoản: {info['seller_bank_acc']}")
            _add_p(doc, f"- Tại ngân hàng: {info['seller_bank_name']}.")
            _add_p(doc, "Xin trân trọng cảm ơn!", space_before=6, space_after=14)

            seller_sig = info.get("seller_sig_url") or info.get("seller_stamp_url") or ""
            write_signatures(
                "NƠI NHẬN:\n- Như trên\n- Lưu VP Công ty",
                "",
                info["seller_name"],
                seller_sig,
                right_sub="(Ký, ghi rõ họ tên, đóng dấu)",
            )

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), filename_prefix

