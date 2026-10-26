"""Generate and export HĐNT (Hợp đồng nguyên tắc) and BBBG (Biên bản bàn giao / giao nhận hàng hoá).

Based on standard template: HĐNT + BBBG VẬN CHUYỂN HTC - TIẾN DŨNG.doc
"""

from __future__ import annotations

import copy
from datetime import date, datetime
import io
import logging
import os
from pathlib import Path
import re
import tempfile
from typing import Any

import docx
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips, Mm, Inches

from html import escape
from frontend.modules.gender_utils import split_rep_honorific, get_honorific
from frontend.modules.transport.contract import (
    derive_company_code,
    _to_data_uri_if_local,
    _contract_bank_accounts,
    _remove_accents,
)
from frontend.modules.invoices.companies import get_company_bank_accounts

logger = logging.getLogger(__name__)

TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "assets" / "HDNT_BBBG_template.docx"

# Keep the source table's total width while adding a plate column before ĐVT.
ITEM_COLUMN_WIDTHS = (500, 3900, 1300, 900, 900, 1350, 1470)  # twips


def _add_vehicle_plate_column(table: docx.table.Table, item_count: int) -> None:
    """Expand the six-column template without splitting its merged total rows."""
    grid_column = OxmlElement("w:gridCol")
    grid_column.set(qn("w:w"), str(ITEM_COLUMN_WIDTHS[2]))
    table._tbl.tblGrid.insert(2, grid_column)
    for row_index, row in enumerate(table.rows):
        if row_index <= item_count:
            plate_cell = copy.deepcopy(row._tr.tc_lst[2])
            row._tr.insert(2, plate_cell)
        else:
            merged = row._tr.tc_lst[1].get_or_add_tcPr().find(qn("w:gridSpan"))
            merged.set(qn("w:val"), str(int(merged.get(qn("w:val"))) + 1))
    plate_header = table.rows[0].cells[2]
    plate_header.text = "BIỂN KIỂM SOÁT (BKS)"
    plate_header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    plate_header.paragraphs[0].runs[0].bold = True


def _format_items_table(table: docx.table.Table, item_count: int) -> None:
    table.autofit = False
    for column, width in zip(table.columns, ITEM_COLUMN_WIDTHS):
        column.width = Twips(width)

    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.insert_element_before(margins, "w:tblLook")
    for side in ("left", "right"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), "50")
        node.set(qn("w:type"), "dxa")

    for row_index, row in enumerate(table.rows):
        if row_index <= item_count:
            cells = row.cells
            for index, width in enumerate(ITEM_COLUMN_WIDTHS):
                cells[index].width = Twips(width)
            unique_cells = cells
        else:
            # Summary rows merge the description through unit-price columns.
            cells = row.cells
            cells[0].width = Twips(ITEM_COLUMN_WIDTHS[0])
            cells[1].width = Twips(sum(ITEM_COLUMN_WIDTHS[1:-1]))
            cells[-1].width = Twips(ITEM_COLUMN_WIDTHS[-1])
            unique_cells = (cells[0], cells[1], cells[-1])

        for index, cell in enumerate(unique_cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            no_wrap = (0 < row_index <= item_count and index in (2, 3, 5, 6)) or (
                row_index > item_count and index == 2)
            if no_wrap:
                tc_pr = cell._tc.get_or_add_tcPr()
                if tc_pr.find(qn("w:noWrap")) is None:
                    tc_pr.append(OxmlElement("w:noWrap"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8.5 if row_index == 0 else 9)


def number_to_vietnamese_words(n: int | float) -> str:
    """Convert an integer number to standard Vietnamese currency words (e.g. 8.640.000 -> Tám triệu sáu trăm bốn mươi nghìn đồng chẵn./.)."""
    n_int = abs(int(round(float(n or 0))))
    if n_int == 0:
        return "Không đồng chẵn./."

    digits = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]
    units = ["", "nghìn", "triệu", "tỷ", "nghìn tỷ", "triệu tỷ"]

    def read_block_3(num: int, show_zero_hundred: bool = False) -> str:
        h = num // 100
        t = (num % 100) // 10
        u = num % 10
        words: list[str] = []
        if h > 0:
            words.append(digits[h] + " trăm")
        elif show_zero_hundred:
            words.append("không trăm")

        if t == 0:
            if u > 0:
                if h > 0 or show_zero_hundred:
                    words.append("lẻ")
                words.append(digits[u])
        elif t == 1:
            words.append("mười")
            if u == 1:
                words.append("một")
            elif u == 5:
                words.append("lăm")
            elif u > 0:
                words.append(digits[u])
        else:
            words.append(digits[t] + " mươi")
            if u == 1:
                words.append("mốt")
            elif u == 4:
                words.append("tư")
            elif u == 5:
                words.append("lăm")
            elif u > 0:
                words.append(digits[u])
        return " ".join(words)

    blocks: list[int] = []
    temp = n_int
    while temp > 0:
        blocks.append(temp % 1000)
        temp //= 1000

    result: list[str] = []
    for i in range(len(blocks) - 1, -1, -1):
        b = blocks[i]
        if b > 0:
            b_str = read_block_3(b, show_zero_hundred=(i < len(blocks) - 1))
            unit = units[i]
            if unit:
                result.append(b_str + " " + unit)
            else:
                result.append(b_str)

    res_str = " ".join(result).strip()
    if not res_str:
        return "Không đồng chẵn./."
    res_str = res_str[0].upper() + res_str[1:] + " đồng chẵn./."
    return res_str


def convert_docx_bytes_to_doc(docx_bytes: bytes) -> bytes:
    """Convert .docx bytes to native Microsoft Word 97-2003 (.doc) binary using Word COM."""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    in_path = ""
    out_path = ""
    try:
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f_in:
            f_in.write(docx_bytes)
            in_path = f_in.name
        out_path = in_path[:-5] + ".doc"

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(os.path.abspath(in_path))
        doc.SaveAs2(os.path.abspath(out_path), FileFormat=0)  # 0 = wdFormatDocument (doc)
        doc.Close()
        word.Quit()

        with open(out_path, "rb") as f_out:
            return f_out.read()
    finally:
        if in_path and os.path.exists(in_path):
            try:
                os.remove(in_path)
            except Exception:
                pass
        if out_path and os.path.exists(out_path):
            try:
                os.remove(out_path)
            except Exception:
                pass
        pythoncom.CoUninitialize()


def _populate_template(trip: dict, customer: dict, seller: dict | None = None) -> tuple[docx.Document, str]:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Template not found at {TEMPLATE_PATH}")

    doc = docx.Document(str(TEMPLATE_PATH))

    # Determine contract date
    raw_date = trip.get("invoice_date") or trip.get("departure_time")
    if isinstance(raw_date, datetime):
        c_date = raw_date.date()
    elif isinstance(raw_date, date):
        c_date = raw_date
    elif isinstance(raw_date, str) and raw_date:
        try:
            c_date = date.fromisoformat(raw_date[:10])
        except Exception:
            c_date = date.today()
    else:
        c_date = date.today()

    day_str = f"{c_date.day:02d}"
    month_str = f"{c_date.month:02d}"
    year_str = str(c_date.year)
    short_year = str(c_date.year)[-2:]

    customer_name = (customer.get("name") or "").strip()
    cust_code = derive_company_code(customer_name)

    if not seller:
        raise ValueError("Chưa chọn hoặc không tìm thấy thông tin công ty bán/vận chuyển (Bên A).")

    seller_name = (seller.get("name") or "").strip()
    if not seller_name:
        raise ValueError("Tên công ty bán/vận chuyển (Bên A) không được để trống trong cơ sở dữ liệu.")

    seller_code = derive_company_code(seller_name)
    contract_no = trip.get("contract_no") or f"{day_str}{month_str}{short_year}/HĐNT/{cust_code}-{seller_code}"

    seller_rep = (seller.get("director_name") or "").strip()
    seller_pos = (seller.get("director_title") or "Giám đốc").strip()
    seller_addr = (seller.get("address") or "").strip()
    seller_phone = (seller.get("phone") or "").strip()
    seller_tax = (seller.get("tax_code") or "").strip()
    seller_bank_acc = (seller.get("bank_account") or "").strip()
    seller_bank_name = (seller.get("bank_name") or "").strip()
    seller_sig_url = (seller.get("signature_url") or "").strip()
    seller_stamp_url = (seller.get("stamp_url") or "").strip()

    customer_rep = (customer.get("director_name") or "").strip()
    customer_pos = (customer.get("director_title") or "Giám đốc").strip()
    customer_addr = (customer.get("address") or "").strip()
    customer_phone = (customer.get("phone") or "").strip()
    customer_tax = (customer.get("tax_code") or "").strip()
    customer_bank_acc = (customer.get("bank_account") or "").strip()
    customer_bank_name = (customer.get("bank_name") or "").strip()
    customer_sig_url = (customer.get("signature_url") or "").strip()
    customer_stamp_url = (customer.get("stamp_url") or "").strip()

    cargo_desc = (trip.get("cargo_description") or "").strip()

    # --- Part 1 (HĐNT) replacement ---
    p = doc.paragraphs[6]
    p.text = f"Số: {contract_no}"
    if p.runs:
        p.runs[0].bold = True

    p = doc.paragraphs[12]
    p.text = f"Hôm nay, ngày {day_str} tháng {month_str} năm {year_str}. Chúng tôi gồm có:"

    # Bên A (Vận chuyển)
    p = doc.paragraphs[13]
    p.text = ""
    r1 = p.add_run("Bên A: (Bên vận chuyển) ")
    r1.bold = True
    r2 = p.add_run(seller_name)
    r2.bold = True

    p = doc.paragraphs[15]
    p.text = f"Địa chỉ            : {seller_addr}"

    p = doc.paragraphs[16]
    p.text = f"Điện thoại       : {seller_phone}"

    p = doc.paragraphs[17]
    p.text = f"Mã số thuế\t: {seller_tax}"

    def bank_lines(company, bank_acc="", bank_name=""):
        accounts = (company or {}).get("bank_accounts") or (
            get_company_bank_accounts(company["id"]) if company and company.get("id") else []
        )
        if not accounts and (bank_acc or bank_name):
            accounts = [{"account_number": bank_acc, "bank_name": bank_name, "branch": ""}]
        if not accounts:
            return "Tài khoản số   \t:"
        lines = []
        for idx, item in enumerate(accounts):
            bank_label = " - ".join(part for part in (item.get("bank_name"), item.get("branch")) if part)
            num = item.get("account_number") or ""
            bank_suffix = f" tại ngân hàng {bank_label}" if bank_label else ""
            if idx == 0:
                lines.append(f"Tài khoản số   \t: {num}{bank_suffix}")
            else:
                lines.append(f"               \t  {num}{bank_suffix}")
        return "\n".join(lines)

    p = doc.paragraphs[18]
    p.text = bank_lines(seller, seller_bank_acc, seller_bank_name)

    def set_rep_paragraph(p, rep, pos, gender=None):
        p.text = ""
        prefix, clean_rep = split_rep_honorific(rep, gender)
        honorific = get_honorific(rep, gender)
        p.add_run(f"Đại diện \t:{honorific} ")
        if clean_rep:
            r = p.add_run(clean_rep)
            r.bold = True
        p.add_run(" \t\tChức vụ: ")
        r_pos = p.add_run(pos or "Giám đốc")
        r_pos.bold = True

    set_rep_paragraph(doc.paragraphs[19], seller_rep, seller_pos, seller.get("director_gender") or seller.get("gender"))

    # Bên B (Thuê vận chuyển)
    p = doc.paragraphs[21]
    p.text = ""
    r1 = p.add_run("BÊN B: (Bên thuê vận chuyển) ")
    r1.bold = True
    r2 = p.add_run(customer_name)
    r2.bold = True

    p = doc.paragraphs[22]
    p.text = f"Địa chỉ     : {customer_addr}"

    p = doc.paragraphs[23]
    p.text = f"Điện thoại\t: {customer_phone}"

    p = doc.paragraphs[24]
    p.text = f"Mã số thuế\t: {customer_tax}"

    p = doc.paragraphs[25]
    p.text = bank_lines(customer, customer_bank_acc, customer_bank_name)

    p = doc.paragraphs[26]
    p.text = ""

    set_rep_paragraph(doc.paragraphs[27], customer_rep, customer_pos, customer.get("director_gender") or customer.get("gender"))

    p = doc.paragraphs[32]
    if cargo_desc:
        p.text = f"+ Tên hàng hóa  \t: {cargo_desc} hoặc các mặt hàng do bên thuê vận chuyển yêu cầu."
    else:
        p.text = "+ Tên hàng hóa  \t: Các mặt hàng do bên thuê vận chuyển yêu cầu."

    p = doc.paragraphs[71]
    p.text = f"- Hợp đồng này có hiệu lực đến hết ngày 31/12/{year_str}"

    # --- Part 2 (BBBG) replacement ---
    p = doc.paragraphs[100]
    p.text = f"     - Căn cứ vào Hợp đồng nguyên tắc số: {contract_no} về việc thỏa thuận vận chuyển hàng hóa được ký ngày {day_str}/{month_str}/{year_str};"

    p = doc.paragraphs[102]
    p.text = f"Hôm nay, Ngày {day_str} tháng {month_str} năm {year_str}, Chúng tôi gồm có:"

    p = doc.paragraphs[104]
    p.text = ""
    r1 = p.add_run("BÊN A: (Bên vận chuyển) ")
    r1.bold = True
    r2 = p.add_run(seller_name)
    r2.bold = True

    p = doc.paragraphs[105]
    p.text = f"Địa chỉ            : {seller_addr}"

    p = doc.paragraphs[106]
    p.text = f"Mã số thuế\t: {seller_tax}"

    p = doc.paragraphs[107]
    set_rep_paragraph(p, seller_rep, seller_pos, seller.get("director_gender") or seller.get("gender"))

    p = doc.paragraphs[109]
    p.text = ""
    r1 = p.add_run("BÊN B: (Bên thuê vận chuyển) ")
    r1.bold = True
    r2 = p.add_run(customer_name)
    r2.bold = True

    p = doc.paragraphs[110]
    p.text = f"Địa chỉ     : {customer_addr}"

    p = doc.paragraphs[111]
    p.text = f"Mã số thuế\t: {customer_tax}"

    p = doc.paragraphs[112]
    set_rep_paragraph(p, customer_rep, customer_pos, customer.get("director_gender") or customer.get("gender"))

    # --- Table 0 (Items Table) ---
    t0 = doc.tables[0]
    segments = trip.get("route_segments") or []
    items: list[dict[str, Any]] = []
    if segments:
        for seg in segments:
            desc = seg.get("description") or ""
            if not desc:
                cargo_part = f" {cargo_desc}" if cargo_desc else ""
                service = trip.get("service_name") or "Cước vận chuyển"
                origin = seg.get("origin", "")
                dest = seg.get("destination", "")
                route_part = f" từ {origin} đi {dest}" if (origin or dest) else ""
                desc = f"{service}{cargo_part}{route_part}".strip()
            amount = int(round(float(seg.get("amount") or 0)))
            items.append({
                "description": desc,
                "plate_number": str(seg.get("plate_number") or trip.get("plate_number") or "").strip(),
                "unit": "Chuyến",
                "quantity": 1,
                "price": amount,
                "amount": amount,
            })
    else:
        desc = trip.get("trip_note") or ""
        if not desc:
            cargo_part = f" {cargo_desc}" if cargo_desc else ""
            service = trip.get("service_name") or "Cước vận chuyển"
            origin = trip.get("origin", "")
            dest = trip.get("destination", "")
            route_part = f" từ {origin} đi {dest}" if (origin or dest) else ""
            desc = f"{service}{cargo_part}{route_part}".strip()
        amount = int(round(float(trip.get("freight_revenue") or 0)))
        items.append({
            "description": desc,
            "plate_number": str(trip.get("plate_number") or "").strip(),
            "unit": "Chuyến",
            "quantity": int(trip.get("trip_quantity") or 1),
            "price": amount,
            "amount": amount,
        })

    num_items = len(items)
    orig_item_rows = 7
    proto_tr = copy.deepcopy(t0.rows[1]._tr)

    if num_items < orig_item_rows:
        for r_idx in reversed(range(1 + num_items, 1 + orig_item_rows)):
            t0._tbl.remove(t0.rows[r_idx]._tr)
    elif num_items > orig_item_rows:
        for _ in range(num_items - orig_item_rows):
            new_tr = copy.deepcopy(proto_tr)
            summary_idx = len(t0.rows) - 3
            t0._tbl.insert(summary_idx, new_tr)

    _add_vehicle_plate_column(t0, num_items)

    subtotal = 0
    for idx, item in enumerate(items):
        row = t0.rows[1 + idx]
        subtotal += item["amount"]
        row.cells[0].text = str(idx + 1)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[1].text = item["description"]
        row.cells[2].text = item["plate_number"]
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[3].text = item["unit"]
        row.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[4].text = str(item["quantity"])
        row.cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[5].text = f"{item['price']:,}".replace(",", ".")
        row.cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        row.cells[6].text = f"{item['amount']:,}".replace(",", ".")
        row.cells[6].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    vat_rate = float(trip.get("vat_rate") if trip.get("vat_rate") is not None else 0.08)
    vat_percent = int(round(vat_rate * 100))
    vat_amount = int(round(subtotal * vat_rate))
    total_amount = subtotal + vat_amount

    # Row -3: CỘNG ĐƠN HÀNG CHƯA THUẾ GTGT
    r_sub = t0.rows[1 + num_items]
    r_sub.cells[6].text = f"{subtotal:,}".replace(",", ".")
    r_sub.cells[6].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Row -2: THUẾ GTGT
    r_vat = t0.rows[2 + num_items]
    r_vat.cells[1].text = f"THUẾ GTGT {vat_percent}%"
    r_vat.cells[6].text = f"{vat_amount:,}".replace(",", ".")
    r_vat.cells[6].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Row -1: TỔNG CỘNG ĐƠN HÀNG
    r_tot = t0.rows[3 + num_items]
    r_tot.cells[6].text = f"{total_amount:,}".replace(",", ".")
    r_tot.cells[6].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    p = doc.paragraphs[115]
    p.text = f"Bằng chữ: {number_to_vietnamese_words(total_amount)}"

    _format_items_table(t0, num_items)

    # --- 1. Populate HĐNT Signatures ---
    p74 = doc.paragraphs[74]
    p75 = doc.paragraphs[75]

    sig_tbl_hdnt = doc.add_table(rows=2, cols=2)
    sig_tbl_hdnt.autofit = False
    for row in sig_tbl_hdnt.rows:
        row.cells[0].width = Twips(5200)
        row.cells[1].width = Twips(5200)

    # Row 0: Titles
    c_a0 = sig_tbl_hdnt.rows[0].cells[0]
    p_a0 = c_a0.paragraphs[0]
    p_a0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_a0.paragraph_format.space_after = Pt(2)
    r = p_a0.add_run("ĐẠI DIỆN BÊN A\n")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    r2 = p_a0.add_run("(Bên vận chuyển)")
    r2.italic = True
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(10.5)

    c_b0 = sig_tbl_hdnt.rows[0].cells[1]
    p_b0 = c_b0.paragraphs[0]
    p_b0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_b0.paragraph_format.space_after = Pt(2)
    r = p_b0.add_run("ĐẠI DIỆN BÊN B\n")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    r2 = p_b0.add_run("(Bên thuê vận chuyển)")
    r2.italic = True
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(10.5)

    # Row 1: Signatures & Names
    c_a1 = sig_tbl_hdnt.rows[1].cells[0]
    _insert_docx_sig(c_a1, seller_sig_url, seller_stamp_url, seller_rep)

    c_b1 = sig_tbl_hdnt.rows[1].cells[1]
    _insert_docx_sig(c_b1, customer_sig_url, customer_stamp_url, customer_rep)

    p74._p.addprevious(sig_tbl_hdnt._tbl)
    p74._p.getparent().remove(p74._p)
    p75._p.getparent().remove(p75._p)

    # --- 2. Populate BBBG Signatures (Table 1) ---
    if len(doc.tables) > 1:
        t1 = doc.tables[1]
        c0 = t1.rows[0].cells[0]
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_after = Pt(2)
        r_sub = p0.add_run("\n(Ký, đóng dấu)")
        r_sub.italic = True
        r_sub.font.name = "Times New Roman"
        r_sub.font.size = Pt(10.5)

        c1 = t1.rows[0].cells[1]
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(2)
        r_sub = p1.add_run("\n(Ký, đóng dấu)")
        r_sub.italic = True
        r_sub.font.name = "Times New Roman"
        r_sub.font.size = Pt(10.5)

        r_sig = t1.add_row()
        _insert_docx_sig(r_sig.cells[0], seller_sig_url, seller_stamp_url, seller_rep)
        _insert_docx_sig(r_sig.cells[1], customer_sig_url, customer_stamp_url, customer_rep)

    return doc, contract_no


def _insert_docx_sig(cell, sig_url: str = "", stamp_url: str = "", rep_name: str = "", width_mm: float = 46.0):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)

    active_url = sig_url or stamp_url
    valid_path = None
    if active_url:
        clean = active_url.split("?")[0].lstrip("/")
        from backend.services import config as cfg
        candidates = [
            Path(cfg.BASE_DIR) / clean,
            Path(clean),
            Path(cfg.BASE_DIR) / "frontend" / clean,
        ]
        for c in candidates:
            if c.exists() and c.is_file():
                valid_path = c
                break

    if valid_path:
        try:
            run = p.add_run()
            run.add_picture(str(valid_path), width=Mm(width_mm))
        except Exception as exc:
            logger.warning("Failed to insert signature picture: %s", exc)
            p.paragraph_format.line_spacing = Pt(18)
            p.add_run("\n\n\n")
    else:
        p.paragraph_format.line_spacing = Pt(18)
        p.add_run("\n\n\n")

    if rep_name:
        p_name = cell.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_name.paragraph_format.space_before = Pt(2)
        p_name.paragraph_format.space_after = Pt(2)
        r = p_name.add_run(rep_name)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)


def _find_bbbg_cut_element(doc_obj: docx.Document):
    quoc_hieu_count = 0
    for p in doc_obj.paragraphs:
        if "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in p.text:
            quoc_hieu_count += 1
            if quoc_hieu_count == 2:
                return p._element
    return None


def render_hdnt_docx(trip: dict, customer: dict, seller: dict | None = None) -> tuple[bytes, str]:
    """Render HĐNT (Hợp đồng nguyên tắc) as docx bytes with signatures."""
    combined_doc, contract_no = _populate_template(trip, customer, seller)

    stream = io.BytesIO()
    combined_doc.save(stream)
    stream.seek(0)

    hdnt_doc = docx.Document(stream)
    body = hdnt_doc._body._element
    bbbg_elem = _find_bbbg_cut_element(hdnt_doc)
    if bbbg_elem is not None:
        idx_bbbg = list(body).index(bbbg_elem)
        for el in list(body)[idx_bbbg:]:
            body.remove(el)

    out = io.BytesIO()
    hdnt_doc.save(out)
    return out.getvalue(), contract_no


def render_bbbg_docx(trip: dict, customer: dict, seller: dict | None = None) -> tuple[bytes, str]:
    """Render BBBG (Biên bản bàn giao / giao nhận) as docx bytes with signatures."""
    combined_doc, contract_no = _populate_template(trip, customer, seller)

    stream = io.BytesIO()
    combined_doc.save(stream)
    stream.seek(0)

    bbbg_doc = docx.Document(stream)
    body = bbbg_doc._body._element
    bbbg_elem = _find_bbbg_cut_element(bbbg_doc)
    if bbbg_elem is not None:
        idx_bbbg = list(body).index(bbbg_elem)
        for el in list(body)[:idx_bbbg]:
            body.remove(el)

    out = io.BytesIO()
    bbbg_doc.save(out)
    return out.getvalue(), contract_no


def render_hdnt_bbbg_combined_docx(trip: dict, customer: dict, seller: dict | None = None) -> tuple[bytes, str]:
    """Render combined HĐNT + BBBG as docx bytes with signatures."""
    combined_doc, contract_no = _populate_template(trip, customer, seller)
    out = io.BytesIO()
    combined_doc.save(out)
    return out.getvalue(), contract_no


def compute_contract_no(trip: dict, customer: dict, seller: dict | None = None) -> str:
    """Compute standard contract number for HĐNT / BBBG."""
    raw_date = trip.get("invoice_date") or trip.get("departure_time")
    if isinstance(raw_date, datetime):
        c_date = raw_date.date()
    elif isinstance(raw_date, date):
        c_date = raw_date
    elif isinstance(raw_date, str) and raw_date:
        try:
            c_date = date.fromisoformat(raw_date[:10])
        except Exception:
            c_date = date.today()
    else:
        c_date = date.today()

    day_str = f"{c_date.day:02d}"
    month_str = f"{c_date.month:02d}"
    short_year = str(c_date.year)[-2:]

    customer_name = ((customer or {}).get("name") or "").strip()
    cust_code = derive_company_code(customer_name)

    seller_name = ((seller or {}).get("name") or "").strip()
    seller_code = derive_company_code(seller_name)
    return f"{day_str}{month_str}{short_year}/HĐNT/{cust_code}-{seller_code}"


def _format_bank_lines(banks: list[dict]) -> str:
    if not banks:
        return ""
    lines = []
    for b in banks:
        num = escape(str(b.get("account_number") or ""))
        bank_name = (b.get("bank_name") or "").strip()
        branch = (b.get("branch") or "").strip()
        label = " - ".join(p for p in (bank_name, branch) if p)
        suffix = f" tại ngân hàng {escape(label)}" if label else ""
        if num:
            lines.append(f'<div style="margin-bottom: 2px;">{num}{suffix}</div>')
    return "".join(lines)


def _render_sig_box(rep_name: str, sig_url: str = "", stamp_url: str = "") -> str:
    sig_data = _to_data_uri_if_local(sig_url) if sig_url else ""
    stamp_data = _to_data_uri_if_local(stamp_url) if stamp_url else ""

    active_data = sig_data or stamp_data
    if active_data:
        sig_space = f'<div style="height: 95px; display: flex; align-items: center; justify-content: center;"><img src="{active_data}" alt="Chữ ký / Con dấu" style="max-height: 90px; max-width: 180px; object-fit: contain;"></div>'
    else:
        sig_space = '<div style="height: 95px;"></div>'

    rep_html = f'<div style="font-weight: bold; font-size: 11pt; margin-top: 4px;">{escape(rep_name)}</div>' if rep_name else '<div style="height: 18px;"></div>'
    return f'{sig_space}{rep_html}'


def render_hdnt_bbbg_html(
    trip: dict,
    customer: dict,
    seller: dict | None = None,
    kind: str = "hdnt",
    preview: bool = True,
    invoice_guid: str | None = None,
) -> str:
    """Render HĐNT (Hợp đồng nguyên tắc) and/or BBBG (Biên bản bàn giao) as standard A4 HTML."""
    raw_date = trip.get("invoice_date") or trip.get("departure_time")
    if isinstance(raw_date, datetime):
        c_date = raw_date.date()
    elif isinstance(raw_date, date):
        c_date = raw_date
    elif isinstance(raw_date, str) and raw_date:
        try:
            c_date = date.fromisoformat(raw_date[:10])
        except Exception:
            c_date = date.today()
    else:
        c_date = date.today()

    day_str = f"{c_date.day:02d}"
    month_str = f"{c_date.month:02d}"
    year_str = str(c_date.year)
    short_year = str(c_date.year)[-2:]

    customer_name = ((customer or {}).get("name") or "").strip()
    cust_code = derive_company_code(customer_name)

    seller = seller or {}
    seller_name = (seller.get("name") or "").strip()
    seller_code = derive_company_code(seller_name)
    contract_no = trip.get("contract_no") or f"{day_str}{month_str}{short_year}/HĐNT/{cust_code}-{seller_code}"

    seller_rep = (seller.get("director_name") or "").strip()
    seller_pos = (seller.get("director_title") or "Giám đốc").strip()
    seller_addr = (seller.get("address") or "").strip()
    seller_phone = (seller.get("phone") or "").strip()
    seller_tax = (seller.get("tax_code") or "").strip()
    seller_sig_url = (seller.get("signature_url") or "").strip()
    seller_stamp_url = (seller.get("stamp_url") or "").strip()
    seller_banks = _contract_bank_accounts(seller, seller.get("bank_account") or "", seller.get("bank_name") or "")

    customer_rep = ((customer or {}).get("director_name") or "").strip()
    customer_pos = ((customer or {}).get("director_title") or "Giám đốc").strip()
    customer_addr = ((customer or {}).get("address") or "").strip()
    customer_phone = ((customer or {}).get("phone") or "").strip()
    customer_tax = ((customer or {}).get("tax_code") or "").strip()
    customer_sig_url = ((customer or {}).get("signature_url") or "").strip()
    customer_stamp_url = ((customer or {}).get("stamp_url") or "").strip()
    customer_banks = _contract_bank_accounts(customer, (customer or {}).get("bank_account") or "", (customer or {}).get("bank_name") or "")

    seller_bank_html = _format_bank_lines(seller_banks)
    customer_bank_html = _format_bank_lines(customer_banks)

    seller_gender = seller.get("director_gender") or seller.get("gender")
    customer_gender = (customer or {}).get("director_gender") or (customer or {}).get("gender")
    seller_prefix, seller_clean_name = split_rep_honorific(seller_rep, seller_gender)
    customer_prefix, customer_clean_name = split_rep_honorific(customer_rep, customer_gender)

    seller_sig_html = _render_sig_box(seller_rep, seller_sig_url, seller_stamp_url)
    customer_sig_html = _render_sig_box(customer_rep, customer_sig_url, customer_stamp_url)

    cargo_desc = (trip.get("cargo_description") or "").strip()
    segments = trip.get("route_segments") or []
    items: list[dict[str, Any]] = []
    if segments:
        for seg in segments:
            desc = seg.get("description") or ""
            if not desc:
                cargo_part = f" {cargo_desc}" if cargo_desc else ""
                service = trip.get("service_name") or "Cước vận chuyển"
                origin = seg.get("origin", "")
                dest = seg.get("destination", "")
                route_part = f" từ {origin} đi {dest}" if (origin or dest) else ""
                desc = f"{service}{cargo_part}{route_part}".strip()
            amount = int(round(float(seg.get("amount") or 0)))
            items.append({
                "description": desc,
                "plate_number": str(seg.get("plate_number") or trip.get("plate_number") or "").strip(),
                "unit": "Chuyến",
                "quantity": 1,
                "price": amount,
                "amount": amount,
            })
    else:
        desc = trip.get("trip_note") or ""
        if not desc:
            cargo_part = f" {cargo_desc}" if cargo_desc else ""
            service = trip.get("service_name") or "Cước vận chuyển"
            origin = trip.get("origin", "")
            dest = trip.get("destination", "")
            route_part = f" từ {origin} đi {dest}" if (origin or dest) else ""
            desc = f"{service}{cargo_part}{route_part}".strip()
        amount = int(round(float(trip.get("freight_revenue") or 0)))
        items.append({
            "description": desc,
            "plate_number": str(trip.get("plate_number") or "").strip(),
            "unit": "Chuyến",
            "quantity": int(trip.get("trip_quantity") or 1),
            "price": amount,
            "amount": amount,
        })

    subtotal = sum(it["amount"] for it in items)
    vat_rate = float(trip.get("vat_rate") if trip.get("vat_rate") is not None else 0.08)
    vat_percent = int(round(vat_rate * 100))
    vat_amount = int(round(subtotal * vat_rate))
    total_amount = subtotal + vat_amount
    total_in_words = number_to_vietnamese_words(total_amount)

    # BBBG table rows
    item_rows_html = []
    for idx, item in enumerate(items):
        item_rows_html.append(f"""
        <tr>
          <td style="text-align: center;">{idx + 1}</td>
          <td>{escape(item["description"])}</td>
          <td style="text-align: center; font-weight: 600;">{escape(item["plate_number"])}</td>
          <td style="text-align: center;">{escape(item["unit"])}</td>
          <td style="text-align: center;">{item["quantity"]}</td>
          <td style="text-align: right;">{f"{item['price']:,}".replace(",", ".")}</td>
          <td style="text-align: right; font-weight: 600;">{f"{item['amount']:,}".replace(",", ".")}</td>
        </tr>
        """)
    items_html_rows = "".join(item_rows_html)

    # HTML Sections
    hdnt_html = f"""
    <div class="page-sheet hdnt-section">
      <div class="header-quoc-hieu">
        <div class="country">CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</div>
        <div class="motto">Độc lập – Tự do – Hạnh phúc</div>
        <div class="divider">---------------o0o------------------</div>
      </div>

      <div class="doc-title">HỢP ĐỒNG NGUYÊN TẮC</div>
      <div class="doc-subtitle">V/v: Thỏa thuận vận chuyển hàng hóa</div>
      <div class="doc-no">Số: {contract_no}</div>

      <div class="can-cu">- Căn cứ Luật Thương mại Nước CHXHCN Việt Nam ban hành ngày 14/06/2005 và có hiệu lực từ ngày 01 tháng 01 năm 2006.</div>
      <div class="can-cu">- Căn cứ Bộ Luật Dân sự số 91/2015/QH13 của Quốc hội Nước CHXHCN Việt Nam khoá XIII, kỳ họp thứ 10 được thông qua ngày 24/11/2015 và có hiệu lực từ ngày 01/01/2017.</div>
      <div class="can-cu">- Căn cứ nhu cầu và khả năng thực tế của hai bên.</div>

      <div style="margin: 8px 0; text-align: justify;">
        Hôm nay, ngày {day_str} tháng {month_str} năm {year_str}. Chúng tôi gồm có:
      </div>

      <table class="party-table">
        <tr>
          <td colspan="2" class="party-title">Bên A: (Bên vận chuyển) <span class="uppercase">{escape(seller_name)}</span></td>
        </tr>
        <tr>
          <td class="col-label">Địa chỉ:</td>
          <td>{escape(seller_addr)}</td>
        </tr>
        <tr>
          <td class="col-label">Điện thoại:</td>
          <td>{escape(seller_phone)}</td>
        </tr>
        <tr>
          <td class="col-label">Mã số thuế:</td>
          <td><span class="bold">{escape(seller_tax)}</span></td>
        </tr>
        <tr>
          <td class="col-label">Tài khoản số:</td>
          <td>{seller_bank_html}</td>
        </tr>
        <tr>
          <td class="col-label">Đại diện:</td>
          <td>{seller_prefix}<span class="bold">{escape(seller_clean_name or seller_rep)}</span> &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Chức vụ: <span class="bold">{escape(seller_pos)}</span></td>
        </tr>

        <tr>
          <td colspan="2" class="party-title" style="padding-top: 10px !important;">BÊN B: (Bên thuê vận chuyển) <span class="uppercase">{escape(customer_name)}</span></td>
        </tr>
        <tr>
          <td class="col-label">Địa chỉ:</td>
          <td>{escape(customer_addr)}</td>
        </tr>
        <tr>
          <td class="col-label">Điện thoại:</td>
          <td>{escape(customer_phone)}</td>
        </tr>
        <tr>
          <td class="col-label">Mã số thuế:</td>
          <td><span class="bold">{escape(customer_tax)}</span></td>
        </tr>
        <tr>
          <td class="col-label">Tài khoản số:</td>
          <td>{customer_bank_html}</td>
        </tr>
        <tr>
          <td class="col-label">Đại diện:</td>
          <td>{customer_prefix}<span class="bold">{escape(customer_clean_name or customer_rep)}</span> &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Chức vụ: <span class="bold">{escape(customer_pos)}</span></td>
        </tr>
      </table>

      <div style="margin: 8px 0; text-align: justify; font-style: italic;">
        Sau khi thỏa thuận hai bên đi đến thống nhất ký kết hợp đồng nguyên tắc với các điều khoản sau:
      </div>

      <div class="section-article">Điều I: NỘI DUNG HỢP ĐỒNG</div>
      <div class="article-p">- Bên B thuê vận chuyển và Bên A nhận cung cấp, điều động phương tiện để vận chuyển hàng hóa theo yêu cầu của Bên B như sau:</div>
      <div class="article-bullet">+ Tên hàng hóa: {escape(cargo_desc) if cargo_desc else "Các mặt hàng do bên thuê vận chuyển yêu cầu."}</div>
      <div class="article-bullet">+ Đơn vị tính: Chuyến</div>
      <div class="article-bullet">+ Số lượng: Theo thông báo từng chuyến/đợt</div>
      <div class="article-bullet">+ Cự ly vận chuyển: Căn cứ theo thỏa thuận của từng lô hàng</div>
      <div class="article-bullet">+ Phương tiện vận chuyển: Ô tô</div>

      <div class="section-article">Điều II: PHƯƠNG THỨC GIAO NHẬN</div>
      <div class="article-p">- Bên A nhận hàng tại các nơi được chỉ định và giao hàng tại các địa điểm, công trình theo yêu cầu của bên B.</div>
      <div class="article-p">- Phương thức kiểm đếm số lượng hàng hóa khi giao nhận: Theo thông báo của bên B trên từng chuyến hàng.</div>
      <div class="article-p">- Bên bốc dỡ hàng hóa: Theo thông báo từng đợt.</div>

      <div class="section-article">Điều III: CƯỚC PHÍ VÀ PHƯƠNG THỨC THANH TOÁN</div>
      <div class="article-p"><strong>3.1. Cước phí:</strong> Căn cứ vào phiếu báo giá hoặc thỏa thuận cụ thể của hai bên qua điện thoại cho mỗi lô hàng.</div>
      <div class="article-p"><strong>3.2. Thanh toán:</strong></div>
      <div class="article-bullet">- Hai bên cùng thống nhất đối chiếu công nợ và xuất hóa đơn GTGT vào cuối của mỗi tháng; hoặc sau khi vận chuyển xong mỗi lô hàng, Bên A căn cứ vào các chứng từ giao nhận hàng do Bên B phát hành có sự xác nhận của bên nhận hàng hoặc thủ kho để lập bảng đối chiếu. Sau khi hai bên ký xác nhận, Bên A tiến hành xuất hóa đơn GTGT cho Bên B để làm cơ sở thanh toán.</div>
      <div class="article-bullet">- Hình thức thanh toán: Thanh toán bằng tiền mặt hoặc chuyển khoản bằng tiền Việt Nam đồng.</div>
      <div class="article-bullet">- Thời hạn thanh toán: Bên B thanh toán cước phí vận chuyển cho Bên A trong vòng 30 ngày làm việc sau khi nhận được đầy đủ các chứng từ do bên B yêu cầu.</div>

      <div class="section-article">Điều IV: TRÁCH NHIỆM CỦA MỖI BÊN</div>
      <div class="article-p"><strong>4.1. Trách nhiệm của bên A:</strong></div>
      <div class="article-bullet">- Tiếp nhận lệnh hoặc thông báo nhận hàng của Bên B. Đại diện của Bên A (tài xế) có trách nhiệm ký vận đơn hoặc biên bản giao nhận cho từng chuyến hàng với đại diện của Bên B tại các địa điểm nhận hoặc giao hàng. Phải đảm bảo phiếu giao nhận có chữ ký, xác nhận của các bên liên quan.</div>
      <div class="article-bullet">- Có trách nhiệm điều động xe nhanh chóng, kịp thời theo yêu cầu của Bên B.</div>
      <div class="article-bullet">- Đảm bảo an toàn hàng hóa trong suốt quá trình vận chuyển. Nếu để xảy ra mất mát, hư hao thì phải bồi thường thiệt hại theo giá bán ra của lô hàng đó cho Bên B, ngoại trừ các trường hợp bất khả kháng như: thiên tai, chiến tranh…</div>
      <div class="article-bullet">- Tuân thủ nội quy, hướng dẫn của Bên B và khách hàng của Bên B.</div>
      <div class="article-bullet">- Giao hàng đúng người nhận trên Biên bản giao nhận và có chữ ký xác nhận đầy đủ theo yêu cầu của Bên B, nếu vi phạm Bên A chịu toàn bộ trách nhiệm.</div>
      <div class="article-bullet">- Sau khi vận chuyển xong lô hàng, Bên A phải gửi cho Bên B bản gốc Biên bản giao nhận cùng với các hóa đơn, chứng từ có liên quan tới các lô hàng.</div>
      <div class="article-bullet">- Hóa đơn GTGT theo bộ tài chính phát hành và Bên A chịu trách nhiệm pháp lý về hóa đơn đó.</div>
      <div class="article-bullet">- Bên A phải chịu trách nhiệm nếu có sự sai trái trong việc sử dụng và phát hành hóa đơn cho Bên B.</div>

      <div class="article-p" style="margin-top: 6px;"><strong>4.2. Trách nhiệm của bên B:</strong></div>
      <div class="article-bullet">- Phát lệnh vận chuyển hoặc cử người đại diện giao hàng cho Bên A trước khi nhận hàng và phải đảm bảo đủ hàng khi đã phát lệnh hoặc thông báo tới cho Bên A.</div>
      <div class="article-bullet">- Phải làm vận đơn hoặc phiếu giao hàng cho từng chuyến hàng, ghi rõ tên hàng, số lượng và có trách nhiệm về những điều mình ghi vào vận đơn giao cho Bên A.</div>
      <div class="article-bullet">- Giao đầy đủ chứng từ pháp lý hợp pháp cần thiết của lô hàng cần vận chuyển cho Bên A.</div>
      <div class="article-bullet">- Hỗ trợ cho Bên A các thủ tục giấy tờ liên quan đến việc giao nhận hàng hóa.</div>

      <div class="section-article">Điều V: ĐIỀU KHOẢN CHUNG</div>
      <div class="article-p">- Trong quá trình thực hiện hợp đồng nếu phát sinh những thỏa thuận khác thì tiến hành lập phụ lục hợp đồng. Phụ lục hợp đồng là một bộ phận không thể tách rời của hợp đồng.</div>
      <div class="article-p">- Hai bên cam kết thực hiện nghiêm túc các điều khoản của hợp đồng. Nếu một trong hai bên gặp khó khăn phải thông báo cho bên kia biết bằng văn bản và cùng hợp tác giải quyết vấn đề. Không bên nào được tự ý chấm dứt hợp đồng trước thời hạn, nếu có phải báo trước cho bên kia ít nhất là 30 ngày và được sự nhất trí của bên kia bằng văn bản.</div>
      <div class="article-p">- Mọi vi phạm nếu không giải quyết được bằng thương lượng, thỏa thuận, sẽ đưa ra Tòa Án nơi bị đơn xét xử, quyết định của tòa án là cuối cùng và bắt buộc đối với cả hai bên, phí tòa do bên thua chịu.</div>

      <div class="section-article">Điều VI: HIỆU LỰC CỦA HỢP ĐỒNG</div>
      <div class="article-p">- Hợp đồng này có hiệu lực đến hết ngày 31/12/{year_str}</div>
      <div class="article-p">- Các thỏa thuận cũng như các sửa đổi bổ sung, được coi như các phụ lục và là một phần không thể tách rời của hợp đồng này.</div>
      <div class="article-p">- Hợp đồng được lập thành 02 (hai) bản, mỗi bên giữ 01 (một) bản có giá trị pháp lý như nhau.</div>

      <table class="sig-table">
        <tr>
          <td>
            <div class="bold" style="font-size: 11.5pt;">ĐẠI DIỆN BÊN A</div>
            <div class="italic" style="font-size: 10.5pt;">(Bên vận chuyển)</div>
            {seller_sig_html}
          </td>
          <td>
            <div class="bold" style="font-size: 11.5pt;">ĐẠI DIỆN BÊN B</div>
            <div class="italic" style="font-size: 10.5pt;">(Bên thuê vận chuyển)</div>
            {customer_sig_html}
          </td>
        </tr>
      </table>
    </div>
    """

    bbbg_html = f"""
    <div class="page-sheet bbbg-section">
      <div class="header-quoc-hieu">
        <div class="country">CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</div>
        <div class="motto">Độc lập – Tự do – Hạnh phúc</div>
        <div class="divider">---------------o0o------------------</div>
      </div>

      <div class="doc-title" style="margin-bottom: 12px;">BIÊN BẢN GIAO NHẬN HÀNG HOÁ</div>

      <div class="can-cu">- Căn cứ vào Hợp đồng nguyên tắc số: <span class="bold">{contract_no}</span> về việc thỏa thuận vận chuyển hàng hóa được ký ngày {day_str}/{month_str}/{year_str};</div>
      <div class="can-cu">- Căn cứ vào hàng hóa giao nhận;</div>

      <div style="margin: 8px 0; text-align: justify;">
        Hôm nay, Ngày {day_str} tháng {month_str} năm {year_str}, Chúng tôi gồm có:
      </div>

      <table class="party-table">
        <tr>
          <td colspan="2" class="party-title">BÊN A: (Bên vận chuyển) <span class="uppercase">{escape(seller_name)}</span></td>
        </tr>
        <tr>
          <td class="col-label">Địa chỉ:</td>
          <td>{escape(seller_addr)}</td>
        </tr>
        <tr>
          <td class="col-label">Mã số thuế:</td>
          <td><span class="bold">{escape(seller_tax)}</span></td>
        </tr>
        <tr>
          <td class="col-label">Đại diện:</td>
          <td>{seller_prefix}<span class="bold">{escape(seller_clean_name or seller_rep)}</span> &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Chức vụ: <span class="bold">{escape(seller_pos)}</span></td>
        </tr>

        <tr>
          <td colspan="2" class="party-title" style="padding-top: 8px !important;">BÊN B: (Bên thuê vận chuyển) <span class="uppercase">{escape(customer_name)}</span></td>
        </tr>
        <tr>
          <td class="col-label">Địa chỉ:</td>
          <td>{escape(customer_addr)}</td>
        </tr>
        <tr>
          <td class="col-label">Mã số thuế:</td>
          <td><span class="bold">{escape(customer_tax)}</span></td>
        </tr>
        <tr>
          <td class="col-label">Đại diện:</td>
          <td>{customer_prefix}<span class="bold">{escape(customer_clean_name or customer_rep)}</span> &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Chức vụ: <span class="bold">{escape(customer_pos)}</span></td>
        </tr>
      </table>

      <div style="margin: 10px 0 6px 0; font-style: italic;">
        Hai bên cùng nhau thống nhất số lượng, giá trị hàng hóa đã giao nhận như sau:
      </div>

      <table class="items-table">
        <thead>
          <tr>
            <th style="width: 38px;">STT</th>
            <th>TÊN HÀNG HÓA, DỊCH VỤ</th>
            <th style="width: 125px;">BIỂN KIỂM SOÁT<br>(BKS)</th>
            <th style="width: 55px;">ĐVT</th>
            <th style="width: 55px;">SỐ LƯỢNG</th>
            <th style="width: 100px;">ĐƠN GIÁ</th>
            <th style="width: 115px;">THÀNH TIỀN</th>
          </tr>
        </thead>
        <tbody>
          {items_html_rows}
          <tr class="summary-row">
            <td colspan="6" style="text-align: right; font-weight: bold; padding: 4px 8px;">CỘNG ĐƠN HÀNG CHƯA THUẾ GTGT:</td>
            <td style="text-align: right; font-weight: bold; padding: 4px 6px;">{f"{subtotal:,}".replace(",", ".")}</td>
          </tr>
          <tr class="summary-row">
            <td colspan="6" style="text-align: right; font-weight: bold; padding: 4px 8px;">THUẾ GTGT {vat_percent}%:</td>
            <td style="text-align: right; font-weight: bold; padding: 4px 6px;">{f"{vat_amount:,}".replace(",", ".")}</td>
          </tr>
          <tr class="summary-row">
            <td colspan="6" style="text-align: right; font-weight: bold; padding: 4px 8px;">TỔNG CỘNG ĐƠN HÀNG:</td>
            <td style="text-align: right; font-weight: bold; color: #b91c1c; padding: 4px 6px;">{f"{total_amount:,}".replace(",", ".")}</td>
          </tr>
        </tbody>
      </table>

      <div style="font-style: italic; margin-top: 6px; font-size: 11pt;">
        <strong>Bằng chữ:</strong> {total_in_words}
      </div>

      <div style="margin-top: 10px; text-align: justify;">
        Bên mua xác nhận Bên bán đã giao cho bên mua đúng chủng loại, số lượng, chất lượng như trên.
      </div>
      <div style="margin-top: 4px; text-align: justify;">
        Hai bên đồng ý, thống nhất ký tên. Biên bản được lập 02 bản, mỗi bên giữ 01 bản để làm chứng từ đối chiếu thanh toán, có giá trị pháp lý như nhau./.
      </div>

      <table class="sig-table">
        <tr>
          <td>
            <div class="bold" style="font-size: 11.5pt;">ĐẠI DIỆN BÊN GIAO HÀNG</div>
            <div class="italic" style="font-size: 10.5pt;">(Ký, đóng dấu)</div>
            {seller_sig_html}
          </td>
          <td>
            <div class="bold" style="font-size: 11.5pt;">ĐẠI DIỆN BÊN NHẬN HÀNG</div>
            <div class="italic" style="font-size: 10.5pt;">(Ký, đóng dấu)</div>
            {customer_sig_html}
          </td>
        </tr>
      </table>
    </div>
    """

    if kind == "hdnt":
        body_content = hdnt_html
        title_str = f"Hợp đồng nguyên tắc - {contract_no}"
    elif kind == "bbbg":
        body_content = bbbg_html
        title_str = f"Biên bản bàn giao - {contract_no}"
    else:  # "both"
        body_content = f"{hdnt_html}\n{bbbg_html}"
        title_str = f"HĐNT & BBBG - {contract_no}"

    top_bar = ""
    if preview:
        top_bar = f"""
        <div class="no-print" style="position: sticky; top: 0; z-index: 1000; background: #0f172a; color: #fff; padding: 10px 18px; display: flex; align-items: center; justify-content: space-between; gap: 12px; box-shadow: 0 4px 12px rgba(0,0,0,0.25);">
          <div style="display: flex; align-items: center; gap: 12px; flex-wrap: wrap;">
            <div style="display: flex; align-items: center; gap: 8px;">
              <span style="background: #2563eb; color: #fff; font-size: 11px; font-weight: 700; padding: 3px 8px; border-radius: 4px; text-transform: uppercase; letter-spacing: 0.5px;">Chứng từ</span>
              <span style="font-weight: 600; font-size: 13.5px; color: #e2e8f0;">Số: <strong style="color: #60a5fa;">{contract_no}</strong></span>
            </div>
            
            <div style="display: flex; align-items: center; gap: 6px;">
              <label for="kind-select" style="font-size: 12.5px; color: #94a3b8; margin: 0; white-space: nowrap;">Loại xem:</label>
              <select id="kind-select" onchange="changeDocKind(this.value)" style="background: #1e293b; color: #f8fafc; border: 1px solid #475569; padding: 5px 10px; border-radius: 5px; font-size: 13px; font-weight: 500; cursor: pointer; outline: none;">
                <option value="hdnt" {"selected" if kind == "hdnt" else ""}>Hợp đồng nguyên tắc (HĐNT)</option>
                <option value="bbbg" {"selected" if kind == "bbbg" else ""}>Biên bản bàn giao (BBBG)</option>
                <option value="both" {"selected" if kind == "both" else ""}>Cả hai (HĐNT + BBBG)</option>
              </select>
            </div>
          </div>
          
          <div style="display: flex; align-items: center; gap: 8px; flex-wrap: wrap;">
            <button type="button" onclick="window.print()" style="background: #2563eb; color: #fff; border: none; padding: 6px 14px; border-radius: 5px; font-weight: 600; cursor: pointer; font-size: 12.5px; display: inline-flex; align-items: center; gap: 5px; transition: background 0.15s;">
              <svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><path d="M6 9V2h12v7M6 18H4a2 2 0 0 1-2-2v-5a2 2 0 0 1 2-2h16a2 2 0 0 1 2 2v5a2 2 0 0 1-2 2h-2"/><rect x="6" y="14" width="12" height="8"/></svg>
              In / Lưu PDF
            </button>
            <button type="button" onclick="downloadAsset('docx')" style="background: #0284c7; color: #fff; border: none; padding: 6px 14px; border-radius: 5px; font-weight: 600; cursor: pointer; font-size: 12.5px; display: inline-flex; align-items: center; gap: 5px; transition: background 0.15s;">
              <svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"/><polyline points="14 2 14 8 20 8"/><line x1="16" y1="13" x2="8" y2="13"/><line x1="16" y1="17" x2="8" y2="17"/></svg>
              Tải Word (.docx)
            </button>
            <button type="button" onclick="downloadAsset('pdf')" style="background: #059669; color: #fff; border: none; padding: 6px 14px; border-radius: 5px; font-weight: 600; cursor: pointer; font-size: 12.5px; display: inline-flex; align-items: center; gap: 5px; transition: background 0.15s;">
              <svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4"/><polyline points="7 10 12 15 17 10"/><line x1="12" y1="15" x2="12" y2="3"/></svg>
              Tải PDF
            </button>
            <button type="button" onclick="closeDocPreview()" style="background: #475569; color: #fff; border: none; padding: 6px 12px; border-radius: 5px; font-size: 12.5px; cursor: pointer; transition: background 0.15s;">
              Đóng
            </button>
          </div>
        </div>
        <script>
          function changeDocKind(k) {{
            try {{
              var url = new URL(window.location.href);
              url.searchParams.set('kind', k);
              window.location.href = url.toString();
            }} catch(e) {{
              window.location.search = '?kind=' + k + '&view=html';
            }}
          }}

          function downloadAsset(fmt) {{
            try {{
              var url = new URL(window.location.href);
              url.searchParams.set('format', fmt);
              url.searchParams.delete('view');
              var link = document.createElement('a');
              link.href = url.toString();
              link.setAttribute('download', '');
              document.body.appendChild(link);
              link.click();
              link.remove();
            }} catch (e) {{
              var sep = window.location.href.indexOf('?') !== -1 ? '&' : '?';
              window.location.href = window.location.href + sep + 'format=' + fmt;
            }}
          }}

          function closeDocPreview() {{
            try {{
              if (window.parent && window.parent !== window && window.parent.$) {{
                window.parent.$('#documentPreviewModal').modal('hide');
                return;
              }}
            }} catch (e) {{}}
            try {{ window.close(); }} catch (e) {{}}
          }}
        </script>
        """

    return f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="utf-8">
<title>{title_str}</title>
<style>
  @page {{
    size: A4 portrait;
    margin: 14mm 16mm 14mm 16mm;
  }}
  * {{
    box-sizing: border-box;
  }}
  html, body {{
    font-family: "Times New Roman", Times, serif;
    font-size: 11pt;
    line-height: 1.35;
    color: #111827;
    margin: 0;
    padding: 0;
    background: #f1f5f9;
  }}
  .preview-wrapper {{
    max-width: 210mm;
    margin: 20px auto;
  }}
  .page-sheet {{
    background: #fff;
    padding: 16mm 18mm;
    margin-bottom: 24px;
    box-shadow: 0 4px 14px rgba(0,0,0,0.08);
    border-radius: 2px;
  }}
  .page-sheet + .page-sheet {{
    page-break-before: always;
    break-before: page;
  }}
  @media print {{
    body {{
      background: #fff !important;
    }}
    .no-print {{
      display: none !important;
    }}
    .preview-wrapper {{
      max-width: 100% !important;
      margin: 0 !important;
      box-shadow: none !important;
    }}
    .page-sheet {{
      padding: 0 !important;
      margin-bottom: 0 !important;
      box-shadow: none !important;
      border-radius: 0 !important;
    }}
    .page-sheet + .page-sheet {{
      page-break-before: always !important;
      break-before: page !important;
    }}
    .sig-table, .items-table {{
      page-break-inside: avoid !important;
      break-inside: avoid !important;
    }}
  }}
  .text-center {{ text-align: center; }}
  .text-right {{ text-align: right; }}
  .text-justify {{ text-align: justify; }}
  .bold {{ font-weight: bold; }}
  .italic {{ font-style: italic; }}
  .uppercase {{ text-transform: uppercase; }}
  
  .header-quoc-hieu {{
    text-align: center;
    line-height: 1.25;
    margin-bottom: 12px;
  }}
  .header-quoc-hieu .country {{
    font-size: 11.5pt;
    font-weight: bold;
  }}
  .header-quoc-hieu .motto {{
    font-size: 11.5pt;
    font-weight: bold;
  }}
  .header-quoc-hieu .divider {{
    font-size: 10pt;
    letter-spacing: 1px;
    margin-top: 2px;
  }}
  
  .doc-title {{
    font-size: 14pt;
    font-weight: bold;
    text-align: center;
    margin-top: 8px;
    margin-bottom: 2px;
  }}
  .doc-subtitle {{
    font-size: 11.5pt;
    font-weight: bold;
    text-align: center;
    margin-bottom: 2px;
  }}
  .doc-no {{
    font-size: 11pt;
    font-weight: bold;
    text-align: center;
    margin-bottom: 12px;
  }}
  
  .can-cu {{
    font-style: italic;
    font-size: 10.5pt;
    margin-bottom: 4px;
    text-align: justify;
    padding-left: 14px;
    text-indent: -14px;
  }}
  
  .party-table {{
    width: 100%;
    border-collapse: collapse;
    margin: 8px 0;
  }}
  .party-table td {{
    padding: 2px 0;
    vertical-align: top;
    font-size: 11pt;
  }}
  .party-title {{
    font-weight: bold;
    padding-top: 6px !important;
    padding-bottom: 2px !important;
  }}
  .col-label {{
    width: 110px;
    white-space: nowrap;
  }}
  
  .section-article {{
    font-weight: bold;
    margin-top: 10px;
    margin-bottom: 3px;
  }}
  .article-p {{
    margin: 3px 0;
    text-align: justify;
  }}
  .article-bullet {{
    margin: 2px 0;
    padding-left: 18px;
    text-indent: -18px;
    text-align: justify;
  }}
  
  .items-table {{
    width: 100%;
    border-collapse: collapse;
    margin: 12px 0 8px 0;
    font-size: 10pt;
  }}
  .items-table th, .items-table td {{
    border: 1px solid #000;
    padding: 4px 6px;
    vertical-align: middle;
  }}
  .items-table th {{
    background-color: #f8fafc;
    font-weight: bold;
    text-align: center;
    font-size: 9.5pt;
  }}
  .items-table tr.summary-row td {{
    font-weight: bold;
  }}
  
  .sig-table {{
    width: 100%;
    border-collapse: collapse;
    margin-top: 16px;
    page-break-inside: avoid;
    break-inside: avoid;
  }}
  .sig-table td {{
    width: 50%;
    text-align: center;
    vertical-align: top;
    padding: 0 10px;
  }}
</style>
</head>
<body>

{top_bar}

<div class="preview-wrapper">
  {body_content}
</div>

</body>
</html>
"""

