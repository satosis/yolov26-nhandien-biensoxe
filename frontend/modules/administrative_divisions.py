"""Vietnamese province/commune directory from the Statistics Office workbook."""

from __future__ import annotations

import hashlib
import io
import re
from threading import Lock
import unicodedata
from urllib.parse import unquote, urljoin, urlsplit

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from psycopg2.extras import execute_values

from backend.services.database import db


SOURCE_PAGE = "https://danhmuchanhchinh.nso.gov.vn/Default.aspx"
SOURCE_DECISION_URL = "https://chinhphu.vn/?classid=1&docid=214409&pageid=27160"
SOURCE_LABEL = "Cục Thống kê - Danh mục tỉnh/thành hiện hành"
COMMUNE_SOURCE_LABEL = "Danh mục xã/phường theo bảng chuyển đổi từ 01/07/2025"
_PROVINCE_CELL = re.compile(r"^(Thành phố|Tỉnh)\s+(.+?)\s*\((\d{2})\)$")
_CODE_IN_PARENS = re.compile(r"^(.+?)\s*\((\d+)\)$")
_sync_lock = Lock()




def _clean_parens(val: object) -> tuple[str, str]:
    if not val:
        return "", ""
    text = unicodedata.normalize("NFC", str(val).strip())
    match = _CODE_IN_PARENS.fullmatch(text)
    if match:
        return match.group(1).strip(), match.group(2).strip()
    return text, ""


def parse_workbook(content: bytes) -> tuple[
    list[tuple[str, str]],
    list[tuple[str, str, str]],
    list[tuple[str, str, str, str, str, str, str, str, str, str, str]],
]:
    if len(content) > 8 * 1024 * 1024 or not content.startswith(b"PK"):
        raise ValueError("Tệp danh mục từ Cục Thống kê không hợp lệ.")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = next((sheet for sheet in workbook.worksheets
                      if "không merge" in sheet.title.lower()
                      and sheet.max_column >= 3 and sheet.max_row > 3000), None)
        if sheet is None:
            raise ValueError("Không tìm thấy bảng tỉnh/xã trong tệp chính thức.")
        provinces: dict[str, str] = {}
        communes: dict[str, tuple[str, str]] = {}
        mappings: list[tuple[str, str, str, str, str, str, str, str, str, str, str]] = []
        for row in sheet.iter_rows(min_row=3, values_only=True):
            if not row or len(row) < 4:
                continue
            province_cell, commune_cell, code_cell = row[:3]
            if not all((province_cell, commune_cell, code_cell)):
                continue
            match = _PROVINCE_CELL.fullmatch(unicodedata.normalize("NFC", str(province_cell).strip()))
            commune_code = str(code_cell).strip().zfill(5)
            commune_name = unicodedata.normalize("NFC", str(commune_cell).strip())
            if commune_code == "24496" and commune_name == "Xa Ea Kly":
                # The mapping workbook omits the accent here; QD 19/2025 gives "Xã Ea Kly".
                commune_name = "Xã Ea Kly"
            if not match or not re.fullmatch(r"\d{5}", commune_code):
                continue
            if not commune_name.startswith(("Xã ", "Phường ", "Đặc khu ")):
                raise ValueError(f"Đơn vị cấp xã không hợp lệ: {commune_name}")
            province_name = f"{match.group(1)} {match.group(2)}"
            province_code = match.group(3)
            if province_code in provinces and provinces[province_code] != province_name:
                raise ValueError(f"Mã tỉnh bị trùng tên: {province_code}")
            previous = communes.get(commune_code)
            if previous and previous != (province_code, commune_name):
                plain = lambda text: "".join(char for char in unicodedata.normalize("NFD", text)
                                              if not unicodedata.combining(char))
                if previous[0] != province_code or plain(previous[1]) != plain(commune_name):
                    raise ValueError(f"Mã xã bị trùng dữ liệu: {commune_code}")
            provinces[province_code] = province_name
            communes.setdefault(commune_code, (province_code, commune_name))

            # Parse old commune conversion mapping
            old_commune_raw = row[3]
            if old_commune_raw:
                old_commune_name = unicodedata.normalize("NFC", str(old_commune_raw).strip())
                old_commune_code = str(row[4]).strip().zfill(5) if len(row) > 4 and row[4] else ""
                note = unicodedata.normalize("NFC", str(row[5]).strip()) if len(row) > 5 and row[5] else ""
                old_dist_name, old_dist_code = _clean_parens(row[6]) if len(row) > 6 else ("", "")
                old_prov_name, old_prov_code = _clean_parens(row[7]) if len(row) > 7 else ("", "")
                mappings.append((
                    province_code, province_name,
                    commune_code, commune_name,
                    old_prov_code, old_prov_name,
                    old_dist_code, old_dist_name,
                    old_commune_code, old_commune_name,
                    note,
                ))

        if len(provinces) != 34 or len(communes) < 3300:
            raise ValueError("Danh mục tải về thiếu tỉnh/xã; chưa cập nhật dữ liệu.")
        return (
            sorted(provinces.items()),
            sorted((code, province_code, name)
                   for code, (province_code, name) in communes.items()),
            mappings,
        )
    finally:
        workbook.close()


def parse_current_provinces(content: bytes) -> list[tuple[str, str]]:
    """Read the 34 current province-level names from NSO's public directory."""
    soup = BeautifulSoup(content, "html.parser")
    provinces: dict[str, str] = {}
    for row in soup.select('tr[id^="ctl00_PlaceHolderMain_grid1N_DXDataRow"]'):
        cells = row.find_all("td", recursive=False)
        if len(cells) < 3:
            continue
        code = cells[1].get_text(" ", strip=True)
        name = unicodedata.normalize("NFC", cells[2].get_text(" ", strip=True))
        if not re.fullmatch(r"\d{2}", code) or not name.startswith(("Tỉnh ", "Thành phố ")):
            raise ValueError("Danh mục tỉnh/thành từ Cục Thống kê không hợp lệ.")
        if code in provinces:
            raise ValueError(f"Mã tỉnh/thành bị trùng: {code}")
        provinces[code] = name
    if len(provinces) != 34:
        raise ValueError("Danh mục tỉnh/thành hiện hành tải về không đủ 34 đơn vị.")
    return sorted(provinces.items())


def discover_workbook_url(content: bytes) -> str:
    """Resolve the current conversion workbook linked by NSO's directory page."""
    soup = BeautifulSoup(content, "html.parser")
    for link in soup.select("a[href]"):
        label = "".join(char for char in unicodedata.normalize("NFD", link.get_text(" ", strip=True))
                        if not unicodedata.combining(char)).casefold().replace("đ", "d")
        if "chuyen doi" not in label or "dvhc" not in label:
            continue
        url = urljoin(SOURCE_PAGE, link["href"])
        parts = urlsplit(url)
        if (parts.scheme == "https"
                and parts.hostname == urlsplit(SOURCE_PAGE).hostname
                and unquote(parts.path).casefold().endswith(".xlsx")):
            return url
    raise ValueError("Không tìm thấy liên kết Excel chuyển đổi ĐVHC trên trang Cục Thống kê.")


def sync_from_official_source() -> dict:
    """Combine current province names with the official commune conversion table."""
    with _sync_lock:
        with requests.get(SOURCE_PAGE, timeout=(8, 35)) as response:
            response.raise_for_status()
            if len(response.content) > 2 * 1024 * 1024:
                raise ValueError("Trang danh mục tỉnh/thành vượt quá kích thước cho phép.")
            source_page_content = response.content
        provinces = parse_current_provinces(source_page_content)
        workbook_url = discover_workbook_url(source_page_content)
        with requests.get(workbook_url, timeout=(8, 35), stream=True) as response:
            response.raise_for_status()
            content = bytearray()
            for part in response.iter_content(chunk_size=65536):
                content.extend(part)
                if len(content) > 8 * 1024 * 1024:
                    raise ValueError("Tệp danh mục vượt quá kích thước cho phép.")
        content = bytes(content)
        old_provinces, communes, mappings = parse_workbook(content)
        if {code for code, _ in provinces} != {code for code, _ in old_provinces}:
            raise ValueError("Mã tỉnh/thành hiện hành không khớp bảng chuyển đổi xã/phường.")
        digest = hashlib.sha256(content + repr(provinces).encode("utf-8")).hexdigest()
        with db._get_connection() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT source_hash FROM public.administrative_sync_state WHERE id = 1")
                previous = cur.fetchone()
                changed = previous is None or previous[0] != digest
                execute_values(cur, """
                    INSERT INTO public.administrative_provinces (code, name) VALUES %s
                    ON CONFLICT (code) DO UPDATE SET
                        name = EXCLUDED.name, is_active = TRUE, updated_at = NOW()
                """, provinces)
                execute_values(cur, """
                    INSERT INTO public.administrative_communes (code, province_code, name) VALUES %s
                    ON CONFLICT (code) DO UPDATE SET
                        province_code = EXCLUDED.province_code,
                        name = EXCLUDED.name, is_active = TRUE, updated_at = NOW()
                """, communes)
                cur.execute("TRUNCATE TABLE public.administrative_commune_mappings RESTART IDENTITY")
                execute_values(cur, """
                    INSERT INTO public.administrative_commune_mappings (
                        new_province_code, new_province_name,
                        new_commune_code, new_commune_name,
                        old_province_code, old_province_name,
                        old_district_code, old_district_name,
                        old_commune_code, old_commune_name,
                        note
                    ) VALUES %s
                """, mappings)
                cur.execute("UPDATE public.administrative_communes SET is_active = FALSE, updated_at = NOW() WHERE NOT (code = ANY(%s))",
                            ([row[0] for row in communes],))
                cur.execute("UPDATE public.administrative_provinces SET is_active = FALSE, updated_at = NOW() WHERE NOT (code = ANY(%s))",
                            ([row[0] for row in provinces],))
                cur.execute("""
                    INSERT INTO public.administrative_sync_state
                        (id, source_url, source_hash, province_count, commune_count, mapping_count, synced_at)
                    VALUES (1, %s, %s, %s, %s, %s, NOW())
                    ON CONFLICT (id) DO UPDATE SET
                        source_url = EXCLUDED.source_url, source_hash = EXCLUDED.source_hash,
                        province_count = EXCLUDED.province_count,
                        commune_count = EXCLUDED.commune_count,
                        mapping_count = EXCLUDED.mapping_count,
                        synced_at = NOW()
                """, (workbook_url, digest, len(provinces), len(communes), len(mappings)))
            conn.commit()
        return {"changed": changed, "provinces": len(provinces), "communes": len(communes), "mappings": len(mappings)}




def list_provinces() -> list[dict]:
    return db._execute_query("""
        SELECT p.code, p.name, COUNT(c.code)::integer AS commune_count
        FROM public.administrative_provinces p
        LEFT JOIN public.administrative_communes c
            ON c.province_code = p.code AND c.is_active
        WHERE p.is_active GROUP BY p.code, p.name ORDER BY p.code
    """, fetchall=True, commit=False) or []


def list_communes(province_code: str = "", query: str = "", page: int = 1, limit: int = 20) -> tuple[list[dict], int]:
    conditions = ["c.is_active", "p.is_active"]
    params: list = []
    if province_code:
        conditions.append("c.province_code = %s")
        params.append(province_code)
    if query:
        conditions.append("(c.name ILIKE %s OR c.code LIKE %s)")
        params.extend((f"%{query}%", f"%{query}%"))
    where = " AND ".join(conditions)
    count = db._execute_query(f"""
        SELECT COUNT(*) AS total FROM public.administrative_communes c
        JOIN public.administrative_provinces p ON p.code = c.province_code
        WHERE {where}
    """, tuple(params), fetch=True, commit=False)["total"]
    rows = db._execute_query(f"""
        SELECT c.code, c.name, p.code AS province_code, p.name AS province_name
        FROM public.administrative_communes c
        JOIN public.administrative_provinces p ON p.code = c.province_code
        WHERE {where} ORDER BY p.code, c.name, c.code LIMIT %s OFFSET %s
    """, (*params, limit, (page - 1) * limit), fetchall=True, commit=False) or []
    return rows, count


def sync_state() -> dict | None:
    return db._execute_query("""
        SELECT *, synced_at AT TIME ZONE 'Asia/Ho_Chi_Minh' AS synced_local
        FROM public.administrative_sync_state WHERE id = 1
    """,
                             fetch=True, commit=False)
