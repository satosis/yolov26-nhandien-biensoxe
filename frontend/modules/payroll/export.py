"""Fill the original multi-sheet salary workbook with a selected month's data."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import re
import warnings
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string


TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "assets" / "staff_payroll_template.xlsx"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS = {"m": MAIN_NS}


def _number(value) -> int | float:
    number = Decimal(value)
    return int(number) if number == number.to_integral_value() else float(number)


def _fill_employee(book, row: dict, salary_row: int, attendance_row: int,
                   bonus_row: int, payroll_row: int, tax_row: int) -> None:
    salary = book.worksheets[0]
    attendance = book["BCC"]
    bonus = book.worksheets[2]
    payroll = book["Bảng lương"]
    tax = book["Thuế TNCN"]

    for col, key in (
        ("B", "full_name"), ("C", "job_title"), ("D", "base_salary"),
        ("E", "phone_allowance"), ("F", "responsibility_allowance"),
        ("G", "meal_allowance"), ("I", "insurance_salary"),
    ):
        value = row[key]
        salary[f"{col}{salary_row}"] = value if isinstance(value, str) else _number(value)

    # The app stores a monthly total, not the individual attendance marks.
    # Remove the sample day's marks so they cannot be mistaken for this month.
    for work_row in (attendance_row, attendance_row + 1):
        for col in range(5, 36):
            attendance.cell(work_row, col).value = None
    attendance[f"AJ{attendance_row}"] = _number(row["work_days"])
    attendance[f"AY{attendance_row}"] = "Tổng công tháng từ hệ thống; chưa có chấm công từng ngày"

    for col, key in (("D", "advance_amount"), ("E", "bonus_amount"), ("F", "trip_wage")):
        bonus[f"{col}{bonus_row}"] = _number(row[key])
    bonus[f"G{bonus_row}"] = _number(row["advance_amount"] + row["bonus_amount"] + row["trip_wage"])
    bonus[f"H{bonus_row}"] = bonus[f"G{bonus_row}"].value
    bonus[f"I{bonus_row}"] = 0

    insurance = Decimal(row["insurance_salary"]) if row["has_month_entry"] else Decimal(0)
    payroll_values = {
        "B": row["full_name"], "C": row["job_title"],
        "D": row["base_salary"], "E": row["responsibility_allowance"],
        "F": row["meal_allowance"], "G": row["advance_amount"],
        "H": row["trip_wage"], "I": row["bonus_amount"],
        "J": row["phone_allowance"], "K": row["work_days"],
        "L": row["standard_days"], "M": row["gross_pay"],
        "N": insurance, "O": insurance * Decimal("0.175"),
        "P": insurance * Decimal("0.03"), "Q": insurance * Decimal("0.01"),
        "R": row["employer_insurance"], "S": insurance * Decimal("0.08"),
        "T": insurance * Decimal("0.015"), "U": insurance * Decimal("0.01"),
        "V": row["employee_insurance"], "W": row["tax_deduction"],
        "X": row["advance_amount"],
        "Y": row["other_deduction"] + row["union_fee"],
        "Z": row["net_pay"], "AF": row["overtime_pay"],
        "AH": 0, "AI": row["overtime_pay"],
        "AK": row["work_factor"],
    }
    for col, value in payroll_values.items():
        payroll[f"{col}{payroll_row}"] = value if isinstance(value, str) else _number(value)
    payroll[f"AJ{payroll_row}"] = row["note"] or None
    tax[f"R{tax_row}"] = _number(row["tax_deduction"])


def _update_payslip(book) -> None:
    """Keep the source payslip's employee selector and align it with monthly pay."""
    sheet = book["Phiếu lương"]
    # Its original income cells refer to full monthly rates rather than
    # prorated values; reference the pay factor in the main sheet instead.
    source_cols = {9: "D", 10: "E", 11: "F", 12: "J"}
    for payslip_row, salary_col in source_cols.items():
        sheet[f"C{payslip_row}"] = (
            f'=IF($B$4="","",IFERROR(INDEX(\'Bảng lương\'!${salary_col}$10:${salary_col}$12,'
            f'MATCH($B$4,\'Bảng lương\'!$B$10:$B$12,0))*'
            f'INDEX(\'Bảng lương\'!$AK$10:$AK$12,'
            f'MATCH($B$4,\'Bảng lương\'!$B$10:$B$12,0)),0))'
        )
    sheet["C13"] = '=IF($B$4="","",IFERROR(INDEX(\'Bảng lương\'!$H$10:$H$12,MATCH($B$4,\'Bảng lương\'!$B$10:$B$12,0))+INDEX(\'Bảng lương\'!$I$10:$I$12,MATCH($B$4,\'Bảng lương\'!$B$10:$B$12,0)),0))'
    sheet["C14"] = '=IF($B$4="","",IFERROR(INDEX(\'Bảng lương\'!$AF$10:$AF$12,MATCH($B$4,\'Bảng lương\'!$B$10:$B$12,0)),0))'
    sheet["F13"] = '=IF($B$4="","",IFERROR(INDEX(\'Bảng lương\'!$X$10:$X$12,MATCH($B$4,\'Bảng lương\'!$B$10:$B$12,0))+INDEX(\'Bảng lương\'!$Y$10:$Y$12,MATCH($B$4,\'Bảng lương\'!$B$10:$B$12,0)),0))'
    sheet["F16"] = 0  # Overtime pay is already included in C14.


def _cell_data(cell, value) -> None:
    """Replace a value in source XML while retaining its original cell style."""
    for child in list(cell):
        if child.tag in (f"{{{MAIN_NS}}}f", f"{{{MAIN_NS}}}v", f"{{{MAIN_NS}}}is"):
            cell.remove(child)
    cell.attrib.pop("t", None)
    if value is None:
        return
    if isinstance(value, str) and value.startswith("="):
        etree.SubElement(cell, f"{{{MAIN_NS}}}f").text = value[1:]
    elif isinstance(value, str):
        cell.set("t", "inlineStr")
        inline = etree.SubElement(cell, f"{{{MAIN_NS}}}is")
        etree.SubElement(inline, f"{{{MAIN_NS}}}t").text = value
    else:
        etree.SubElement(cell, f"{{{MAIN_NS}}}v").text = str(value)


def _patch_sheet(xml: bytes, values: dict[str, object], *, clear_formula_caches: bool) -> bytes:
    root = etree.fromstring(xml)
    data = root.find("m:sheetData", NS)
    if data is None:
        return xml
    row_nodes = {int(row.get("r")): row for row in data.findall("m:row", NS)}
    for coord, value in values.items():
        col, row_num = coordinate_from_string(coord)
        row = row_nodes.get(row_num)
        if row is None:
            row = etree.Element(f"{{{MAIN_NS}}}row", r=str(row_num))
            after = next((node for node in data if node.tag == f"{{{MAIN_NS}}}row"
                          and int(node.get("r")) > row_num), None)
            data.insert(data.index(after) if after is not None else len(data), row)
            row_nodes[row_num] = row
        cell = next((node for node in row if node.tag == f"{{{MAIN_NS}}}c"
                     and node.get("r") == coord), None)
        if cell is None:
            cell = etree.Element(f"{{{MAIN_NS}}}c", r=coord)
            col_index = column_index_from_string(col)
            after = next((node for node in row if node.tag == f"{{{MAIN_NS}}}c"
                          and column_index_from_string(coordinate_from_string(node.get("r"))[0]) > col_index), None)
            row.insert(row.index(after) if after is not None else len(row), cell)
        _cell_data(cell, value)
    if clear_formula_caches:
        for cell in root.findall(".//m:sheetData/m:row/m:c", NS):
            if cell.find("m:f", NS) is not None:
                cached = cell.find("m:v", NS)
                if cached is not None:
                    cell.remove(cached)
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True)


def _patch_workbook(xml: bytes) -> bytes:
    root = etree.fromstring(xml)
    calc = root.find("m:calcPr", NS)
    if calc is None:
        calc = etree.SubElement(root, f"{{{MAIN_NS}}}calcPr")
    calc.set("calcMode", "auto")
    calc.set("fullCalcOnLoad", "1")
    calc.set("forceFullCalc", "1")
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True)


def _save_without_changing_template(book, before: list[dict]) -> bytes:
    """Preserve the source ZIP parts, including Excel extensions unsupported by openpyxl."""
    changes = []
    for sheet, old in zip(book.worksheets, before):
        changes.append({cell.coordinate: cell.value for cell in sheet._cells.values()
                        if cell.value != old.get(cell.coordinate)})
    output = BytesIO()
    with ZipFile(TEMPLATE_PATH) as source, ZipFile(output, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", info.filename)
            if match:
                index = int(match.group(1)) - 1
                data = _patch_sheet(data, changes[index], clear_formula_caches=True)
            elif info.filename == "xl/workbook.xml":
                data = _patch_workbook(data)
            target.writestr(info, data)
    return output.getvalue()


def build_staff_payroll_workbook(month_start: date, rows: list[dict]) -> bytes:
    """Return the supplied workbook with only month and employee data replaced."""
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        book = load_workbook(TEMPLATE_PATH)
    if len(rows) > 2:
        raise ValueError("File mẫu hiện có hai dòng nhân sự; chưa thể xuất thêm người.")
    before = [{cell.coordinate: cell.value for cell in sheet._cells.values()}
              for sheet in book.worksheets]

    attendance = book["BCC"]
    attendance["O4"] = month_start.month
    attendance["O5"] = month_start.year
    attendance["AJ5"] = _number(rows[0]["standard_days"]) if rows else 26
    slots = [(8, 9, 5, 10, 5), (10, 12, 7, 12, 7)]
    for row, slot in zip(rows, slots):
        _fill_employee(book, row, *slot)
    payroll = book["Bảng lương"]
    total_columns = ("D", "E", "F", "G", "H", "I", "J", "K", "M", "N", "O", "P", "Q", "R",
                     "S", "T", "U", "V", "W", "X", "Y", "Z", "AF")
    for col in total_columns:
        payroll[f"{col}13"] = sum(Decimal(str(payroll[f"{col}{r}"].value or 0)) for r in (10, 12))
    _update_payslip(book)
    return _save_without_changing_template(book, before)
