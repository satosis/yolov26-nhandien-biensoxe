"""Router for Tax and Accounting Work Checklist & Dashboard."""

from __future__ import annotations

import io
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, Form, Query, Request
from fastapi.responses import HTMLResponse, JSONResponse, Response, RedirectResponse

from frontend.deps import get_menu_context, require_menu, templates
from frontend.modules.enterprise import tax_checklist_service as service

router = APIRouter(prefix="/enterprise/tax-checklist", tags=["TaxChecklist"])


@router.get("", include_in_schema=False)
@router.get("/")
async def tax_checklist_page(
    request: Request,
    period: str = Query("2026-Q3"),
    view: str = Query("overview"),  # 'overview' (Image 2) or 'matrix' (Image 1)
    q: Optional[str] = Query(None),
    user: dict = Depends(require_menu("tax_checklist")),
):
    ctx = get_menu_context(request, user)

    # Fetch checklist data
    data = service.get_checklist_matrix(period, q_search=q)
    periods = service.get_available_periods()
    active_period = next((p for p in periods if p["key"] == period), periods[1] if len(periods) > 1 else periods[0])

    urgent_tasks = service.get_urgent_tasks(period)
    notifications = service.get_notifications(period)

    ctx.update({
        "request": request,
        "current_menu": "tax_checklist",
        "current_submenu": "tax_checklist",
        "page_title": "Kiểm soát công việc Kế toán & Thuế",
        "period": period,
        "active_period": active_period,
        "periods": periods,
        "view": view,
        "q": q or "",
        "clients": data["clients"],
        "column_totals": data["column_totals"],
        "overall_stats": data["overall_stats"],
        "groups": data["groups"],
        "all_items": data["all_items"],
        "urgent_tasks": urgent_tasks,
        "notifications": notifications,
    })

    return templates.TemplateResponse("pages/enterprise/tax_checklist.html", ctx)


@router.post("/cycle-status")
async def cycle_status_api(
    request: Request,
    client_id: int = Form(...),
    period: str = Form(...),
    item_key: str = Form(...),
    user: dict = Depends(require_menu("tax_checklist")),
):
    """Toggle item status to next in cycle."""
    res = service.cycle_item_status(client_id, period, item_key)
    return JSONResponse(res)


@router.post("/cycle-group-status")
async def cycle_group_status_api(
    request: Request,
    client_id: int = Form(...),
    period: str = Form(...),
    group_id: str = Form(...),
    user: dict = Depends(require_menu("tax_checklist")),
):
    """Toggle all items in a group or single group-level status."""
    res = service.cycle_group_status(client_id, period, group_id)
    return JSONResponse(res)


@router.post("/set-status")
async def set_status_api(
    request: Request,
    client_id: int = Form(...),
    period: str = Form(...),
    item_key: str = Form(...),
    status: str = Form(...),
    user: dict = Depends(require_menu("tax_checklist")),
):
    """Set explicit item status."""
    res = service.set_item_status(client_id, period, item_key, status)
    return JSONResponse(res)


@router.post("/update-note")
async def update_note_api(
    request: Request,
    client_id: int = Form(...),
    period: str = Form(...),
    note: str = Form(""),
    user: dict = Depends(require_menu("tax_checklist")),
):
    """Update general note for client."""
    service.update_client_note(client_id, period, note)
    return JSONResponse({"ok": True})


@router.post("/client/add")
async def add_client_api(
    request: Request,
    code: str = Form(...),
    name: str = Form(...),
    tax_code: str = Form(""),
    contact_info: str = Form(""),
    note: str = Form(""),
    period: str = Form("2026-Q3"),
    view: str = Form("overview"),
    user: dict = Depends(require_menu("tax_checklist")),
):
    """Add a new client and redirect back."""
    res = service.add_client(code, name, tax_code, contact_info, note)
    redir_url = f"/admin/enterprise/tax-checklist?period={period}&view={view}"
    if not res.get("ok"):
        return RedirectResponse(f"{redir_url}&error={quote(res.get('message', 'Lỗi'))}", status_code=303)
    return RedirectResponse(f"{redir_url}&success={quote('Thêm khách hàng thành công!')}", status_code=303)


@router.post("/task/toggle")
async def toggle_task_api(
    request: Request,
    task_id: int = Form(...),
    user: dict = Depends(require_menu("tax_checklist")),
):
    """Toggle task completion."""
    service.toggle_urgent_task(task_id)
    return JSONResponse({"ok": True})


@router.post("/task/add")
async def add_task_api(
    request: Request,
    period: str = Form(...),
    title: str = Form(...),
    client_name: str = Form(...),
    deadline: str = Form(...),
    user: dict = Depends(require_menu("tax_checklist")),
):
    """Add urgent task."""
    service.add_urgent_task(period, title, client_name, deadline)
    return RedirectResponse(f"/admin/enterprise/tax-checklist?period={period}&view=overview&success={quote('Đã thêm việc cần xử lý!')}", status_code=303)


@router.get("/export-excel")
async def export_excel_api(
    period: str = Query("2026-Q3"),
    user: dict = Depends(require_menu("tax_checklist")),
):
    """Export checklist matrix as beautiful formatted Excel spreadsheet."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    data = service.get_checklist_matrix(period)
    clients = data["clients"]
    groups = data["groups"]
    all_items = data["all_items"]
    column_totals = data["column_totals"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Checklist_{period}"

    # Palette
    font_main = Font(name="Times New Roman", size=11)
    font_title = Font(name="Times New Roman", size=15, bold=True, color="FFFFFF")
    font_group = Font(name="Times New Roman", size=11, bold=True)
    font_header = Font(name="Times New Roman", size=10, bold=True)
    font_done = Font(name="Segoe UI Symbol", size=12, bold=True, color="16A34A")
    font_not_done = Font(name="Segoe UI Symbol", size=11, bold=True, color="DC2626")
    font_na = Font(name="Times New Roman", size=10, color="9CA3AF")

    fill_title = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_fixed = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_summary_row = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    group_fills = {
        "gtgt": PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"),
        "tncn": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        "tndn": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "sao_ke": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
        "hop_dong": PatternFill(start_color="CFFAFE", end_color="CFFAFE", fill_type="solid"),
        "so_sach": PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid"),
    }

    thin = Side(border_style="thin", color="CBD5E1")
    thick = Side(border_style="medium", color="64748B")
    box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 1. Title Banner
    total_cols = 3 + len(all_items) + 3
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    cell_title = ws.cell(row=2, column=1, value=f"CHECKLIST KIỂM SOÁT CÔNG VIỆC KẾ TOÁN – THUẾ ({period})")
    cell_title.font = font_title
    cell_title.fill = fill_title
    cell_title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 36

    # 2. Row 3: Group Headers
    ws.merge_cells("A3:A4")
    ws.cell(row=3, column=1, value="STT").alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row=3, column=1).fill = fill_fixed
    ws.cell(row=3, column=1).font = font_group

    ws.merge_cells("B3:B4")
    ws.cell(row=3, column=2, value="Mã KH").alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=3, column=2).fill = fill_fixed
    ws.cell(row=3, column=2).font = font_group

    ws.merge_cells("C3:C4")
    ws.cell(row=3, column=3, value="Tên khách hàng").alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=3, column=3).fill = fill_fixed
    ws.cell(row=3, column=3).font = font_group

    cur_col = 4
    for g in groups:
        g_len = len(g["items"])
        end_col = cur_col + g_len - 1
        ws.merge_cells(start_row=3, start_column=cur_col, end_row=3, end_column=end_col)
        g_cell = ws.cell(row=3, column=cur_col, value=g["name"])
        g_cell.alignment = Alignment(horizontal="center", vertical="center")
        g_cell.fill = group_fills.get(g["id"], fill_fixed)
        g_cell.font = font_group

        # Row 4: Item names
        for it in g["items"]:
            it_cell = ws.cell(row=4, column=cur_col, value=it["name"])
            it_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            it_cell.fill = group_fills.get(g["id"], fill_fixed)
            it_cell.font = font_header
            cur_col += 1

    # Final summary columns
    ws.merge_cells(start_row=3, start_column=cur_col, end_row=4, end_column=cur_col)
    c_tot = ws.cell(row=3, column=cur_col, value="Tổng hoàn thành")
    c_tot.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c_tot.fill = fill_fixed
    c_tot.font = font_group
    cur_col += 1

    ws.merge_cells(start_row=3, start_column=cur_col, end_row=4, end_column=cur_col)
    c_pct = ws.cell(row=3, column=cur_col, value="Tỷ lệ %")
    c_pct.alignment = Alignment(horizontal="center", vertical="center")
    c_pct.fill = fill_fixed
    c_pct.font = font_group
    cur_col += 1

    ws.merge_cells(start_row=3, start_column=cur_col, end_row=4, end_column=cur_col)
    c_note = ws.cell(row=3, column=cur_col, value="Ghi chú")
    c_note.alignment = Alignment(horizontal="center", vertical="center")
    c_note.fill = fill_fixed
    c_note.font = font_group

    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 42

    # 3. Data Rows
    row_idx = 5
    for idx, c in enumerate(clients, 1):
        ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_idx, column=2, value=c["code"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_idx, column=3, value=c["name"]).alignment = Alignment(horizontal="left", vertical="center")

        col_pos = 4
        for it in all_items:
            k = it["key"]
            st = c["statuses"].get(k, "not_done")
            val = "✔" if st == "done" else ("⏳" if st == "in_progress" else ("N/A" if st == "na" else "✘"))
            fnt = font_done if st == "done" else (font_na if st == "na" else font_not_done)
            cell = ws.cell(row=row_idx, column=col_pos, value=val)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = fnt
            col_pos += 1

        # Summary columns
        ws.cell(row=row_idx, column=col_pos, value=f"{c['done_count']}/{c['total_items']}").alignment = Alignment(horizontal="center", vertical="center")
        col_pos += 1

        ws.cell(row=row_idx, column=col_pos, value=f"{c['percentage']}%").alignment = Alignment(horizontal="center", vertical="center")
        col_pos += 1

        ws.cell(row=row_idx, column=col_pos, value=c["note"]).alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[row_idx].height = 26
        row_idx += 1

    # 4. Bottom Total Rows
    # Row: Tổng hoàn thành theo cột
    ws.cell(row=row_idx, column=1, value="").fill = fill_summary_row
    ws.cell(row=row_idx, column=2, value="").fill = fill_summary_row
    lbl1 = ws.cell(row=row_idx, column=3, value="Tổng hoàn thành")
    lbl1.alignment = Alignment(horizontal="right", vertical="center")
    lbl1.font = font_group
    lbl1.fill = fill_summary_row

    col_pos = 4
    for it in all_items:
        k = it["key"]
        tot = column_totals.get(k, {})
        c_v = ws.cell(row=row_idx, column=col_pos, value=f"{tot.get('done_count', 0)}/{tot.get('total', 0)}")
        c_v.alignment = Alignment(horizontal="center", vertical="center")
        c_v.font = font_group
        c_v.fill = fill_summary_row
        col_pos += 1

    ws.cell(row=row_idx, column=col_pos, value="-").fill = fill_summary_row
    ws.cell(row=row_idx, column=col_pos + 1, value="-").fill = fill_summary_row
    ws.cell(row=row_idx, column=col_pos + 2, value="-").fill = fill_summary_row
    ws.row_dimensions[row_idx].height = 26
    row_idx += 1

    # Row: Tỷ lệ hoàn thành %
    ws.cell(row=row_idx, column=1, value="").fill = fill_summary_row
    ws.cell(row=row_idx, column=2, value="").fill = fill_summary_row
    lbl2 = ws.cell(row=row_idx, column=3, value="Tỷ lệ hoàn thành")
    lbl2.alignment = Alignment(horizontal="right", vertical="center")
    lbl2.font = font_group
    lbl2.fill = fill_summary_row

    col_pos = 4
    for it in all_items:
        k = it["key"]
        tot = column_totals.get(k, {})
        c_v = ws.cell(row=row_idx, column=col_pos, value=f"{tot.get('pct', 0)}%")
        c_v.alignment = Alignment(horizontal="center", vertical="center")
        c_v.font = font_group
        c_v.fill = fill_summary_row
        col_pos += 1

    ws.cell(row=row_idx, column=col_pos, value="-").fill = fill_summary_row
    ws.cell(row=row_idx, column=col_pos + 1, value="-").fill = fill_summary_row
    ws.cell(row=row_idx, column=col_pos + 2, value="-").fill = fill_summary_row
    ws.row_dimensions[row_idx].height = 26

    # Borders for all cells in table
    for r in range(3, row_idx + 1):
        for c_i in range(1, total_cols + 1):
            cell = ws.cell(row=r, column=c_i)
            cell.border = box_border
            if not cell.font or cell.font.name != "Segoe UI Symbol":
                curr_font = cell.font or font_main
                cell.font = Font(name="Times New Roman", size=curr_font.size, bold=curr_font.bold, color=curr_font.color)

    # Column widths
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 11
    ws.column_dimensions["C"].width = 24
    for c_idx in range(4, 4 + len(all_items)):
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = 13
    ws.column_dimensions[get_column_letter(4 + len(all_items))].width = 16
    ws.column_dimensions[get_column_letter(5 + len(all_items))].width = 12
    ws.column_dimensions[get_column_letter(6 + len(all_items))].width = 24

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"Checklist_Ketoan_Thue_{period}.xlsx"
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
